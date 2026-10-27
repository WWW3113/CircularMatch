#!/usr/bin/env python3
"""Expand the four reports into parameter-level evidence without new inference."""
from __future__ import annotations
import argparse,csv,hashlib,json,math,statistics as st
from pathlib import Path
from collections import defaultdict
from datetime import datetime
from zoneinfo import ZoneInfo
import xlsxwriter
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
OLD=ROOT/'results/professor_reports_20261026'
OUT=ROOT/'results/parameter_details_20261028'
TITLES={1:'匹配容差與弱配對補救',2:'DTM平滑的局部改善與代價',3:'加入多尺度幾何微調',4:'多候選與垂直結構驗證'}
SCENES=['Plot1','Plot2','Plot3','Plot4','ETH']
PAPER_EP=dict(zip(SCENES,[.052558,.114808,.048756,.043560,.153427]))
MISSING='待補'

def digest(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def finite(x):
 try:
  f=float(x);return f if math.isfinite(f) else None
 except (ValueError,TypeError):return None
def mean(v):
 x=[n for a in v if (n:=finite(a)) is not None];return st.mean(x) if x else None
def scene(row):return row.get('scene') or ('Plot'+row['plot'] if 'plot' in row else 'ETH')
def key(row):return scene(row),row['pair'],int(row['seed'])
def stats(rows):
 assert rows and len({key(r) for r in rows})==len(rows)
 good=[r for r in rows if r['success']=='1'];ep=[r for r in rows if finite(r['e_p_m']) is not None]
 out=dict(n=len(rows),ok=len(good),rate=len(good)/len(rows),finite=len(ep),missing=sum(r.get('pose_ok',r.get('transform_ok','0'))!='1' for r in rows),
          ep=mean(r['e_p_m'] for r in ep),ep_success=mean(r['e_p_m'] for r in good),
          ep_median=st.median([float(r['e_p_m']) for r in ep]) if ep else None,
          rot=mean(r['rot_err_deg'] for r in rows),trans=mean(r['trans_err_m'] for r in rows),
          time=mean(r.get('online_ms',r.get('time_ms')) for r in rows))
 for col in ['n_matches','n_correct_matches','ransac_inliers','recovery_triggered','recovery_accepted','recovery_hypotheses','recovery_support','proposal_count','refined_candidates','refine_ms','descriptor_ms','matching_ms','solve_ms']:
  vals=[r.get(col) for r in rows]
  # Zero in an inactive solver/stage is a placeholder, not measured support.
  if col=='ransac_inliers' and all(r.get('solver')=='svd' for r in rows):out[col]=None
  else:out[col]=mean(vals)
 out['recovery_triggered_total']=sum(int(r.get('recovery_triggered') or 0) for r in rows) if any(r.get('recovery_triggered') not in [None,''] for r in rows) else None
 out['recovery_accepted_total']=sum(int(r.get('recovery_accepted') or 0) for r in rows) if any(r.get('recovery_accepted') not in [None,''] for r in rows) else None
 return out

class Report:
 def __init__(self,index):
  self.index=index;self.sources={};self.data={};self.comparisons=[];self.params=[];self.groups={}
  OUT.mkdir(parents=True,exist_ok=True)
 def load(self,sid,relative):
  new_archive=OUT/'evidence'/(sid+'_'+Path(relative).name)
  archived=OLD/'evidence'/(sid+'_'+Path(relative).name)
  p=new_archive if new_archive.exists() else (archived if archived.exists() else ROOT/relative)
  if not p.exists():p=OLD/'evidence'/(sid+'_'+Path(relative).name)
  rows=list(csv.DictReader(p.open()));self.data[sid]=rows
  self.sources[sid]=dict(path=str(p.relative_to(ROOT)),sha256=digest(p),rows=len(rows))
  return rows
 def select(self,sid,**criteria):
  return [r for r in self.data[sid] if all((scene(r) if k=='scene' else r[k])==str(v) for k,v in criteria.items())]
 def comp(self,cid,title,before,after,a,b,level='單項受控',note='',time_fair=True):
  assert {key(r) for r in a}=={key(r) for r in b},cid
  da={key(r):r for r in a};db={key(r):r for r in b};common=[k for k in da if da[k]['success']==db[k]['success']=='1']
  record=dict(id=cid,title=title,before=before,after=after,level=level,note=note,time_fair=time_fair,
              scope=f"{'Tongji-trees' if scene(a[0])!='ETH' else 'ETH-tree'}／{scene(a[0])}；{len({r['pair'] for r in a})}配對；seed{min(int(r['seed']) for r in a)}–{max(int(r['seed']) for r in a)}",
              a=stats(a),b=stats(b),
              gained=sum(da[k]['success']=='0' and db[k]['success']=='1' for k in da),
              lost=sum(da[k]['success']=='1' and db[k]['success']=='0' for k in da),common=len(common),
              common_before=mean(da[k]['e_p_m'] for k in common),common_after=mean(db[k]['e_p_m'] for k in common),
              same_finite={key(r) for r in a if finite(r['e_p_m']) is not None}=={key(r) for r in b if finite(r['e_p_m']) is not None})
  for side,rows in [('a',a),('b',b)]:
   refs=[]
   for sid,source in self.data.items():
    positions={id(r):n+2 for n,r in enumerate(source)}
    lines=[positions[id(r)] for r in rows if id(r) in positions]
    if lines:refs.append(dict(source=sid,csv_lines=lines))
   assert sum(len(x['csv_lines']) for x in refs)==len(rows)
   record[side+'_sources']=refs
  self.comparisons.append(record);return cid
 def param(self,pid,name,unit,before,after,adopt,level,ids,why,paper='待補',direct=True):
  self.params.append(dict(id=pid,name=name,unit=unit,before=before,after=after,adopt=adopt,level=level,ids=ids,why=why,paper=paper,direct=direct))
 def summary(self,ids):
  rr=[]
  for cid in ids:
   c=next(c for c in self.comparisons if c['id']==cid);a,b=c['a'],c['b']
   ep=f"；e_p {a['ep']*100:.3f}→{b['ep']*100:.3f}cm" if a['ep'] is not None and b['ep'] is not None else '；e_p待補'
   rr.append(f"{cid}：{a['ok']}/{a['n']}→{b['ok']}/{b['n']}（{(b['rate']-a['rate'])*100:+.3f}pp）{ep}")
  return '\n'.join(rr)

def prepare1():
 r=Report(1)
 for sid,p in [('F','results/eth_trees/tuning_20261025/runs/fresh_rows.csv'),('V','results/eth_trees/failure_followup_20261025/runs/validation_rows.csv'),('S','results/eth_trees/tuning_20261025/runs/screen_rows.csv'),('U','results/unified_20261025/registration_rows.csv')]:r.load(sid,p)
 archive=OUT/'evidence/S_screen_rows.csv';archive.parent.mkdir(parents=True,exist_ok=True)
 if not archive.exists():archive.write_bytes((ROOT/r.sources['S']['path']).read_bytes())
 r.sources['S']['path']=str(archive.relative_to(ROOT))
 # Flatten only completed discovery outputs. Point error was intentionally skipped.
 discovery=[];files=[]
 for method in ['baseline','triangle_scarce','evidence_scarce','evidence_weak']:
  for p in sorted((ROOT/'results/eth_trees/failure_followup_20261025/runs'/method).glob('*.csv')):
   seed,pair=p.stem.split('_',1);files.append(dict(path=str(p.relative_to(ROOT)),sha256=digest(p)))
   for row in csv.DictReader(p.open()):discovery.append(dict(row,method=method,seed=seed,pair=pair))
 assert len(discovery)==1200
 snapshot=OUT/'evidence/R1_recovery_discovery_rows.csv';snapshot.parent.mkdir(exist_ok=True,parents=True)
 with snapshot.open('w',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(dict.fromkeys(k for row in discovery for k in row)));w.writeheader();w.writerows(discovery)
 r.load('R1',str(snapshot.relative_to(ROOT)));r.sources['R1']['original_files']=files
 # True threshold/solver controls on the same fresh extracted trees.
 th=[];sol=[]
 for solver,tag in [('svd','S'),('svd+ransac','R')]:
  cid=r.comp('F-TH-'+tag,'只改CN距離容差',f'0.05m／{solver}',f'0.10m／{solver}',r.select('F',method='fresh_control',solver=solver),r.select('F',method='fresh_matching',solver=solver),note='提取樹位完全相同；無新增recovery。time_ms為descriptor+matching+solve，非提取。')
  th.append(cid)
 for method,tag in [('fresh_control','05'),('fresh_matching','10')]:
  threshold='0.05m' if tag=='05' else '0.10m'
  sol.append(r.comp('F-SOL-'+tag,'只更換剛體solver',f'{threshold}／SVD',f'{threshold}／SVD+RANSAC',r.select('F',method=method,solver='svd'),r.select('F',method=method,solver='svd+ransac'),level='單項方法切換',note='只切solver；RANSAC內部1000次/0.30m/seed1一起啟用，沒有各常數消融。'))
 r.param('P01','CN threshold','m','0.05','0.10','採用','單一參數',th,'同步影響gatekeeper、DC與triangle的距離一致性；不改Thscore。','0.05')
 r.param('P02','registration solver','方法','Kabsch／SVD','Kabsch＋RANSAC','採用','單項方法',sol,'能排除不符合3D剛體共識的樹位；不是擬合圓柱的RANSAC。')
 for pid,name,unit,val,why in [('P03','registration RANSAC iterations','次','1000','沒有1000對其他迭代數比較；只能引用solver整組。'),('P04','registration RANSAC inlier_dist','m','0.30','使用完整3D殘差；沒有距離門檻獨立掃描。'),('P05','registration RANSAC seed','種子','1','固定solver隨機種子；非提取seed11–20。')]:r.param(pid,name,unit,'不啟用',val,'採用','缺單參數消融',sol,why,direct=False)
 # Recovery controls: direct trigger change exists in completed discovery runs.
 dev=[]
 for mode in ['triangle_scarce','evidence_scarce','evidence_weak']:
  dev.append(r.comp('R-ADD-'+mode,'開發：啟用補救方法','recovery off',mode,r.select('R1',method='baseline',solver='svd+ransac'),r.select('R1',method=mode,solver='svd+ransac'),level='單項方法／內部整組',note='同ETH seed11–20已用於選方案。開發省略e_p，NaN不是0；時間僅舊time_ms。'))
 trigger=r.comp('R-TRIGGER','只改evidence觸發對數','trigger=3','trigger=6',r.select('R1',method='evidence_scarce',solver='svd+ransac'),r.select('R1',method='evidence_weak',solver='svd+ransac'),note='mode=evidence、top3、最多60邊、至少6支持及0.30m距離固定。這是開發資料，非盲測。')
 mode=r.comp('R-MODE','相同trigger=3只改補救mode','triangles','evidence',r.select('R1',method='triangle_scarce',solver='svd+ransac'),r.select('R1',method='evidence_scarce',solver='svd+ransac'),level='單項方法切換',note='mode切換代表整個內部算法不同，不是只動一個幾何常數；e_p當時未計算。')
 val=[]
 for solver,tag in [('svd','S'),('svd+ransac','R')]:
  val.append(r.comp('V-REC-'+tag,'凍結後新seed：僅新增weak recovery','10cm／recovery off','10cm／evidence trigger6',r.select('V',method='baseline',solver=solver),r.select('V',method='recovery',solver=solver),level='單項方法／內部整組',note='seed21–30僅隨機性檢查，仍是同六站森林。GT不進推論，但離線開發用過GT選方案。'))
 r.param('P06','CN recovery mode','方法','off','evidence','採用','方法整組',val,'整個weak recovery新增的效果；不能分配給每個內部常數。',direct=True)
 r.param('P07','recovery trigger_matches','原CN對數','3','6','採用','單一參數',[trigger],'原match對數<門檻才觸發；與至少3對能解Kabsch的數學下限不同。')
 r.param('P08','recovery mode（候選比較）','方法','triangles','evidence','採用evidence','單項方法',[mode],'固定trigger3比較兩算法；成功數相同，不宣稱兩者功能等價。')
 for pid,name,unit,val0,why in [('P09','recovery top_candidates','個','3','沒有top3對top1/5的對照。'),('P10','recovery max_evidence_pairs','條邊','60','沒有60對其他邊數的對照。'),('P11','recovery min_support','個3D支持','6','與trigger_matches=6不同用途，沒有min_support獨立消融。'),('P12','recovery inlier_distance','m','0.30','共識使用3D距離，不是GT正確對應的XY檢查。'),('P13','recovery XY length tolerance','m','2×CNthreshold=0.20','三元組相容邊長差，隨CNthreshold連動；沒有獨立改該係數。'),('P14','recovery hypothesis / solver','方法','相容三元組＋6DOF Kabsch','新增內部幾何處理，沒有與recovery其他條件分離。')]:r.param(pid,name,unit,'無新增recovery',val0,'採用','缺單參數消融',val,why,direct=False)
 # Keep unsuccessful extraction changes and separate their two changed fields.
 vert=[];cyl=[]
 for baseline,vm,cm0,tag in [('fresh_control','fresh_vertical085','fresh_ransac999','05'),('fresh_matching','fresh_vertical085_matching','fresh_ransac999_matching','10')]:
  for solver,stag in [('svd','S'),('svd+ransac','R')]:
   a=r.select('F',method=baseline,solver=solver)
   vert.append(r.comp('F-VERT-'+tag+stag,'提取verticality單項變更',baseline,vm,a,r.select('F',method=vm,solver=solver),note='法向篩選改變提取樹位，其他指定條件固定；不等於v3驗證normal.z。'))
   cyl.append(r.comp('F-CYL-'+tag+stag,'圓柱RANSAC兩常數同時變更',baseline,cm0,a,r.select('F',method=cm0,solver=solver),level='兩參數整組',note='iterations與probability同時改，缺任一單獨對照；不是註冊RANSAC。'))
 r.param('P15','extraction normals.vert_threshold','verticality','0.90','0.85','未採用','單一參數',vert,'放寬近垂直表面保留，在主10cm/RANSAC分支142→130；變差也保留。','>0.9')
 r.param('P16','cylinder.max_iterations','次','1000','3000','未採用','兩參數同改',cyl,'與probability一起改；10cm/RANSAC142→141，不能說3000本身的效果。',direct=False)
 r.param('P17','cylinder.probability','機率','0.99','0.999','未採用','兩參數同改',cyl,'與iterations一起改；缺0.999單獨效果。',direct=False)
 # Every other completed matching-parameter sweep gets its own before/after result.
 screens=[('threshold075','CN threshold','m','0.05','0.075','未採用0.075','0.05'),('threshold100','CN threshold（早期探索）','m','0.05','0.10','後續採用0.10','0.05'),('score20','CN th_score','分數','40','20','未採用','40'),('score60','CN th_score','分數','40','60','未採用','40'),('min_dc4','CN min_dc','個扇區','3','4','未採用','待補／文字偽碼邊界有歧義'),('sectors93','CN N','維度','183','93','未採用','183'),('merge015','tree merge_radius','m','0（不合併）','0.15','未採用','待補'),('merge030','tree merge_radius','m','0（不合併）','0.30','未採用','待補')]
 for i,(method,name,unit,b,a,adopt,paper) in enumerate(screens,18):
  ids=[]
  for low,high,tag in [(1,5,'DEV'),(6,10,'OLD')]:
   for solver,stag in [('svd','S'),('svd+ransac','R')]:
    pred=lambda x:low<=int(x['seed'])<=high
    aa=list(filter(pred,r.select('S',method='control',solver=solver)));bb=list(filter(pred,r.select('S',method=method,solver=solver)))
    ids.append(r.comp('S-'+method+'-'+tag+stag,'早期探索：'+name,b,a,aa,bb,note='固定同solver，其他matching條件同control。seed1–5開發、6–10既有seed檢查，非未見森林。'))
  r.param('P'+str(i).zfill(2),name,unit,b,a,adopt,'單一參數探索',ids,'只有本輪範圍內的條件效果；不套用到未完成的其他森林。',paper)
 # Main report transfer: whole coarse bundle, not falsely attributed to tolerance alone.
 combined=[]
 for s in SCENES:
  combined.append(r.comp('U-'+s,'主表：完整粗方法整組', 'reference5_raw','eth10_raw',r.select('U',scene=s,method='reference5_raw'),r.select('U',scene=s,method='eth10_raw'),level='三項方法整組',note='threshold、solver、recovery同時不同；相同提取、step、ctol0.2。有限集合改變時不算誤差下降率。'))
 r.param('P26','主報告整體coarse方法','方法','5cm＋SVD＋recovery off','10cm＋RANSAC＋weak recovery','採用','完整整組',combined,'跨Tongji四Plot與ETH的整組差異；非單一threshold的跨資料集效果。','0.05m；solver與新增recovery待補',direct=False)
 return r

def build(r):
 index=r.index;path=OUT/f'report_{index:02d}_{TITLES[index]}_逐項參數與前後數據.xlsx'
 wb=xlsxwriter.Workbook(path);wb.set_properties({'title':f'第{index}次：逐參數前後數據','comments':'Only completed experiments. GT-free inference, not GT-free development.'})
 fmt={}
 for name,opts in dict(title={'bold':True,'font_size':18,'font_color':'203D56'},head={'bold':True,'bg_color':'E7EDF2'},text={},num={'num_format':'0.000000'},integer={'num_format':'0'},pct={'num_format':'0.000%'},pp={'num_format':'0.000" pp"'},note={'font_size':10,'font_color':'555555'},missing={'font_color':'A35A00'},link={'font_color':'0563C1','underline':1}).items():
  fmt[name]=wb.add_format(dict(font_name='Microsoft JhengHei',font_size=12,text_wrap=True,valign='vcenter')|opts)
 # Font size is explicitly overwritten for title/note.
 pages={name:wb.add_worksheet(name) for name in ['逐項參數證據','每項修改前後數據','論文與場景總覽']}
 for ws in pages.values():
  ws.hide_gridlines(2);ws.set_zoom(85);ws.set_landscape();ws.set_paper(8);ws.fit_to_pages(1,0);ws.set_margins(.25,.25,.35,.35);ws.set_header('&C&12第'+str(index)+'次報告：逐項參數與數據');ws.set_footer('&L缺值不是零；pp為百分點&R&P / &N')
 checks=[];counts={'formulas':0}
 def write(ws,row,col,value,kind='text'):
  val=MISSING if value is None else value;ws.write(row,col,val,fmt['missing'] if val==MISSING else fmt[kind]);checks.append(dict(sheet=ws.name,row=row,col=col,expected=val))
 def formula(ws,row,col,expr,value,kind='num'):
  val=MISSING if value is None else value;ws.write_formula(row,col,expr,fmt[kind],val);checks.append(dict(sheet=ws.name,row=row,col=col,expected=val,formula=expr));counts['formulas']+=1
 def merge(ws,row,col,last,value,kind='note',height=42):
  ws.merge_range(row,col,row,last,value,fmt[kind]);ws.set_row(row,height)
 common='成功：e_R<2°且e_t<0.5m；e_p為全來源點Eq8，finite mean保留可計算失敗；無變換仍算失敗。GT不進單次推論，但離線開發看過GT選參。'
 # Detailed comparison blocks, linked from the parameter catalogue.
 ws=pages['每項修改前後數據'];ws.set_column('A:A',28);ws.set_column('B:B',12);ws.set_column('C:F',19);ws.set_column('G:G',27);ws.set_column('H:H',60);ws.freeze_panes(5,2)
 merge(ws,0,0,7,'每項改動：完整修改前／後數據','title',36)
 merge(ws,1,0,7,common,height=44)
 merge(ws,2,0,7,'先由第一頁的參數ID／實驗索引跳入。各block是可追溯的受控對照或整組對照；整組數字不是每個內部參數的獨立效果。誤差與時間下降率=(前−後)/前；success差為後−前的pp。',height=44)
 row=5;anchors={}
 metrics=[('試驗總數','次','n',1),('成功數','次','ok',1),('成功率','比例','rate',1),('有限e_p數','次','finite',1),('缺變換數','次','missing',1),('全部有限mean e_p','cm','ep',100),('有限median e_p','cm','ep_median',100),('成功限定mean e_p','cm','ep_success',100),('mean e_R（含有限失敗）','degree','rot',1),('mean e_t（含有限失敗）','cm','trans',100),('mean online／舊time_ms','ms','time',1),('平均matching對應數','對','n_matches',1),('平均GT-XY正確對應','對','n_correct_matches',1),('平均註冊RANSAC inliers','對','ransac_inliers',1),('recovery觸發總數','次','recovery_triggered_total',1),('recovery採用總數','次','recovery_accepted_total',1),('平均recovery支持數','個','recovery_support',1),('平均候選pose數','個','proposal_count',1),('平均精化候選數','個','refined_candidates',1),('平均refine_ms','ms','refine_ms',1)]
 for c in r.comparisons:
  anchors[c['id']]=row;ws.set_h_pagebreaks([*getattr(ws,'hbreaks',[]),row])
  merge(ws,row,0,7,c['id']+'｜'+c['title'],'head',28);row+=1
  merge(ws,row,0,7,c['scope']+'；修改前='+c['before']+'；修改後='+c['after']+'；證據='+c['level'],height=40);row+=1
  merge(ws,row,0,7,c['note'],height=40);row+=1
  for j,x in enumerate(['指標','單位','修改前','修改後','差：後−前','下降率：以修改前為基準','統計母體／N','讀法與限制']):write(ws,row,j,x,'head')
  row+=1
  for title,unit,k,mult in metrics:
   av,bv=c['a'].get(k),c['b'].get(k)
   if av is None and bv is None and k not in ['ep','ep_success','time','n','ok','rate','finite','missing','rot','trans']:continue
   a=av*mult if av is not None else None;b=bv*mult if bv is not None else None
   write(ws,row,0,title);write(ws,row,1,unit);kind='pct' if k=='rate' else ('integer' if k in ['n','ok','finite','missing','recovery_triggered_total','recovery_accepted_total'] else 'num');write(ws,row,2,a,kind);write(ws,row,3,b,kind)
   rr=row+1;delta=(b-a)*(100 if k=='rate' else 1) if a is not None and b is not None else None
   formula(ws,row,4,f'=IF(AND(ISNUMBER(C{rr}),ISNUMBER(D{rr})),(D{rr}-C{rr})'+('*100' if k=='rate' else '')+',"待補")',delta,'pp' if k=='rate' else 'num')
   comparable=k in ['ep','ep_median'] and c['same_finite'] or k=='time' and c['time_fair']
   decline=(a-b)/a if comparable and a not in [None,0] and b is not None else None
   if comparable:formula(ws,row,5,f'=IF(AND(ISNUMBER(C{rr}),ISNUMBER(D{rr}),C{rr}<>0),(C{rr}-D{rr})/C{rr},"待補")',decline,'pct')
   else:write(ws,row,5,None)
   if k in ['ep','ep_median']:population=f"finite {c['a']['finite']}/{c['a']['n']} → {c['b']['finite']}/{c['b']['n']}";note='包含有限失敗；有限集合不同不算下降率。'
   elif k=='ep_success':population=f"成功 {c['a']['ok']} → {c['b']['ok']}";note='成功集合可能改變，另看共同成功。'
   else:population=f"總N {c['a']['n']} → {c['b']['n']}";note=''
   if k=='rate':note='差值單位pp；不是相對百分比。'
   if k=='n_correct_matches':note='GT後XY<0.30m；只作事後診斷，不檢查z一致性。'
   if k=='time':note='不含提取、cache及GT評估；'+('同批次描述性比較，舊CPU型號／並行干擾未齊。' if c['time_fair'] else '跨批次／epoch，速度下降率待補。')
   if k=='ransac_inliers':note='SVD未使用RANSAC留待補，不把CSV預設0當真實inlier數。'
   write(ws,row,6,population);write(ws,row,7,note);ws.set_row(row,34);row+=1
  for title,k in [('新增成功','gained'),('退步失敗','lost'),('共同成功數','common')]:
   write(ws,row,0,title);write(ws,row,1,'次');write(ws,row,3,c[k],'integer');write(ws,row,7,'同pair＋seed的配對比較');row+=1
  a=c['common_before'];b=c['common_after'];write(ws,row,0,'共同成功mean e_p');write(ws,row,1,'cm');write(ws,row,2,a*100 if a is not None else None,'num');write(ws,row,3,b*100 if b is not None else None,'num');rr=row+1
  formula(ws,row,4,f'=IF(AND(ISNUMBER(C{rr}),ISNUMBER(D{rr})),D{rr}-C{rr},"待補")',(b-a)*100 if a is not None and b is not None else None)
  formula(ws,row,5,f'=IF(AND(ISNUMBER(C{rr}),ISNUMBER(D{rr}),C{rr}<>0),(C{rr}-D{rr})/C{rr},"待補")',(a-b)/a if a not in [None,0] and b is not None else None,'pct');write(ws,row,6,f"共同成功N={c['common']}");write(ws,row,7,'僅共同成功，有存活者偏差，不能取代全體成功率。');row+=1
  refs='；'.join(side+':'+','.join(x['source'] for x in c[side+'_sources']) for side in ['a','b'])
  merge(ws,row,0,7,'來源 '+refs+'；精確CSV行號與SHA見隨附manifest。'+c['level']+'的結果，不能超出對照條件解讀。',height=34);row+=3
 ws.print_area(0,0,row,7);ws.repeat_rows(0,2)
 # Parameter values + independent evidence and bundle evidence in separate columns.
 ws=pages['逐項參數證據'];widths=[10,32,12,22,25,19,23,58,58,48]
 for j,w in enumerate(widths):ws.set_column(j,j,w)
 merge(ws,0,0,9,f'report_{index:02d}｜每一項參數改動與數據證據','title',38)
 merge(ws,1,0,9,'每列一項參數／方法。H欄只放能隔離該項的既有對照；沒有就寫待補。I欄可引用相關整組結果，但不能當成該參數的獨立貢獻。點J欄跳到完整前後數據。',height=46)
 merge(ws,2,0,9,common,height=40)
 headers=['ID','參數／方法名稱','單位','修改前','修改後','最後採用？','證據類型','單項／單方法前後數據','相關整組前後數據（非獨立效果）','實驗索引與作用／限制']
 for j,h in enumerate(headers):write(ws,4,j,h,'head')
 for n,p in enumerate(r.params,5):
  independent=r.summary(p['ids']) if p['direct'] else '待補：沒有只改這一個參數的實驗。'
  bundle=r.summary(p['ids']) if not p['direct'] else '詳見H欄；不可超出該對照條件。'
  values=[p['id'],p['name'],p['unit'],p['before'],p['after'],p['adopt'],p['level'],independent,bundle,', '.join(p['ids'])+'\n'+p['why']+'\n論文：'+p['paper']]
  for j,v in enumerate(values):write(ws,n,j,v)
  ws.write_url(n,9,"internal:'每項修改前後數據'!A"+str(anchors[p['ids'][0]]+1),fmt['link'],values[9])
  checks[-1]['expected']=values[-1]
  ws.set_row(n,min(225,max(90,22*max(independent.count('\n')+1,bundle.count('\n')+1,4))))
 ws.autofilter(4,0,4+len(r.params),9);ws.freeze_panes(5,2);ws.repeat_rows(0,4);ws.print_area(0,0,len(r.params)+6,9)
 # The paper comparison remains a reference, not an author-code rerun.
 ws=pages['論文與場景總覽'];ws.set_column('A:A',24);ws.set_column('B:C',18);ws.set_column('D:H',23);ws.set_column('I:J',31)
 merge(ws,0,0,9,f'第{index}次｜論文與本地主比較','title',36)
 merge(ws,1,0,9,'Tongji四Plot屬同一dataset，ETH為另一dataset。作者原作重跑、論文配對分母與成功門檻待補；本地reference5不是作者程式。數字下降不代表嚴格超越作者實作。',height=46)
 merge(ws,2,0,9,common,height=40)
 headers=['dataset／場景','論文成功率','本地修改前成功率','修改後成功率','差：後−前 pp','論文e_p cm','前finite mean cm','後finite mean cm','vs論文下降率','前後樣本／比較限制']
 for j,h in enumerate(headers):write(ws,4,j,h,'head')
 summary_comps=[c for c in r.comparisons if c['id'].startswith('U-')]
 for row,c in enumerate(summary_comps,5):
  s=next(s for s in SCENES if c['id']=='U-'+s);a,b=c['a'],c['b'];rr=row+1
  write(ws,row,0,'ETH-tree' if s=='ETH' else 'Tongji／'+s);write(ws,row,1,1,'pct');write(ws,row,2,a['rate'],'pct');write(ws,row,3,b['rate'],'pct')
  formula(ws,row,4,f'=(D{rr}-C{rr})*100',(b['rate']-a['rate'])*100,'pp');write(ws,row,5,PAPER_EP[s]*100,'num');write(ws,row,6,a['ep']*100 if a['ep'] is not None else None,'num');write(ws,row,7,b['ep']*100 if b['ep'] is not None else None,'num')
  if b['finite']==b['n']:formula(ws,row,8,f'=IF(AND(ISNUMBER(F{rr}),ISNUMBER(H{rr}),F{rr}<>0),(F{rr}-H{rr})/F{rr},"待補")',(PAPER_EP[s]-b['ep'])/PAPER_EP[s],'pct')
  else:write(ws,row,8,None)
  write(ws,row,9,f"成功 {a['ok']}/{a['n']}→{b['ok']}/{b['n']}；finite {a['finite']}→{b['finite']}；整組比較")
  ws.set_row(row,52)
 merge(ws,12,0,9,'論文來源：上傳PDF p4 Table I／Eq8；本地來源U（逐例重算）。論文平均e_p對照僅供參考；缺變換不填0。分母或有限集合不同時，不以兩個平均數冒充相同案例的效果。',height=48)
 merge(ws,14,0,9,'主表修改前後版本：'+(summary_comps[0]['before']+' → '+summary_comps[0]['after'] if summary_comps else '見逐項頁')+'。本次沒有新增配準試驗；逐項歸因不足處如實待補。',height=44)
 if summary_comps:
  chart=wb.add_chart({'type':'column'});chart.set_title({'name':'成功率：每場景分開'});chart.set_y_axis({'name':'成功率','min':0,'max':1,'num_format':'0%'});chart.set_legend({'position':'bottom'})
  for col,title in [(2,'修改前'),(3,'修改後')]:chart.add_series({'name':title,'categories':[ws.name,5,0,9,0],'values':[ws.name,5,col,9,col],'fill':{'color':'7D8C98' if col==2 else '305B81'},'data_labels':{'value':True,'num_format':'0%'}})
  ws.insert_chart('A18',chart,{'x_scale':1.45,'y_scale':1.1})
  chart2=wb.add_chart({'type':'column'});chart2.set_title({'name':'全有限mean e_p（包含有限失敗）'});chart2.set_y_axis({'name':'e_p（cm）','min':0});chart2.set_legend({'position':'bottom'})
  for col,title in [(5,'論文報告值'),(6,'本地修改前'),(7,'修改後')]:chart2.add_series({'name':title,'categories':[ws.name,5,0,9,0],'values':[ws.name,5,col,9,col]})
  ws.insert_chart('F18',chart2,{'x_scale':1.45,'y_scale':1.1})
 source_row=42
 merge(ws,source_row,0,9,'數據來源與SHA256：逐block的a/b表示修改前／後；CSV篩選行號另存manifest，以下來源可追溯到專案。','head',32)
 for sid,info in r.sources.items():
  source_row+=1;merge(ws,source_row,0,9,sid+'：'+info['path']+'；'+str(info['rows'])+'列；SHA256='+info['sha256'],height=36)
 source_row+=2
 merge(ws,source_row,0,9,'所有參數作用與單位依config/default.ini、include/cm/registration.hpp、src/cn_recovery.cpp、src/cn_matching.cpp及凍結plan核對。論文：Sun等，DOI 10.1109/LGRS.2025.3650098；PDF p2方法、p3 Algorithm1、p4 TableI/II與Eq8、p5 TableIII。原論文時間硬體i7-9700F 3GHz、OpenMP執行緒數待補；本地舊批次OMP4、CPU型號未齊，禁止與paper做公平速度下降率。',height=64)
 ws.freeze_panes(5,1);ws.print_area(0,0,source_row+1,9)
 wb.close()
 manifest=dict(generated_at=datetime.now(ZoneInfo('Asia/Taipei')).isoformat(),report=index,file=path.name,sha256=digest(path),parameters=r.params,comparisons=r.comparisons,sources=r.sources,checks=checks,formula_count=counts['formulas'],notes='No new inference. Single-parameter evidence is separated from shared bundle evidence. GT used in development, not in per-pair inference.')
 mp=OUT/f'report_{index:02d}_manifest.json';mp.write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
 # Actual reopening validates cached formulas and source-literal tables.
 check_book(path,manifest)
 print(json.dumps(dict(file=str(path),parameters=len(r.params),comparisons=len(r.comparisons),formula_count=counts['formulas'],checks=len(checks),sha256=digest(path)),ensure_ascii=False))
 return path

def check_book(path,manifest):
 book=load_workbook(path,data_only=False);cache=load_workbook(path,data_only=True)
 assert len(book.sheetnames)==3 and not book._external_links
 for check in manifest['checks']:
  row,col=check['row']+1,check['col']+1;expected=check['expected'];value=cache[check['sheet']].cell(row,col).value
  assert (math.isclose(value,expected,rel_tol=1e-10,abs_tol=1e-10) if isinstance(expected,(int,float)) else (expected==value or expected=='' and value is None)),(check,value)
  if 'formula' in check:assert book[check['sheet']].cell(row,col).data_type=='f'
 for ws in cache:
  for row in ws:
   assert all(cell.data_type!='e' for cell in row)
 assert sum(len(ws._charts) for ws in book)==2
 book.close();cache.close()

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--report',type=int,default=1);args=ap.parse_args()
 if args.report==1:r=prepare1()
 else:raise ValueError('Report not implemented yet')
 build(r)
if __name__=='__main__':main()
