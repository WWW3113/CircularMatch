#!/usr/bin/env python3
"""Four evidence-led professor reports from completed runs; no new experiment."""
from pathlib import Path
import csv, json, hashlib, math, statistics as st, subprocess, shutil,os
import unicodedata
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.pagebreak import Break

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'results/professor_reports_20261026';RUNS=OUT/'runs';RUNS.mkdir(exist_ok=True,parents=True)
PAPER=Path(os.environ.get('CIRCULARMATCH_PAPER','/workspace/attachments/0da7b2b7-8aa0-4620-a7ef-7d834adc5316/Sun …等 - 2026 - CircularMatch Efficient Registration of Forest Terrestrial Point Clouds Using Circular Neighborhood.pdf'))
P_EP=dict(Plot1=.052558,Plot2=.114808,Plot3=.048756,Plot4=.043560,ETH=.153427)
P_MATCH=dict(Plot1=20.718,Plot2=30.699,Plot3=4.5596,Plot4=19.522,ETH=3.0405)
P_TOTAL=dict(Plot1=11.657,Plot2=16.712,Plot3=8.3584,Plot4=9.4516,ETH=8.6992)
SCENES=list(P_EP);MISSING='待補';SOURCES={};GROUPS={};CHECKS={};RANGES={};CHARTS={};BOOKS={}
TITLES={1:'匹配容差與弱配對補救',2:'DTM平滑的局部改善與代價',3:'加入多尺度幾何微調',4:'多候選與垂直結構驗證'}
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def read(path):
 with (ROOT/path).open() as f:return list(csv.DictReader(f))
def source(sid,path,description):
 p=ROOT/path
 if not p.exists():p=OUT/'evidence'/(sid+'_'+Path(path).name)
 SOURCES[sid]=dict(path=path,sha256=sha(p),bytes=p.stat().st_size,description=description)
 if p.suffix=='.csv':
  with p.open() as f:return list(csv.DictReader(f))
 return json.loads(p.read_text()) if p.suffix=='.json' else p.read_text()
U=source('U','results/unified_20261025/registration_rows.csv','統一協定：9版本×550案例，seed31–40，全來源點Eq8')
D=source('D','results/tongji_trees/dtm_height_20261025/registration_rows.csv','DTM固定軸線定位測試；主要solver=svd+ransac')
DF=source('DF','results/tongji_trees/dtm_height_20261025/full_pipeline_rows.csv','DTM完整重提取，只有Plot2')
F=source('F','results/eth_trees/tuning_20261025/runs/fresh_rows.csv','ETH seed11–20，matching 2×2 controlled experiment')
V=source('V','results/eth_trees/failure_followup_20261025/runs/validation_rows.csv','ETH seed21–30，固定10cm/RANSAC之補救前後')
H=source('H','results/tongji_trees/dtm_height_20261025/fixed_pair_height_diagnostic.csv','固定原CN配對的GT高度診斷，僅評估使用')
COLD=source('C','results/unified_20261025/cold_benchmark.json','12筆冷啟動；僅4筆v3同epoch matched comparison用於時間下降率')
CFG=source('CFG','results/unified_20261025/unified_config.json','固定參數，原提取DTM未套smooth')
EPOCH=source('HW','results/unified_20261025/current_epoch.json','新epoch硬體；不能套用為舊epoch的已知CPU型號')
for sid,path,desc in [
 ('AUTHOR','results/unified_20261025/author_source_audit.json','作者程式可得性稽核，稽核當時只有README'),
 ('DP','results/tongji_trees/dtm_height_20261025/plan.json','定位候選、開發範圍、固定條件'),
 ('DFP','results/tongji_trees/dtm_height_20261025/full_pipeline_plan.json','完整流程僅Plot2'),
 ('VP','results/eth_trees/failure_followup_20261025/frozen_recovery.json','先凍結補救，再算seed21–30'),
 ('UP','results/unified_20261025/protocol.json','55配對與成功門檻'),
 ('V3P','results/unified_20261025/v3_protocol.json','v3事先固定的整組修改'),
 ('EVAL','results/unified_20261025/evaluator_full_cloud_audit.json','19,633,280點Python獨立核對Eq8'),
 ('US','results/unified_20261025/scene_summary.csv','已發布場景摘要，獨立交叉核對本次重算'),
 ('DSP','results/tongji_trees/dtm_height_20261025/full_pipeline_paired.csv','既有成對新增成功／退步紀錄'),
 ('PAIR','results/unified_20261025/ablation.csv','既有成對消融紀錄'),
 ('CODE_MATCH','src/cn_matching.cpp','距離容差如何進入DC、gatekeeper、triangle'),
 ('CODE_REC','src/cn_recovery.cpp','弱配對補救；GT不進入推論'),
 ('CODE_REF','apps/geometry_refinement.cpp','原refinement、3h驗證距離'),
 ('CODE_V3','apps/geometry_refinement_structural.cpp','垂直表面篩選及驗證gate一起改'),
 ('CODE_DTM','src/dtm.cpp','局部平面／補洞機制'),
 ('CODE_TIME','apps/cm_register.cpp','舊time_ms含descriptor+matching+recovery+solver，無提取與GT評分'),
 ('CODE_REG','include/cm/registration.hpp','SVD為本地選擇；RANSAC與error定義'),
 ('BASECFG','config/default.ini','未指定部分及額外cylinder檢查，不能當嚴格作者baseline')]:source(sid,path,desc)
for sid,path,desc in [
 ('DDEV','results/tongji_trees/dtm_height_20261025/development_summary.csv','4個DTM開發候選，Plot2 seed31–35'),
 ('DAUD','results/tongji_trees/dtm_height_20261025/baseline_position_audit.json','200份原樹位一致性與DTM殘差'),
 ('DTILT','results/tongji_trees/dtm_height_20261025/gt_tilt_diagnostic.csv','GT垂直軸夾角，只作診斷'),
 ('DFEX','results/tongji_trees/dtm_height_20261025/full_extraction_summary.csv','50份完整提取樹數變化'),
 ('DATA','results/unified_20261025/dataset_manifest.json','26站、55配對、單位與GT方向'),
 ('FS','results/eth_trees/tuning_20261025/screen_summary.csv','探索seed1–5與既有seed6–10，不能當新森林'),
 ('FD','results/eth_trees/failure_followup_20261025/discovery_summary.csv','補救候選開發結果'),
 ('L','results/liang_trees/transfer_20261025/summary.csv','只做處理後資料樹位提取，非配準結果'),
 ('K','results/whu_park/transfer_20261025/summary.csv','Park粗配準診斷；不是WHU-FGI'),
 ('KS','results/whu_park/sparse_stems_20261025/validation_summary.csv','Park稀疏樹幹支線，非最終v3'),
 ('GS','results/eth_trees/tuning_20261025/fresh_summary.csv','包含6站graph_mst，任務不同不與兩站直接比較')]:source(sid,path,desc)
source('LOG','results/unified_20261025/RESEARCH_LOG.md','失敗分析、資料接觸歷史與假設界線')
source('EXEC','results/unified_20261025/execution_config.json','舊批次固定4OMP；舊epoch CPU型號待補')
if PAPER.exists():
 SOURCES['P']=dict(path='使用者上傳CircularMatch原論文PDF（不隨報告重傳）',sha256=sha(PAPER),doi='10.1109/LGRS.2025.3650098',pages=5,description='PDF p2方法；p3 Algorithm1；p4 TableI/II與Eq8；p5 TableIII。表格數值逐頁視讀核對。')
else:SOURCES['P']=json.loads((OUT/'source_manifest.json').read_text())['sources']['P']
def num(x):
 try:v=float(x);return v if math.isfinite(v) else None
 except (ValueError,TypeError):return None
def avg(x):
 a=[v for v in map(num,x) if v is not None];return st.mean(a) if a else None
def scene(r):return r.get('scene') or ('Plot'+r['plot'] if 'plot' in r else 'ETH')
def key(r):return (scene(r),r['pair'],int(r.get('seed',r.get('label'))))
def group(gid,sid,data,predicate,selector):
 rr=[r for r in data if predicate(r)];assert rr,gid
 assert len({key(r) for r in rr})==len(rr),gid
 ep=[num(r['e_p_m']) for r in rr];finite=[x for x in ep if x is not None];good=[r for r in rr if r['success']=='1']
 pairs={r['pair'] for r in rr};seeds={int(r['seed']) for r in rr}
 stat=dict(n=len(rr),ok=len(good),rate=len(good)/len(rr),pairs=len(pairs),seeds=sorted(seeds),ep=st.mean(finite) if finite else None,ep_n=len(finite),ep_success=avg(r['e_p_m'] for r in good),ep_median=st.median(finite) if finite else None,rot=avg(r['rot_err_deg'] for r in rr),trans=avg(r['trans_err_m'] for r in rr),time=avg(r.get('online_ms',r.get('time_ms')) for r in rr),all_seed_pairs=sum(all(r['success']=='1' for r in rr if r['pair']==p) for p in pairs))
 GROUPS[gid]=dict(source=sid,selector=selector,csv_lines=[i+2 for i,r in enumerate(data) if predicate(r)],stats=stat)
 return dict(id=gid,rows=rr,**stat)
def ug(scene_name,method):
 gid='U:'+scene_name+':'+method
 if gid in GROUPS:return dict(id=gid,rows=[r for r in U if r['scene']==scene_name and r['method']==method],**GROUPS[gid]['stats'])
 return group(gid,'U',U,lambda r:r['scene']==scene_name and r['method']==method,dict(scene=scene_name,method=method,seeds='31–40'))
def dg(scene_name,method):return group('D:'+scene_name+':'+method,'D',D,lambda r:scene(r)==scene_name and r['method']==method and r['solver']=='svd+ransac',dict(scene=scene_name,method=method,solver='svd+ransac',seeds='31–40'))
def joint(a,b):
 da={key(r):r for r in a['rows']};db={key(r):r for r in b['rows']};assert da.keys()==db.keys()
 common=[k for k in da if da[k]['success']=='1' and db[k]['success']=='1']
 return dict(n=len(da),gained=sum(da[k]['success']=='0' and db[k]['success']=='1' for k in da),lost=sum(da[k]['success']=='1' and db[k]['success']=='0' for k in da),common=len(common),before=avg(da[k]['e_p_m'] for k in common),after=avg(db[k]['e_p_m'] for k in common))
def cm(v):return 100*v if v is not None else MISSING
def f(v,d=3):return f'{v:.{d}f}' if isinstance(v,(int,float)) else str(v)
def label(s):return 'ETH-tree' if s=='ETH' else 'Tongji／'+s
BASE={s:ug(s,'reference5_raw') for s in SCENES}
SETS={1:[dict(scene=s,scope='',ref=BASE[s],before=BASE[s],after=ug(s,'eth10_raw')) for s in SCENES],2:[],3:[dict(scene=s,scope='',ref=BASE[s],before=ug(s,'eth10_raw'),after=ug(s,'eth10_refined')) for s in SCENES],4:[dict(scene=s,scope='',ref=BASE[s],before=ug(s,'multi_v2_refined'),after=ug(s,'multi_v3_refined')) for s in SCENES]}
for s in SCENES[:4]:SETS[2].append(dict(scene=s,scope='僅定位',ref=BASE[s],before=dg(s,'baseline'),after=dg(s,'selected')))
SETS[2].append(dict(scene='Plot2',scope='完整重提取',ref=BASE['Plot2'],before=dg('Plot2','baseline'),after=group('DF:Plot2:full','DF',DF,lambda r:r['solver']=='svd+ransac',dict(scene='Plot2',method='full_pipeline',solver='svd+ransac',seeds='31–40'))))
SETS[2].append(dict(scene='ETH',scope='本輪未測',ref=BASE['ETH'],before=None,after=None))
FA=[]
for method in ['fresh_control','fresh_matching']:
 for solver in ['svd','svd+ransac']:
  FA.append(group('F:'+method+':'+solver,'F',F,lambda r,m=method,so=solver:r['method']==m and r['solver']==so,dict(method=method,solver=solver,seeds='11–20')))
VA=group('V:baseline','V',V,lambda r:r['method']=='baseline' and r['solver']=='svd+ransac',dict(method='baseline',solver='svd+ransac',seeds='21–30'))
VB=group('V:recovery','V',V,lambda r:r['method']=='recovery' and r['solver']=='svd+ransac',dict(method='recovery',solver='svd+ransac',seeds='21–30'))
VJ=joint(VA,VB)

# Cross-check independently published summaries before rendering.
for r in read('results/unified_20261025/scene_summary.csv'):
 g=ug(r['scene'],r['method']);assert g['ok']==int(r['success_trials']) and g['ep_n']==int(r['finite_error_cases'])
 assert math.isclose(g['ep'],float(r['mean_finite_ep_m']),rel_tol=1e-12,abs_tol=1e-12)
for r in read('results/tongji_trees/dtm_height_20261025/summary.csv'):
 if r['group'] in SCENES and r['solver']=='svd+ransac':
  g=dg(r['group'],r['method']);assert g['ok']==int(r['success']);assert math.isclose(g['ep_success'],float(r['mean_success_ep_m']),abs_tol=1e-10)

FONT='Microsoft JhengHei';GRAY='E7E6E6';NAVY='305B81';CHECKS={str(i):[] for i in TITLES}
def put(ws,row,values,rid):
 lines=1
 for j,v in enumerate(values,1):
  c=ws.cell(row,j,v);c.font=Font(name=FONT,size=13);c.alignment=Alignment(vertical='center',wrap_text=True)
  if isinstance(v,str) and not v.startswith('='):
   visual=sum(2 if unicodedata.east_asian_width(x) in 'WF' else 1 for x in v)
   lines=max(lines,math.ceil(visual/max(5,(ws.column_dimensions[col(j)].width-2)*11/13)))
  if isinstance(v,float):c.number_format='0.000'
  if isinstance(v,(int,float)):CHECKS[str(rid)].append(dict(sheet=ws.title,cell=c.coordinate,expected=v,kind='literal'))
 ws.row_dimensions[row].height=max(38,min(150,lines*18+6))

def formula(ws,r,c,expr,expected,rid,fmt='0.0%'):
 cell=ws.cell(r,c,expr);cell.number_format=fmt;CHECKS[str(rid)].append(dict(sheet=ws.title,cell=cell.coordinate,expected=expected,kind='formula'))
def decline(ws,row,dest,before,after,bv,av,rid,allowed=True):
 expr=f'=IF(AND(ISNUMBER({before}),ISNUMBER({after}),{before}<>0),({before}-{after})/{before},"待補")'
 if not allowed:
  ws.cell(row,dest,MISSING);CHECKS[str(rid)].append(dict(sheet=ws.title,cell=ws.cell(row,dest).coordinate,expected=MISSING,kind='unavailable'));return
 expected=(bv-av)/bv if isinstance(bv,(int,float)) and isinstance(av,(int,float)) and bv!=0 else MISSING
 formula(ws,row,dest,expr,expected,rid)
def note(ws,row,text):
 ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=6 if ws.title=='參數方法改動' else 8);c=ws.cell(row,1,text);c.font=Font(name=FONT,size=10,color='444444');c.alignment=Alignment(wrap_text=True,vertical='top');ws.row_dimensions[row].height=42
 return row+2

def table(ws,row,title,heads,data,rid,keyname):
 ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=6 if ws.title=='參數方法改動' else 8);ws.cell(row,1,title).font=Font(name=FONT,size=14,bold=True);ws.cell(row,1).alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[row].height=36
 header=row+1;put(ws,header,heads,rid)
 for c in ws[header][:len(heads)]:c.font=Font(name=FONT,size=13,bold=True);c.fill=PatternFill('solid',fgColor=GRAY);c.border=Border(bottom=Side(style='thin',color='999999'))
 for k,record in enumerate(data):put(ws,header+1+k,record,rid)
 end=header+len(data);RANGES.setdefault(str(rid),{})[keyname]=dict(sheet=ws.title,header=header,start=header+1,end=end,ref=f'A{header}:{col(len(heads))}{end}')
 return header+1,end+2

def setup(rid):
 b=Workbook();b.remove(b.active)
 for title in ['結果比較','參數方法改動','證據與圖表']:
  w=b.create_sheet(title);w.sheet_view.showGridLines=False;w.freeze_panes='C5';w.sheet_view.zoomScale=100
  widths=[30,28,28,36,44,25] if title=='參數方法改動' else [28,21,19,23,23,23,18,28]
  for i,width in enumerate(widths,1):w.column_dimensions[col(i)].width=width
  w.merge_cells('A1:F1' if title=='參數方法改動' else 'A1:H1');w['A1']=f'第{rid}次｜{TITLES[rid]}';w['A1'].font=Font(name=FONT,size=18,bold=True);w.row_dimensions[1].height=32
  note(w,2,'兩個dataset：Tongji四Plot＋ETH；本地5cm是參考實作，作者原作重跑待補。下降率=(基準−修改)/基準；成功率差異用pp。')
  w.page_setup.orientation='landscape';w.page_setup.paperSize=w.PAPERSIZE_A4;w.page_setup.fitToWidth=1;w.page_setup.fitToHeight=0;w.sheet_properties.pageSetUpPr.fitToPage=True;w.print_title_rows='1:3'
 BOOKS[rid]=b;CHARTS[str(rid)]=[];return b

def chart(rid,title,src,keyname,columns,anchor,unit,labels=None,percent=False,limit=None):
 target=BOOKS[rid]['證據與圖表'];tab=RANGES[str(rid)][keyname];c=BarChart();c.type='col';c.title=title;c.y_axis.title=unit;c.x_axis.title='場景／版本';c.width=27;c.height=11;c.style=10;c.y_axis.scaling.min=0
 if percent:c.y_axis.scaling.max=1;c.y_axis.numFmt='0%'
 if limit:c.y_axis.scaling.max=limit
 for j in columns:c.add_data(Reference(src,min_col=j,min_row=tab['header'],max_row=tab['end']),titles_from_data=True)
 c.set_categories(Reference(src,min_col=1,min_row=tab['start'],max_row=tab['end']))
 for i,s in enumerate(c.series):s.graphicalProperties.solidFill=['A6A6A6','305B81','666666','C5D4DF'][i%4];s.graphicalProperties.line.noFill=True
 c.dLbls=DataLabelList();c.dLbls.showVal=True;c.dLbls.showCatName=False;c.dLbls.showSerName=False;c.dLbls.showLegendKey=False;c.dLbls.showPercent=False;c.dLbls.showBubbleSize=False;c.dLbls.showLeaderLines=False;c.dLbls.dLblPos='outEnd';c.dLbls.numFmt='0%' if percent else '0.00';c.legend.position='b'
 if not percent:c.y_axis.numFmt='0.0' if unit=='e_p cm' else '0'
 target.add_chart(c,anchor)
 target.row_breaks.append(Break(id=int(anchor[1:])-1))
 CHARTS[str(rid)].append(dict(title=title,sheet=target.title,anchor=anchor,source_sheet=src.title,range=tab['ref'],min_row=tab['start'],max_row=tab['end'],columns=columns,unit=unit))

ROLES={1:('5cm＋SVD，補救關閉','10cm＋RANSAC＋弱配對補救'),2:('原DTM，10cm＋RANSAC＋補救','smooth；依列區分僅定位／完整提取'),3:('ETH固定粗配準，無refinement','相同粗配準＋多尺度refinement'),4:('多候選v2＋全表面驗證','同候選引擎v3＋結構驗證整組')}
def main_tables(rid):
 w=BOOKS[rid]['結果比較'];entries=SETS[rid];before_name,after_name=ROLES[rid]
 r=4;data=[]
 for e in entries:
  a,b=e['before'],e['after'];data.append([label(e['scene'])+('／'+e['scope'] if e['scope'] else ''),f'{(a or e["ref"])["pairs"]}配對×{len((a or e["ref"])["seeds"])}seed',1,None,None,None,None,None])
 start,r=table(w,r,'成功率：修改前='+before_name+'；修改後='+after_name,['dataset／場景','唯一配對×seed','論文成功率\n分母待補','5cm參考','修改前','修改後','vs論文 pp','vs修改前 pp'],data,rid,'success')
 for i,e in enumerate(entries,start):
  w.cell(i,3).number_format='0.0%'
  for j,g in [(4,e['ref']),(5,e['before']),(6,e['after'])]:
   if g:
    fmt=f'"{g["ok"]}/{g["n"]} ("0.0%")"';formula(w,i,j,f'=IF({g["n"]}>0,{g["ok"]}/{g["n"]},"待補")',g['rate'],rid,fmt)
   else:w.cell(i,j,MISSING)
  for j,base,bv in [(7,'C',1),(8,'E',e['before']['rate'] if e['before'] else None)]:
   val=(e['after']['rate']-bv)*100 if e['after'] and bv is not None else MISSING
   formula(w,i,j,f'=IF(AND(ISNUMBER({base}{i}),ISNUMBER(F{i})),(F{i}-{base}{i})*100,"待補")',val,rid,'0.0" pp"')
 r=note(w,r,'來源P：PDF p4 Table I；本地U/D/DF，精確篩選與CSV列號見source_manifest.json。每個seed是同一配對重跑，不能當成獨立森林。論文與專案成功門檻未確認一致；作者原作重跑待補。')
 data=[]
 for e in entries:
  a,b=e['before'],e['after'];rn=e['ref'];data.append([label(e['scene'])+('／'+e['scope'] if e['scope'] else ''),100*P_EP[e['scene']],cm(rn['ep']),cm(a['ep']) if a else MISSING,cm(b['ep']) if b else MISSING,None,None,'／'.join(f'{x["ep_n"]}/{x["n"]}' if x else MISSING for x in [rn,a,b])])
 start,r=table(w,r,'平均e_p：全部有限誤差，包含能計算的失敗；單位cm',['dataset／場景','論文平均e_p','5cm參考','修改前','修改後','vs論文下降率','vs前下降率','有限N：參考／前／後'],data,rid,'error')
 for i,e in enumerate(entries,start):
  a,b=e['before'],e['after']
  same=bool(a and b and {key(x) for x in a['rows'] if num(x['e_p_m']) is not None}=={key(x) for x in b['rows'] if num(x['e_p_m']) is not None})
  decline(w,i,6,f'B{i}',f'E{i}',100*P_EP[e['scene']],cm(b['ep']) if b else None,rid,allowed=bool(b and b['ep_n']==b['n']))
  decline(w,i,7,f'D{i}',f'E{i}',cm(a['ep']) if a else None,cm(b['ep']) if b else None,rid,allowed=same)
 r=note(w,r,'P：p4 Eq.(8)/Table I。全來源點e_p再跨case取mean，含有限失敗；缺變換不填0。修改後缺輸出則vs論文下降率待補；本地有限case集合不同則vs前下降率待補，另見第三頁共同成功比較。論文數字僅參考，協定未完全核對。')
 data=[]
 for e in entries:
  a,b=e['before'],e['after'];rn=e['ref'];data.append([label(e['scene'])+('／'+e['scope'] if e['scope'] else ''),rn['rot'],a['rot'] if a else MISSING,b['rot'] if b else MISSING,cm(rn['trans']),cm(a['trans']) if a else MISSING,cm(b['trans']) if b else MISSING,'論文e_R／e_t：待補'])
 _,r=table(w,r,'旋轉與平移：保留可計算失敗；不是論文額外指標',['dataset／場景','參考 e_R 度','前 e_R 度','後 e_R 度','參考 e_t cm','前 e_t cm','後 e_t cm','論文欄位狀態'],data,rid,'pose')
 r=note(w,r,'來源同上原始CSV之rot_err_deg、trans_err_m；e_R=R_gtᵀR_est的旋轉角，e_t=||t_est−t_gt||。Table I只列e_p及成功率，未列e_R/e_t，故原論文欄待補。本地成功為e_R<2°且e_t<50cm。')
 data=[]
 for e in entries:
  a,b=e['before'],e['after'];data.append([label(e['scene'])+('／'+e['scope'] if e['scope'] else ''),P_MATCH[e['scene']],P_TOTAL[e['scene']],e['ref']['time'],a['time'] if a else MISSING,b['time'] if b else MISSING,None,'見下方計時限制'])
 start,r=table(w,r,'時間：論文matching與total分開；本地皆為配對內部時間',['dataset／場景','論文matching ms','論文total s','參考online ms','前online ms','後online ms','本地下降率','範圍／硬體限制'],data,rid,'time')
 for i,e in enumerate(entries,start):
  a,b=e['before'],e['after'];decline(w,i,7,f'E{i}',f'F{i}',a['time'] if a else None,b['time'] if b else None,rid,allowed=rid in [1,3])
 r=note(w,r,'P：p4 Table II、p5 Table III；i7-9700F 3GHz，OpenMP執行緒數待補。U online含讀樹位／descriptor／matching／求解／refinement，不含幾何cache載入、原始點雲提取及GT評分；舊time_ms含descriptor＋matching＋recovery＋solver。論文速度下降率一律待補，範圍／硬體不同。')
 note(w,r,('R1/R3：前後為同一v1批次、4 OpenMP threads；可描述本地online代價，舊epoch CPU型號待補，不能與論文作硬體一致的速度宣稱。' if rid in [1,3] else 'R2：本地前後跨批次、舊硬體未完整記錄；R4：v2/v3跨epoch且採集並行度不同。以上本地online下降率也待補。R4第三頁另列同新epoch、串行冷啟動的代表配對。'))
 w.auto_filter.ref=RANGES[str(rid)]['success']['ref']

COMMON_PARAMS=[['Thscore（分數）','40','40','40','判斷高分／三角回退，維持固定','P p4／CFG'],['w（無單位）','10','10','10','FMscore權重，維持固定','P Eq7 p4／CFG'],['N（長度參數）','183','183','183','CN鄰域編碼解析度，維持固定','P p4／CFG'],['提取體素／法向半徑 m','0.01／0.10','0.01／0.10','0.01／0.10','樹幹提取不隨dataset改','P p2／BASECFG'],['提取verticality','>0.9','>0.9','>0.9','全部場景共用；與v3驗證門檻不同','P Eq4 p2／CFG']]
PARAMS={1:[['匹配距離容差 m','0.05','0.05','0.10','容許鄰域距離噪聲，可能增加錯配','P p4／CODE_MATCH'],['配準solver','待補（只述Euclidean transform）','Kabsch／SVD','Kabsch＋RANSAC','排除不符合剛體變換的對應','CODE_REG'],['配準RANSAC','待補','關閉','1000次；3D inlier<0.30m；seed=1','整組coarse設定，非圓柱RANSAC','CFG'],['弱配對補救','待補（非本專案補救）','關閉','原CN對應<6啟動','原本證據不足時再檢查共識','VP／CODE_REC'],['補救搜尋／支持','待補','無','前3候選／最多60邊／至少6支持','水平邊長差<0.20m，完整6DOF','VP'],['其他固定條件','作者細節待補','20cm分群、step距離抽樣','同前；不加ICP','多條件主比較只能歸因整組','CFG／BASECFG']],2:[['DTM cell m','待補（未列數值）','0.50','0.50','網格大小固定','DP／BASECFG'],['fit_radius 格','待補','2（約1m）','6（約3m）','局部平面使用更大地面鄰域','DP／CODE_DTM'],['fill_max_radius 格','待補','5（約2.5m）','6（約3m）','補洞上限同步變更，非單參數消融','DP'],['ground_select','待補','supported_lowest','相同','最低支持點選法固定','DP'],['推論輸入','作者完整流程','既有圓柱軸線','前4列僅改交點；Plot2另完整重提取','不同範圍不可混稱完整流程','DP／DFP'],['coarse設定','匹配0.05m；solver待補','0.10m＋RANSAC＋補救','相同','本次不動matching或solver','DP']],3:[['refinement架構','文中粗配準；未列本組ICP參數','無surface refinement','新增完整6DOF點對平面','修正剩餘旋轉／平移誤差','CODE_REF'],['refinement voxel m','未報告本組設定','無','0.50→0.20→0.10','新增階段，原提取1cm保持固定','CFG'],['每尺度最大迭代 次','待補','0','30→25→20','由粗至細，增加逐配對成本','CFG'],['法向／鄰域','待補','無新階段','半徑2.5h；至少6點；曲率≤0.1','抑制不穩定表面法向','CFG'],['幾何對應與robust','待補','無新階段','距離<3h、|normal dot|≥0.7、保留70%','一對一target＋Huber，整組加入','CFG'],['點數預算／退化保護','待補','無','來源最多20000；評分雙向各8000','Hessian ratio1e−8；惡化保留初值','CFG／CODE_REF'],['公平控制','作者程式待補','5cm參考／ETH固定粗配準','兩者都另有相同refinement','不能把ICP只給候選版','U']],4:[['候選擴張 v1→v2','待補','5候選／60邊／最多16姿態／精化3個','20候選／120邊／最多256姿態／精化8個','同時改搜尋及支持分數，不是單一top-k','method_v2_frozen／CODE_REF'],['v2新增水平初始值','未報告此設計','相容3D triples≤10000','最多7140雙點yaw seeds＋10000 triples','只限制初始值，最終仍6DOF','CFG'],['v2→v3候選引擎','待補','同上v2','保留同一v2引擎','本次主差異在驗證整組','V3P'],['驗證用表面','待補','全部方向','|normal.z|<0.30；兩邊各≥30點','降低大平面重疊影響；不足退回全部','CODE_V3'],['驗證NN gate m','待補','3h：1.50／0.60／0.30','min(h,0.20)：0.20／0.20／0.10','距離門檻同步收緊，不能獨立歸因篩選','CODE_V3'],['優化與接受判斷','待補','所有方向6DOF點對平面','優化同前；新評分也影響接受／保留','不只改最終排序','CODE_V3'],['DTM smooth','作者細節待補','原DTM，非report02 smooth','相同','避免把兩輪方法拼在一起','CFG']]}

def params(rid):
 w=BOOKS[rid]['參數方法改動'];_,row=table(w,4,'原論文明列值與本次實際改動',['參數／方法（單位）','原論文','本地修改前','本次修改後','目的與代價','來源'],COMMON_PARAMS+PARAMS[rid],rid,'params')
 note(w,row,'原論文未交代的值標待補，不能把本地預設數字當作者設定。各版本共用本地額外圓柱檢查、20cm分群與scanner origin=(0,0,0)假設；細節BASECFG。Thscore/w/N與step抽樣維持固定。')

def evidence(rid):
 w=BOOKS[rid]['證據與圖表'];r=4
 if rid==1:
  data=[[('5cm' if i<2 else '10cm')+'／'+('SVD' if i%2==0 else 'RANSAC'),g['ok'],g['n'],g['ok']/g['n'],avg(x['n_matches'] for x in g['rows']),avg(x['n_correct_matches'] for x in g['rows']),g['time'],'ETH seed11–20'] for i,g in enumerate(FA)]
  start,r=table(w,r,'同樹位2×2：容差與solver分開控制',['設定','成功','試驗','成功率','平均對應數','平均GT正確對應','平均online ms','範圍'],data,rid,'factorial')
  for i,g in enumerate(FA,start):formula(w,i,4,f'=IF(C{i}>0,B{i}/C{i},"待補")',g['rate'],rid)
  r=note(w,r,'F raw；提取條件固定。GT正確對應指套GT後XY距離<0.30m，僅事後診斷；不是配準RANSAC的3D inlier。seed11–20、21–30、31–40分輪報告。')
  data=[['補救前',VA['ok'],VA['n'],VA['rate'],cm(VJ['before']),VA['time'],'共同成功142','V seed21–30'],['補救後',VB['ok'],VB['n'],VB['rate'],cm(VJ['after']),VB['time'],'新增8／退步0','V seed21–30']]
  start,r=table(w,r,'固定10cm／RANSAC後，只新增弱配對補救',['版本','成功','試驗','成功率','共同成功mean e_p cm','online ms','成對證據','範圍'],data,rid,'recovery')
  for i,g in enumerate([VA,VB],start):formula(w,i,4,f'=IF(C{i}>0,B{i}/C{i},"待補")',g['rate'],rid)
  r=note(w,r,f'V raw重算mean：{cm(VJ["before"]):.3f}→{cm(VJ["after"]):.3f}cm。paired_validation.json的common_ep欄由原script計算median，不能直接當mean。本表未使用該欄作平均。')
  data=[]
  for name in ['fresh_matching','fresh_vertical085_matching','fresh_ransac999_matching']:
   g=group('F:negative:'+name,'F',F,lambda x,n=name:x['method']==n and x['solver']=='svd+ransac',dict(method=name,solver='svd+ransac',seeds='11–20'))
   data.append([name,g['ok'],g['n'],g['rate'],0.85 if '085' in name else 0.9,3000 if '999' in name else 1000,0.999 if '999' in name else 0.99,'同10cm＋RANSAC'])
  start,r=table(w,r,'已有負面對照：提取放寬或加資源不一定更好',['設定','成功','試驗','成功率','verticality','圓柱迭代上限','圓柱機率','固定coarse'],data,rid,'negative')
  for i,record in enumerate(data,start):formula(w,i,4,f'=IF(C{i}>0,B{i}/C{i},"待補")',record[3],rid)
  r=note(w,r,'F raw，seed11–20；142→130或141/150。圓柱迭代上限與停止機率同時修改，只能歸因該整組資源設定。其他探索版本見實驗盤點與數據口徑.txt。')
  anchor=r+2;chart(rid,'粗配準整組：五場景成功率',BOOKS[rid]['結果比較'],'success',[4,6],f'A{anchor}','成功率',percent=True)
  chart(rid,'ETH controlled 2×2（同150試驗）',w,'factorial',[4],f'A{anchor+24}','成功率',percent=True)
 elif rid==2:
  data=[]
  for e in SETS[2][:-1]:
   j=joint(e['before'],e['after']);data.append([label(e['scene'])+'／'+e['scope'],j['gained'],j['lost'],j['common'],cm(j['before']),cm(j['after']),None,'固定同pair／seed'])
  start,r=table(w,r,'新增成功與退步並列；共同成功精度',['範圍','新增成功','退步失敗','共同成功N','前mean e_p cm','後mean e_p cm','下降率','集合'],data,rid,'paired')
  for i,e in enumerate(SETS[2][:-1],start):
   j=joint(e['before'],e['after']);decline(w,i,7,f'E{i}',f'F{i}',cm(j['before']),cm(j['after']),rid)
  r=note(w,r,'D/DF raw；Plot2完整列=新增23、退步9、共同成功25。其他三Plot未完成新DTM全流程；ETH此改動待補。')
  h0=st.median(float(x['old_median_abs_z_m']) for x in H);h1=st.median(float(x['new_median_abs_z_m']) for x in H);hn=sum(float(x['new_median_abs_z_m'])<float(x['old_median_abs_z_m']) for x in H)
  _,r=table(w,r,'固定原CN配對的高度診斷（不是e_p）',['指標','原值cm','新值cm','改善案例','總案例','用途','證據強度','來源'],[['每case |Δz|中位數，再跨case中位數',100*h0,100*h1,hn,100,'事後GT診斷','支持部分高度不一致改善','H']],rid,'height')
  r=note(w,r,'H：原CN配對保持不變；沒有人工地形標註，因此無法證明絕對DTM誤差降低。fit_radius與fill_max_radius一起改，無單參數因果結論。')
  data=[['原設定',34/100],['只改交點',48/100],['完整重提取',48/100]]
  _,r=table(w,r,'Plot2三種流程，成功集合不同',['方法','成功率'],data,rid,'full_compare')
  for i,k in enumerate([34,48,48],RANGES[str(rid)]['full_compare']['start']):formula(w,i,2,f'={k}/100',k/100,rid)
  anchor=r+1;chart(rid,'Plot2：同為48%不代表同一批成功',w,'full_compare',[2],f'A{anchor}','成功率',percent=True)
  tab=RANGES['2']['success'];saved=tab['end'];tab['end']=tab['start']+3
  chart(rid,'四Plot定位測試（非全新提取流程）',BOOKS[rid]['結果比較'],'success',[5,6],f'A{anchor+24}','成功率',percent=True);tab['end']=saved
 elif rid==3:
  data=[]
  for e in SETS[3]:
   j=joint(e['before'],e['after']);data.append([label(e['scene']),j['gained'],j['lost'],j['common'],cm(j['before']),cm(j['after']),None,'同粗配準、同樹位'])
  start,r=table(w,r,'相同初始化：refinement前後的共同成功案例',['場景','新增成功','退步失敗','共同成功N','前mean e_p cm','後mean e_p cm','下降率','控制'],data,rid,'paired')
  for i,e in enumerate(SETS[3],start):
   j=joint(e['before'],e['after']);decline(w,i,7,f'E{i}',f'F{i}',cm(j['before']),cm(j['after']),rid)
  r=note(w,r,'U raw；primary使用eth10_raw→eth10_refined。不是把兩個成功集合的平均直接相減。無法輸出變換的case不能靠此局部微調自行產生初值。')
  data=[]
  for s in SCENES:
   a=ug(s,'reference5_raw');b=ug(s,'reference5_refined');j=joint(a,b);data.append([label(s),a['ok'],b['ok'],a['n'],cm(a['ep']),cm(b['ep']),j['gained'],j['lost']])
  _,r=table(w,r,'公平控制：5cm參考也加完全相同refinement',['場景','參考成功','加微調成功','試驗','前有限mean cm','後有限mean cm','新增成功','退步'],data,rid,'fair')
  r=note(w,r,'U reference5_raw/reference5_refined；Plot4仍有缺變換及大錯配，mean有限集合必須連同成功數閱讀。新增refinement是一組設定，未逐一隔離三尺度、Huber、trim的效果。')
  anchor=r+1;chart(rid,'固定粗配準：微調前後成功率',BOOKS[rid]['結果比較'],'success',[5,6],f'A{anchor}','成功率',percent=True)
  chart(rid,'共同成功案例的mean e_p（同一集合）',w,'paired',[5,6],f'A{anchor+24}','e_p cm')
  chart(rid,'逐配對online時間（含refinement）',BOOKS[rid]['結果比較'],'time',[5,6],f'A{anchor+48}','ms')
 else:
  data=[]
  for m in ['multi_refined','multi_v2_refined','multi_v3_refined']:
   g=ug('Plot4',m);data.append([m,g['ok'],g['n'],g['rate'],cm(g['ep']),cm(g['ep_median']),g['ep_n'],'U Plot4 seed31–40'])
  start,r=table(w,r,'Plot4：擴張失敗與修正都保留',['版本','成功','試驗','成功率','所有有限mean cm','有限median cm','有限N','範圍'],data,rid,'versions')
  for i,record in enumerate(data,start):formula(w,i,4,f'=IF(C{i}>0,B{i}/C{i},"待補")',record[3],rid)
  r=note(w,r,'U；v1/v2/v3全100例都有有限e_p，因此此處mean可直接比較。v2的median仍小，卻有22%失敗，不能只報median。')
  data=[]
  for s in SCENES:
   a=ug(s,'reference5_refined_v3');b=ug(s,'eth10_refined_v3');c=ug(s,'multi_v3_refined');j=joint(b,c);data.append([label(s),a['ok'],b['ok'],c['ok'],c['n'],j['gained'],j['lost'],'同一v3 backend'])
  _,r=table(w,r,'公平backend：多初始值額外救回多少？',['場景','5cm＋同微調','ETH固定＋同微調','統一v3成功','試驗','較ETH新增','較ETH退步','控制'],data,rid,'fair')
  r=note(w,r,'U；總數507/550、546/550、550/550。統一版相對ETH固定＋同v3微調額外4次，全部在Plot4；不把全部改善都歸給新候選。')
  rr=[next(x for x in U if x['scene']=='Plot4' and x['pair']=='S4-S2' and x['seed']=='31' and x['method']==m) for m in ['multi_refined','multi_v2_refined','multi_v3_refined']]
  data=[[x['method'],float(x['geo_overlap']) if i<2 else MISSING,100*float(x['e_p_m']),int(x['success']),'S4-S2',31,'all surface' if i<2 else '新定義不直接比','U'] for i,x in enumerate(rr)]
  _,r=table(w,r,'一個已完成案例：overlap較大卻錯得更遠',['版本','同定義overlap','e_p cm','成功0/1','配對','seed','overlap範圍','來源'],data,rid,'case')
  r=note(w,r,'v1/v2舊geo_overlap定義相同；v3改了點集合與gate，同定義overlap待補。這支持舊驗證分數有歧義；地面重複是機制解釋，尚無只改垂直篩選／只改gate的2×2。')
  data=[]
  for ds in ['ETH','Tongji']:
   a=next(x for x in COLD if x['dataset']==ds and x['method']=='reference5_refined_v3');b=next(x for x in COLD if x['dataset']==ds and x['method']=='multi_v3_refined');assert a['runtime_epoch']==b['runtime_epoch']
   data.append([ds+'／'+a['scene']+'/'+a['pair'],a['wall_s'],b['wall_s'],None,a['peak_child_rss_kib']/1024,b['peak_child_rss_kib']/1024,1,'同epoch、native串行'])
  start,r=table(w,r,'代表配對冷啟動：兩站重提取＋geometry＋配對',['資料／固定配對','參考＋同微調 s','統一v3 s','下降率','參考峰值 MiB','v3峰值 MiB','每法測量次數','範圍'],data,rid,'cold')
  for i,row in enumerate(data,start):decline(w,i,4,f'B{i}',f'C{i}',row[1],row[2],rid)
  r=note(w,r,'C/HW：Xeon Platinum8573C、CPU配額4、32GiB、每程序4OMP。seed31，每資料每法1個固定pair，OS cache未清。下降率是這次觀測差異，不證速度排名；峰值為最大native子程序RSS，不含Python及OS page cache。')
  anchor=r+1;chart(rid,'Plot4：v1→v2→v3成功率',w,'versions',[4],f'A{anchor}','成功率',percent=True)
  chart(rid,'五場景e_p：論文報告值與統一v3',BOOKS[rid]['結果比較'],'error',[2,5],f'A{anchor+24}','e_p cm')
  chart(rid,'冷啟動代表配對（每法各1次）',w,'cold',[2,3],f'A{anchor+48}','seconds')
 for sheet in BOOKS[rid]:
  maxchart=max([int(x['anchor'][1:])+23 for x in CHARTS[str(rid)] if x['sheet']==sheet.title]+[sheet.max_row])
  sheet.print_area=f'A1:{"F" if sheet.title=="參數方法改動" else "H"}{maxchart}'
 return r

for rid in TITLES:
 setup(rid);main_tables(rid);params(rid);evidence(rid)
 for name in ['error','pose','time']:
  section=RANGES[str(rid)][name];BOOKS[rid]['結果比較'].row_breaks.append(Break(id=section['header']-2))
 BOOKS[rid]['參數方法改動'].page_setup.fitToHeight=1
 # All missing data remain text, never a fabricated numeric zero.
 BOOKS[rid].calculation.fullCalcOnLoad=True
 BOOKS[rid].save(RUNS/f'report_{rid:02d}_{TITLES[rid]}.xlsx')

# Detailed provenance stays outside the projection workbooks.
snapshots=OUT/'evidence';snapshots.mkdir(exist_ok=True)
for sid,record in SOURCES.items():
 if sid=='P' or not record['path'].startswith('results/'):continue
 path=ROOT/record['path']
 if not path.exists():path=snapshots/(sid+'_'+Path(record['path']).name)
 if path.suffix not in ['.csv','.json']:continue
 destination=snapshots/(sid+'_'+Path(record['path']).name)
 if path.resolve()!=destination.resolve():shutil.copyfile(path,destination)
 record['snapshot']=str(destination.relative_to(OUT));assert sha(destination)==record['sha256']
manifest=dict(generated_at=datetime.now(ZoneInfo('Asia/Taipei')).isoformat(timespec='seconds'),source_commit=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),datasets={'Tongji-trees':['Plot1','Plot2','Plot3','Plot4'],'ETH-tree':['ETH']},paper_values=dict(ep_m=P_EP,matching_ms=P_MATCH,total_s=P_TOTAL,success_rate=1.0,pair_counts=MISSING,rotation_error=MISSING,success_threshold=MISSING),sources=SOURCES,groups=GROUPS,table_ranges=RANGES,charts=CHARTS,validation_expectations=CHECKS)
(OUT/'source_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
print('Generated 4 source workbooks and traceable manifest',flush=True)
