#!/usr/bin/env python3
"""Position-only DTM study; isolate axes from terrain and keep the old run intact."""
import argparse
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
import json
import math
from pathlib import Path
import statistics as stats
import subprocess
import threading
import time
import numpy as np
from openpyxl import Workbook
from make_study_workbook import table
import tongji_study as old

ROOT, DATA = old.ROOT, old.DATA
OUT = ROOT/'results/tongji_trees/dtm_height_20261025'
RUNS = OUT/'runs'
PROFILES = ['baseline','smooth','coarse','quantile1']
PARAMETERS = {
    'baseline': dict(cell=.5, fit_radius=2, fill_max_radius=5, ground_select='supported_lowest', percentile=5),
    'smooth': dict(cell=.5, fit_radius=6, fill_max_radius=6, ground_select='supported_lowest', percentile=5),
    'coarse': dict(cell=1., fit_radius=3, fill_max_radius=6, ground_select='supported_lowest', percentile=5),
    'quantile1': dict(cell=.5, fit_radius=2, fill_max_radius=5, ground_select='percentile', percentile=1),
}
NAMES = {'baseline':'原地形設定','smooth':'較大平滑範圍','coarse':'較粗地形網格','quantile1':'第1百分位地面'}
LOCK = threading.Lock()

def execute(label,args):
    path=RUNS/'logs'/f'{label}.log';path.parent.mkdir(parents=True,exist_ok=True)
    start=time.monotonic();cmd=([str(old.TOOL)] if old.TOOL else [])+list(map(str,args))
    print('START',label,flush=True)
    with path.open('w') as log:r=subprocess.run(cmd,cwd=ROOT,stdout=log,stderr=subprocess.STDOUT)
    with LOCK:
        with (RUNS/'commands.jsonl').open('a') as f:f.write(json.dumps(dict(label=label,command=cmd,exit_code=r.returncode,seconds=time.monotonic()-start))+'\n')
    if r.returncode:raise RuntimeError(path)
    print('DONE',label,round(time.monotonic()-start,1),flush=True)

def parallel(fn,jobs):
    with ThreadPoolExecutor(max_workers=4) as pool:return list(pool.map(fn,jobs))

def freeze():
    old.plan();OUT.mkdir(parents=True,exist_ok=True);RUNS.mkdir(exist_ok=True)
    p=dict(profiles=PARAMETERS,development='Plot2, all10 pairs, seeds31-35',
           within_plot_validation='Plot2, same stations, seeds36-40; not independent forests',
           cross_plot_validation='Plot1,3,4, all30 pairs, seeds31-40; their old baseline outcomes are already known',
           selection='maximum development successes of svd+ransac; tie order smooth,coarse,quantile1; retain baseline unless strictly improved',
           primary_solver='svd+ransac',secondary_solver='svd',threshold=.10,recovery='evidence',recovery_trigger=6,
           success='rotation<2deg AND translation<0.5m',
           isolation='same existing cylinders and accepted tree list; only recompute axis/DTM intersection; not a full extraction rerun',
           failed_intersections='drop from candidate output; preserve baseline; expose every drop',
           evaluation='development skips point error; selected profile gets full source-cloud point error after freeze',
           no_gt_in_repositioning=True,no_gt_in_inference=True,
           previous_plan_sha256=old.sha(old.OUT/'plan.json'),
           fixed_sha256={str(path.relative_to(ROOT)):old.sha(path) for path in [ROOT/'build/cm_dtm_reposition',ROOT/'apps/cm_dtm_reposition.cpp',ROOT/'build/cm_register',ROOT/'config/default.ini']},
           prior_results_sha256=old.sha(old.OUT/'registration_rows.csv'))
    path=OUT/'plan.json'
    if path.exists():assert {k:v for k,v in json.loads(path.read_text()).items() if k!='frozen_at_utc'}==p
    else:p['frozen_at_utc']=datetime.now(timezone.utc).isoformat();path.write_text(json.dumps(p,indent=2)+'\n')
    print('FROZEN development and validation protocol',flush=True)

def plan():
    p=json.loads((OUT/'plan.json').read_text());old.plan()
    for path,digest in p['fixed_sha256'].items():assert old.sha(ROOT/path)==digest
    assert old.sha(old.OUT/'registration_rows.csv')==p['prior_results_sha256']
    return p

def reposition(plots,profiles):
    plan()
    def one(job):
        plot,scan=job;dest=RUNS/'positions'/f'Plot{plot}'/f'S{scan}';dest.mkdir(parents=True,exist_ok=True)
        marker=dest/('complete_'+'-'.join(profiles)+'.json')
        outputs=[dest/profile/f'step_seed{seed}_trees.csv' for profile in profiles for seed in old.SEEDS]
        if marker.exists() and all(p.exists() for p in outputs):return
        original=old.RUNS/'extraction'/f'Plot{plot}'/f'S{scan}'
        batch=dest/'jobs.txt';batch.write_text(''.join(f'{original}/step_seed{seed}_trees.csv {original}/step_seed{seed}_clusters.csv\n' for seed in old.SEEDS))
        execute(f'position_P{plot}_S{scan}',[ROOT/'build/cm_dtm_reposition',f'--input={DATA}/Plot{plot}/S{scan}.ply',f'--batch={batch}',f'--out={dest}',f'--profiles={",".join(profiles)}','--config=config/default.ini','--scanner.x0=0','--scanner.y0=0','--scanner.z0=0'])
        assert all(p.exists() for p in outputs)
        marker.write_text(json.dumps(dict(profiles=profiles,files={p.name+':'+p.parent.name:old.sha(p) for p in outputs}))+'\n')
    parallel(one,[(plot,scan) for plot in plots for scan in range(1,6)])

def registration(profile,plots,seeds,phase,skip_point_error):
    p=plan();pairs=json.loads((old.OUT/'plan.json').read_text())['pairs']
    def one(job):
        plot,source=job;dest=RUNS/phase/profile/f'Plot{plot}';dest.mkdir(parents=True,exist_ok=True);files=[];lines=[]
        for pair in pairs:
            if pair['plot']!=plot or pair['source']!=source:continue
            target=pair['target']
            for seed in seeds:
                base=RUNS/'positions'/f'Plot{plot}';out=dest/f'{seed}_S{source}-S{target}.csv'
                src=base/f'S{source}'/profile/f'step_seed{seed}_trees.csv';tgt=base/f'S{target}'/profile/f'step_seed{seed}_trees.csv'
                assert src.exists() and tgt.exists()
                label=f'Plot{plot}_{seed}_S{source}-S{target}'
                lines.append(f'{src} {tgt} {DATA}/Plot{plot}/S{source}.ply {DATA}/Plot{plot}/S{source}_S{target}.txt {out} {label}')
                files.append((seed,target,out))
        batch=dest/f'S{source}.txt';batch.write_text('\n'.join(lines)+'\n');marker=dest/f'S{source}.complete'
        if not marker.exists() or not all(p.exists() for _,_,p in files):
            args=[ROOT/'build/cm_register',f'--batch={batch}','--thresholds=0.10','--recovery=evidence','--recovery_trigger=6']
            if skip_point_error:args+=['--skip_point_error=1']
            execute(f'{phase}_{profile}_P{plot}_S{source}',args);marker.write_text('completed\n')
        return [dict(profile=profile,plot=plot,seed=seed,pair=f'S{source}-S{target}',**r) for seed,target,path in files for r in old.read_csv(path)]
    return [r for part in parallel(one,[(plot,source) for plot in plots for source in range(2,6)]) for r in part]

def develop():
    reposition([2],PROFILES)
    rows=[r for profile in PROFILES for r in registration(profile,[2],range(31,36),'development',True)]
    assert len(rows)==400;old.write_csv(RUNS/'development_rows.csv',rows)
    baseline=[r for r in rows if r['profile']=='baseline']
    prior={(int(r['plot']),int(r['seed']),r['pair'],r['solver']):r for r in old.read_csv(old.OUT/'registration_rows.csv') if r['method']=='recovery10cm'}
    fields=['success','transform_ok','n_matches','n_correct_matches','recovery_triggered','recovery_accepted']+[f'r{i}{j}' for i in range(3) for j in range(3)]+[f't{i}' for i in range(3)]
    for r in baseline:
        original=prior[(int(r['plot']),int(r['seed']),r['pair'],r['solver'])]
        assert all(r[k]==original[k] for k in fields),'baseline output differs'
    summary=[]
    for profile in PROFILES:
        for solver in ['svd','svd+ransac']:
            rr=[r for r in rows if r['profile']==profile and r['solver']==solver]
            summary.append(dict(profile=profile,solver=solver,cases=len(rr),success=sum(int(r['success']) for r in rr)))
    old.write_csv(OUT/'development_summary.csv',summary)
    print(json.dumps(summary),flush=True)

def select():
    plan();summary=old.read_csv(OUT/'development_summary.csv')
    scores={r['profile']:int(r['success']) for r in summary if r['solver']=='svd+ransac'}
    best=max(PROFILES[1:],key=lambda profile:scores[profile]);chosen=best if scores[best]>scores['baseline'] else 'baseline'
    decision=dict(selected=chosen,scores=scores,development_rows_sha256=old.sha(RUNS/'development_rows.csv'),frozen_at_utc=datetime.now(timezone.utc).isoformat(),validation_results_not_used=True)
    path=OUT/'selected.json'
    if path.exists():assert {k:v for k,v in json.loads(path.read_text()).items() if k!='frozen_at_utc'}=={k:v for k,v in decision.items() if k!='frozen_at_utc'}
    else:path.write_text(json.dumps(decision,indent=2)+'\n')
    print('SELECTED',chosen,scores,flush=True)

def selected():
    plan();p=json.loads((OUT/'selected.json').read_text());assert old.sha(RUNS/'development_rows.csv')==p['development_rows_sha256'];return p['selected']

def validate():
    profile=selected()
    if profile=='baseline':
        rows=[dict(profile='baseline',**r) for r in old.read_csv(old.OUT/'registration_rows.csv') if r['method']=='recovery10cm']
    else:
        reposition([1,3,4],['baseline',profile])
        rows=registration(profile,[1,2,3,4],old.SEEDS,'final',False)
    assert len(rows)==800;old.write_csv(RUNS/'selected_rows.csv',rows)

def full():
    """Follow-up on Plot2: recompute every extraction stage with the frozen choice."""
    profile=selected();assert profile=='smooth','This follow-up protocol was specified for smooth only'
    path=OUT/'full_pipeline_plan.json'
    protocol=dict(selected_profile=profile,plots=[2],seeds=old.SEEDS,fit_radius=6,fill_max_radius=6,
                  selection_frozen_sha256=old.sha(OUT/'selected.json'),
                  scope='post-selection full-pipeline follow-up; same Plot2, not new independent validation',
                  parameters_reselected=False)
    if path.exists():assert {k:v for k,v in json.loads(path.read_text()).items() if k!='frozen_at_utc'}==protocol
    else:protocol['frozen_at_utc']=datetime.now(timezone.utc).isoformat();path.write_text(json.dumps(protocol,indent=2)+'\n')
    def one(scan):
        dest=RUNS/'full_extraction/Plot2'/f'S{scan}';dest.mkdir(parents=True,exist_ok=True)
        outputs=[dest/f'step_seed{seed}_{kind}.csv' for seed in old.SEEDS for kind in ['trees','clusters','retention_bins']]
        marker=dest/'complete.json'
        if marker.exists() and all(p.exists() for p in outputs):return
        execute(f'full_extract_P2_S{scan}',[ROOT/'build/cm_extract_batch',f'--input={DATA}/Plot2/S{scan}.ply',f'--out={dest}',
            '--config=config/default.ini','--version=step','--cluster.tolerance=0.20','--scanner.x0=0','--scanner.y0=0','--scanner.z0=0',
            '--experiment.seed_base=31','--experiment.n_seeds=10','--dtm.fit_radius=6','--dtm.fill_max_radius=6'])
        assert all(p.exists() for p in outputs);marker.write_text('completed\n')
    parallel(one,range(1,6))
    pairs=json.loads((old.OUT/'plan.json').read_text())['pairs']
    def register_source(source):
        dest=RUNS/'full_registration';dest.mkdir(exist_ok=True);lines=[];outputs=[]
        for pair in pairs:
            if pair['plot']!=2 or pair['source']!=source:continue
            target=pair['target']
            for seed in old.SEEDS:
                base=RUNS/'full_extraction/Plot2';out=dest/f'{seed}_S{source}-S{target}.csv';label=f'Plot2_{seed}_S{source}-S{target}'
                lines.append(f'{base}/S{source}/step_seed{seed}_trees.csv {base}/S{target}/step_seed{seed}_trees.csv {DATA}/Plot2/S{source}.ply {DATA}/Plot2/S{source}_S{target}.txt {out} {label}')
                outputs.append((seed,target,out))
        batch=dest/f'S{source}.txt';batch.write_text('\n'.join(lines)+'\n');marker=dest/f'S{source}.complete'
        if not marker.exists() or not all(p.exists() for _,_,p in outputs):
            execute(f'full_register_P2_S{source}',[ROOT/'build/cm_register',f'--batch={batch}','--thresholds=0.10','--recovery=evidence','--recovery_trigger=6']);marker.write_text('completed\n')
        return [dict(method='full_pipeline',profile=profile,plot=2,seed=seed,pair=f'S{source}-S{target}',stratum=stratum(dict(plot=2,seed=seed)),**r) for seed,target,p in outputs for r in old.read_csv(p)]
    rows=[r for part in parallel(register_source,range(2,6)) for r in part];assert len(rows)==200
    old.write_csv(OUT/'full_pipeline_rows.csv',rows)
    print('FULL PIPELINE ROWS',len(rows),flush=True)

def stratum(row):
    return ('development' if int(row['seed'])<=35 else 'within_plot_validation') if int(row['plot'])==2 else 'cross_plot_validation'

def audit():
    """Fixed-pair height diagnosis and unchanged-input checks; evaluation only."""
    assert selected()=='smooth'
    expected={(int(r['plot']),int(r['scan']),int(r['seed'])):r['sha256'] for r in old.read_csv(old.OUT/'extraction_summary.csv')}
    inputs=[];tree_count=0;max_residual=0.;baseline_files=0
    for plot in range(1,5):
        for scan in range(1,6):
            for seed in old.SEEDS:
                original=old.RUNS/f'extraction/Plot{plot}/S{scan}';tree=original/f'step_seed{seed}_trees.csv';cluster=original/f'step_seed{seed}_clusters.csv'
                digest=old.sha(tree);assert digest==expected[(plot,scan,seed)]
                inputs.append(dict(plot=plot,scan=scan,seed=seed,tree_sha256=digest,cluster_sha256=old.sha(cluster)))
                base=RUNS/f'positions/Plot{plot}/S{scan}/baseline'/tree.name
                assert old.read_csv(base)==old.read_csv(tree);baseline_files+=1
                for r in old.read_csv(str(base)+'.audit.csv'):
                    assert r['intersection_ok']=='1';tree_count+=1;max_residual=max(max_residual,abs(float(r['old_dtm_residual_m'])))
    (OUT/'input_axes_audit.json').write_text(json.dumps(dict(previous_200_tree_hashes_unchanged=True,inputs=inputs),indent=2)+'\n')
    (OUT/'baseline_position_audit.json').write_text(json.dumps(dict(scans=20,seed_files_checked=baseline_files,baseline_coordinate_rows_identical=True,tree_rows=tree_count,max_abs_dtm_residual_m=max_residual),indent=2)+'\n')
    pose=[]
    for scan in range(2,6):
        t=np.loadtxt(DATA/f'Plot2/S{scan}_S1.txt');pose.append(dict(scan=scan,gt_vertical_axis_tilt_deg=float(np.degrees(np.arccos(np.clip(t[2,2],-1,1))))))
    old.write_csv(OUT/'gt_tilt_diagnostic.csv',pose)
    dest=RUNS/'height_diagnostics';dest.mkdir(exist_ok=True);pairs=[p for p in json.loads((old.OUT/'plan.json').read_text())['pairs'] if p['plot']==2]
    lines=[];traces=[]
    for pair in pairs:
        source,target=pair['source'],pair['target']
        for seed in old.SEEDS:
            base=old.RUNS/'extraction/Plot2';label=f'{seed}_S{source}-S{target}'
            lines.append(f'{base}/S{source}/step_seed{seed}_trees.csv {base}/S{target}/step_seed{seed}_trees.csv {DATA}/Plot2/S{source}.ply {DATA}/Plot2/S{source}_S{target}.txt {dest}/{label}.csv {label}')
            traces.append(dest/f'{label}_thr0.100000_descriptors.csv')
    batch=dest/'jobs.txt';batch.write_text('\n'.join(lines)+'\n')
    if not all(p.exists() for p in traces):execute('height_trace_P2',[ROOT/'build/cm_register',f'--batch={batch}','--thresholds=0.10','--skip_point_error=1',f'--diagnostics={dest}'])
    rows=[]
    for pair in pairs:
        source,target=pair['source'],pair['target'];gt=np.loadtxt(DATA/f'Plot2/S{source}_S{target}.txt')
        for seed in old.SEEDS:
            matched=[r for r in old.read_csv(dest/f'{seed}_S{source}-S{target}_thr0.100000_descriptors.csv') if r['accepted']=='1']
            original=[old.read_csv(old.RUNS/f'extraction/Plot2/S{k}/step_seed{seed}_trees.csv') for k in [source,target]]
            newer=[{r['cluster_id']:r for r in old.read_csv(RUNS/f'positions/Plot2/S{k}/smooth/step_seed{seed}_trees.csv')} for k in [source,target]]
            before=[];after=[]
            def xyz(row):return np.array([float(row[k]) for k in ['x','y','z']])
            for r in matched:
                a,b=original[0][int(r['source'])],original[1][int(r['target'])]
                if a['cluster_id'] not in newer[0] or b['cluster_id'] not in newer[1]:continue
                a2,b2=newer[0][a['cluster_id']],newer[1][b['cluster_id']]
                before.append(gt[:3,:3]@xyz(a)+gt[:3,3]-xyz(b));after.append(gt[:3,:3]@xyz(a2)+gt[:3,3]-xyz(b2))
            if not before:continue
            x,y=np.array(before),np.array(after)
            rows.append(dict(plot=2,seed=seed,pair=f'S{source}-S{target}',original_pairs=len(matched),compared_pairs=len(x),old_median_abs_z_m=float(np.median(abs(x[:,2]))),new_median_abs_z_m=float(np.median(abs(y[:,2]))),old_median_xy_m=float(np.median(np.linalg.norm(x[:,:2],axis=1))),new_median_xy_m=float(np.median(np.linalg.norm(y[:,:2],axis=1)))))
    assert len(rows)==100;old.write_csv(OUT/'fixed_pair_height_diagnostic.csv',rows)

def report():
    profile=selected();fresh=old.read_csv(RUNS/'selected_rows.csv');prior=[r for r in old.read_csv(old.OUT/'registration_rows.csv') if r['method']=='recovery10cm']
    rows=[]
    for method,source in [('baseline',prior),('selected',fresh)]:
        for raw in source:
            r=dict(raw);r.update(method=method,profile='baseline' if method=='baseline' else profile,stratum=stratum(raw))
            assert int(r['success'])==int(r['transform_ok']=='1' and float(r['rot_err_deg'])<2 and float(r['trans_err_m'])<.5)
            rows.append(r)
    keys=list(dict.fromkeys(k for r in rows for k in r));rows=[{k:r.get(k,'') for k in keys} for r in rows];old.write_csv(OUT/'registration_rows.csv',rows)
    summary=[];paired=[]
    groups=[('all',lambda r:True)]+[(f'Plot{p}',lambda r,p=p:int(r['plot'])==p) for p in range(1,5)]+[(s,lambda r,s=s:r['stratum']==s) for s in ['development','within_plot_validation','cross_plot_validation']]
    for group,where in groups:
        for solver in ['svd','svd+ransac']:
            for method in ['baseline','selected']:
                rr=[r for r in rows if where(r) and r['method']==method and r['solver']==solver];good=[float(r['e_p_m']) for r in rr if r['success']=='1']
                summary.append(dict(group=group,method=method,profile='baseline' if method=='baseline' else profile,solver=solver,cases=len(rr),success=sum(int(r['success']) for r in rr),rate=sum(int(r['success']) for r in rr)/len(rr),mean_success_ep_m=stats.mean(good) if good else ''))
            left={(int(r['plot']),int(r['seed']),r['pair']):r for r in rows if where(r) and r['method']=='baseline' and r['solver']==solver}
            right={(int(r['plot']),int(r['seed']),r['pair']):r for r in rows if where(r) and r['method']=='selected' and r['solver']==solver};assert left.keys()==right.keys()
            common=[(a,right[k]) for k,a in left.items() if a['success']==right[k]['success']=='1']
            paired.append(dict(group=group,solver=solver,cases=len(left),gained=sum(a['success']=='0' and right[k]['success']=='1' for k,a in left.items()),lost=sum(a['success']=='1' and right[k]['success']=='0' for k,a in left.items()),common_success=len(common),mean_ep_change_m=stats.mean(float(b['e_p_m'])-float(a['e_p_m']) for a,b in common) if common else ''))
    old.write_csv(OUT/'summary.csv',summary);old.write_csv(OUT/'paired_comparison.csv',paired)
    full_rows=old.read_csv(OUT/'full_pipeline_rows.csv') if (OUT/'full_pipeline_rows.csv').exists() else []
    full_summary=[];full_paired=[]
    if full_rows:
        assert len(full_rows)==200
        p2=[dict(r,method='position_only' if r['method']=='selected' else 'baseline') for r in rows if int(r['plot'])==2]+full_rows
        for group,where in [('Plot2_all',lambda r:True),('development',lambda r:r['stratum']=='development'),('within_plot_validation',lambda r:r['stratum']=='within_plot_validation')]:
            for solver in ['svd','svd+ransac']:
                for method in ['baseline','position_only','full_pipeline']:
                    rr=[r for r in p2 if where(r) and r['solver']==solver and r['method']==method]
                    for r in rr:assert int(r['success'])==int(r['transform_ok']=='1' and float(r['rot_err_deg'])<2 and float(r['trans_err_m'])<.5)
                    ep=[float(r['e_p_m']) for r in rr if r['success']=='1']
                    full_summary.append(dict(group=group,method=method,solver=solver,cases=len(rr),success=sum(int(r['success']) for r in rr),rate=sum(int(r['success']) for r in rr)/len(rr),mean_success_ep_m=stats.mean(ep) if ep else ''))
                left={(r['seed'],r['pair']):r for r in p2 if where(r) and r['solver']==solver and r['method']=='baseline'}
                for method in ['position_only','full_pipeline']:
                    right={(r['seed'],r['pair']):r for r in p2 if where(r) and r['solver']==solver and r['method']==method};assert left.keys()==right.keys()
                    common=[(a,right[k]) for k,a in left.items() if a['success']==right[k]['success']=='1']
                    full_paired.append(dict(group=group,method=method,solver=solver,cases=len(left),gained=sum(a['success']=='0' and right[k]['success']=='1' for k,a in left.items()),lost=sum(a['success']=='1' and right[k]['success']=='0' for k,a in left.items()),common_success=len(common),mean_ep_change_m=stats.mean(float(b['e_p_m'])-float(a['e_p_m']) for a,b in common) if common else ''))
        old.write_csv(OUT/'full_pipeline_summary.csv',full_summary);old.write_csv(OUT/'full_pipeline_paired.csv',full_paired)
        extractions=[]
        for scan in range(1,6):
            for seed in old.SEEDS:
                path=RUNS/f'full_extraction/Plot2/S{scan}/step_seed{seed}_trees.csv';trees=old.read_csv(path)
                assert all(math.isfinite(float(r[k])) for r in trees for k in ['x','y','z'])
                extractions.append(dict(scan=scan,seed=seed,old_trees=len(old.read_csv(old.RUNS/f'extraction/Plot2/S{scan}/step_seed{seed}_trees.csv')),new_trees=len(trees),new_tree_sha256=old.sha(path)))
        old.write_csv(OUT/'full_extraction_summary.csv',extractions)
    audits=[]
    for path in sorted((RUNS/'positions').glob('*/*/*/*_trees.csv.audit.csv')):
        records=old.read_csv(path)
        if not records:continue
        audits.append(dict(plot=path.parents[2].name,scan=path.parents[1].name,profile=path.parent.name,seed=int(path.name.split('_')[1][4:]),input_trees=len(records),intersection_failed=sum(r['intersection_ok']=='0' for r in records),median_abs_height_change_m=stats.median(abs(float(r['new_z'])-float(r['old_z'])) for r in records),median_old_dtm_residual_m=stats.median(abs(float(r['old_dtm_residual_m'])) for r in records if math.isfinite(float(r['old_dtm_residual_m'])))))
    old.write_csv(OUT/'position_audit.csv',audits)
    notes=[['這次改什麼','先固定上一輪樹幹軸線，只修改DTM並重算地面交點；再以選定設定在Plot2重跑完整提取與配準，兩種實驗分開列。'],['原設定','地形網格0.5m、平面擬合半徑2格、最低支持點；匹配10cm＋三維補救＋配準RANSAC不變。'],['候選設定','smooth：擬合半徑2→6格、填補上限5→6格；coarse：網格0.5→1m、擬合半徑2→3格、填補上限5→6格；quantile1：最低支持點改成第1百分位。'],['選定方法',NAMES[profile]],['怎麼選','只用Plot2、seed31–35的50案例選定位階段成功數最高者；平手優先smooth、coarse、quantile1；沒有改善就保留原設定。'],['驗證怎麼分','Plot2 seed36–40另列，仍共用站點；Plot1、3、4共300案例為不同站點的定位階段驗證。所有樣區的舊基準結果先前已看過。'],['完整流程範圍','完整提取只追測Plot2五站、seed31–40。這是同一資料的後續測試，不能將其他樣區的定位結果當作完整流程驗證。'],['成功標準','旋轉<2°且平移<0.5m；平均點雲誤差取source全部點，只對成功案例彙總。'],['怎麼看退步','新增成功越多、退步失敗越少越好；誤差差值是新減舊，僅比較兩法皆成功案例，負數較好。'],['限制','定位階段重用既有樹幹；完整流程會改變入選點、法向量及擬合樹幹。兩者可能得到不同結果。'],['真值使用','DTM工具不接收GT；GT只用於開發組選設定與事後評分，沒有用於單次推論。'],['與論文關係','對照為上一輪10cm＋補救，並非原作者程式。成功率相同也不能宣稱嚴格超越論文。'],['原資料','上一輪結果完整保留；沒有存取或修改任何新的Google雲端檔案。']]
    notes.insert(2,['平滑範圍變多少','每格0.5m；擬合半徑2→6格，約1→3m；填補上限5→6格。'])
    notes.insert(3,['為何影響旋轉','各站的DTM高度不一致時，三維SVD可能用傾斜補償高度差；加大平滑範圍可減少部分這類差異。'])
    labels={'all':'全部（含開發組）','Plot2_all':'Plot2全部','development':'Plot2開發組','within_plot_validation':'Plot2其餘seed','cross_plot_validation':'Plot1／3／4驗證'}
    def label(s):return labels.get(s,s)
    def main_sheets(book):
        method_names={'baseline':'原設定','position_only':'只改地面交點','full_pipeline':'新地形＋完整重提取'}
        if full_summary:
            ws=table(book,'完整流程Plot2',['範圍','方法','成功數','案例數','成功率','成功案例平均誤差cm'],[[label(r['group']),method_names[r['method']],r['success'],r['cases'],r['rate'],100*r['mean_success_ep_m'] if r['mean_success_ep_m']!='' else None] for r in full_summary if r['solver']=='svd+ransac'],[30,40,18,18,20,32])
            for row in ws.iter_rows(min_row=2):row[4].number_format='0.0%'
            table(book,'完整流程改善',['範圍','新方法','新增成功','退步失敗','兩法皆成功','平均誤差變化cm（負數較好）'],[[label(r['group']),method_names[r['method']],r['gained'],r['lost'],r['common_success'],100*r['mean_ep_change_m'] if r['mean_ep_change_m']!='' else None] for r in full_paired if r['solver']=='svd+ransac'],[30,40,20,20,24,38])
        primary=[r for r in summary if r['solver']=='svd+ransac']
        ws=table(book,'定位階段結果',['範圍','方法','成功數','案例數','成功率','成功案例平均誤差cm'],[[label(r['group']),'原設定' if r['method']=='baseline' else NAMES[profile],r['success'],r['cases'],r['rate'],100*r['mean_success_ep_m'] if r['mean_success_ep_m']!='' else None] for r in primary],[30,35,18,18,20,32])
        for row in ws.iter_rows(min_row=2):row[4].number_format='0.0%'
        table(book,'改了什麼',['項目','說明'],notes,[30,130])
        table(book,'改善與退步',['範圍','新增成功','退步失敗','兩法皆成功','平均誤差變化cm（負數較好）'],[[label(r['group']),r['gained'],r['lost'],r['common_success'],100*r['mean_ep_change_m'] if r['mean_ep_change_m']!='' else None] for r in paired if r['solver']=='svd+ransac'],[32,20,20,24,38])
        dev=old.read_csv(OUT/'development_summary.csv');table(book,'候選怎麼選',['候選設定','解變換方式','開發成功數','案例數'],[[NAMES[r['profile']],r['solver'],int(r['success']),int(r['cases'])] for r in dev],[35,25,22,20])
    simple=Workbook();simple.remove(simple.active);main_sheets(simple);simple.save(OUT/'comparison_simple.xlsx')
    full=Workbook();full.remove(full.active);main_sheets(full)
    table(full,'欄位說明',['欄位','白話說明'],[
        ['method / profile','baseline為原設定；selected為只重算交點；full_pipeline為完整重提取。smooth為0.5m網格、擬合半徑6格。'],
        ['plot / pair / seed','樣區、source到target配對、距離抽樣亂數；比較方法時三欄必須相同。'],
        ['stratum','development為Plot2 seed31–35；within_plot_validation為Plot2 seed36–40；cross_plot_validation為其他三樣區的定位階段。'],
        ['success / transform_ok','通過旋轉及平移門檻／有求出變換，1表示是；有變換仍可能配錯。'],
        ['rot_err_deg / trans_err_m','相對GT的旋轉誤差（度）及平移誤差（m），越小越好。'],
        ['e_p_m','source全部原始點在估計變換與GT下的位置差平均，單位m；開發候選未計算者留空。'],
        ['n_correct_matches','GT事後評估水平XY距離<30cm的配對數，不保證樹位高度正確。'],
        ['time_ms','原生匹配與解變換時間，不含完整提取、讀檔或全點誤差；不與論文總時間直接比較。'],
        ['gained / lost','同案例原失敗變成功／原成功變失敗。'],
        ['mean_ep_change_m','只對兩法皆成功案例，新減舊的平均誤差差值；負值較好。'],
        ['median_old_dtm_residual_m','原樹位z到該profile地形面的垂直差絕對值中位數；候選列並非到原地形面的差。'],
        ['樹位數','通過本次提取的候選數，沒有人工標註可確認實際樹木數或召回率。']
    ],[42,140])
    def value(x):
        if x=='':return None
        try:f=float(x);return f if math.isfinite(f) else None
        except (ValueError,TypeError):return x
    details=[('兩種solver摘要',summary),('逐次配準結果',rows),('地形交點變動',audits),('開發逐次結果',old.read_csv(RUNS/'development_rows.csv'))]
    if full_rows:details.extend([('完整流程逐次',full_rows),('完整流程兩solver',full_summary),('完整提取樹位數',extractions)])
    diagnostic=OUT/'fixed_pair_height_diagnostic.csv'
    if diagnostic.exists():details.append(('固定配對高度診斷',old.read_csv(diagnostic)))
    for title,records in details:
        h=list(records[0]);table(full,title,h,[[value(r.get(k,'')) for k in h] for r in records],[30]*len(h))
    full.save(OUT/'comparison_full.xlsx')
    files=['comparison_simple.xlsx','comparison_full.xlsx','registration_rows.csv','summary.csv','paired_comparison.csv','position_audit.csv']
    files.extend(name for name in ['full_pipeline_rows.csv','full_pipeline_summary.csv','full_pipeline_paired.csv','full_extraction_summary.csv','fixed_pair_height_diagnostic.csv','input_axes_audit.json','baseline_position_audit.json','gt_tilt_diagnostic.csv'] if (OUT/name).exists())
    (OUT/'validation.json').write_text(json.dumps(dict(final_rows=len(rows),full_pipeline_rows=len(full_rows),development_rows=400,selected_profile=profile,original_native_binaries_and_defaults_unchanged=True,analysis_script_sha256=old.sha(Path(__file__)),files={name:old.sha(OUT/name) for name in files}),indent=2)+'\n')
    print(json.dumps([r for r in summary if r['solver']=='svd+ransac'],indent=2),flush=True)

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('phase',choices=['freeze','develop','select','validate','audit','full','report']);args=parser.parse_args();globals()[args.phase]()
