#!/usr/bin/env python3
"""Publishable accuracy-complete workbook preview; never marks the round done."""
import hashlib
import fcntl
import json
from pathlib import Path
import shutil
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import load_workbook
from openpyxl.styles import Alignment,Font
import unified_report as r
import unified_study as s

def main():
    original=s.OUT;dest=original/'progress';dest.mkdir(exist_ok=True)
    rows=r.collect();summary,datasets,seeds,paired,goals=r.summaries(rows)
    station,_,resources=r.costs()
    cold=[json.loads(p.read_text()) for p in sorted((s.RUNS/'cold_benchmark').glob('*/*/complete.json'))]
    station_count=len(list((s.RUNS/'feature_benchmark').glob('*/*/complete.json')))
    copies=['versions.json','method_v3_frozen.json','eth_official10_summary.csv','registration_rows.csv','scene_summary.csv','dataset_summary.csv','ablation.csv','paper_reference_comparison.csv']
    for name in copies:shutil.copyfile(original/name,dest/name)
    r.OUT=dest
    r.make_books(rows,summary,datasets,seeds,paired,goals,station,cold,resources)
    note=f'階段版：精度4950筆完整；成本尚未完成（冷啟動{len(cold)}/12、單站成本{station_count}/26）。正式版將在背景量測及核驗完成後另外發布。'
    books=[]
    for kind in ['simple','full']:
        old=dest/f'comparison_{kind}.xlsx';book=load_workbook(old);sheet=book['先看結論']
        sheet.insert_rows(2);sheet.cell(2,1,'階段版狀態');sheet.cell(2,2,note)
        for c in sheet[2]:c.font=Font(name='Arial',size=10,bold=True);c.alignment=Alignment(wrap_text=True,vertical='top')
        sheet.row_dimensions[2].height=48;sheet.auto_filter.ref=sheet.dimensions
        for row in sheet.iter_rows(min_row=3):
            if row[0].value in ['成本解讀','改善的代價']:row[1].value='成本未完成；此階段版只列已取得的紀錄，未量測項目不當作0。';row[1].alignment=Alignment(wrap_text=True,vertical='top')
            if row[0].value=='研究完成狀態':row[1].value='本輪仍在計算成本；作者原作baseline與未見森林驗證也尚未完成。'
        target=dest/f'comparison_{kind}_progress.xlsx';book.save(target);old.unlink()
        verified=load_workbook(target)
        assert not any(c.data_type in ['e','f'] for w in verified for row in w for c in row)
        assert all(w.freeze_panes=='A2' and w.auto_filter.ref for w in verified)
        if kind=='full':assert verified['全部逐次結果'].max_row==4951
        books.append(dict(file=target.name,bytes=target.stat().st_size,sha256=s.sha(target)))
    for name in copies+['validation.json']:(dest/name).unlink()
    metadata=dict(provisional=True,round_complete=False,accuracy_rows=4950,primary_trials=550,primary_success=550,scenes=5,datasets=2,cold_complete=len(cold),cold_planned=12,station_costs_complete=station_count,station_costs_planned=26,generated_at=datetime.now(ZoneInfo('Asia/Taipei')).isoformat(timespec='seconds'),tests_passed=93,workbooks=books,source_sha256={name:s.sha(original/name) for name in copies if name!='versions.json'})
    (dest/'validation.json').write_text(json.dumps(metadata,ensure_ascii=False,indent=2)+'\n')
    (dest/'README.md').write_text('# 精度完整、成本待補的階段版\n\n'+note+'\n\n同一固定兩站方法在55配對×10seed全部通過專案門檻；五場景平均全來源點Eq8誤差均低於論文報告值至少10%。這是報告數字比較，作者未釋出程式，未完成嚴格原作對照；兩來源均已參與開發，未證未見森林一般化。\n\n正式版的comparison_full.xlsx及comparison_simple.xlsx將由finalize_unified.py於12/12冷啟動及26/26單站成本齊全、控制器結束後生成並核驗。此目錄只保留階段版，不冒稱整輪已完成。\n')
    print(json.dumps(metadata,ensure_ascii=False,indent=2),flush=True)

if __name__=='__main__':
    with (s.RUNS/'progress-report.lock').open('a') as lock:
        fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
        main()
