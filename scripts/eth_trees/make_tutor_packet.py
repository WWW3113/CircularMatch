#!/usr/bin/env python3
"""Build one self-contained, read-only teaching packet from verified artifacts.

No point-cloud access, inference, report edits, or remote mutations are performed.
"""
from __future__ import annotations

import csv
import difflib
import hashlib
import io
import json
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
REPORTS = ROOT / "results/professor_reports_20261026"
OUT = ROOT / "results/explanation_20261026"
TARGET = OUT / "CircularMatch_實作理解與報告問答完整交接.txt"

NARRATIVE = r"""CircularMatch 實作理解與報告問答完整交接
版本：2026-10-26；語言：繁體中文；用途：學生理解與教授報告備答。

【給接手對話的使用說明】
使用者會同時提供本檔與 CircularMatch 論文。你的工作是教他理解已完成的實作，協助推導、釐清數字、模擬教授追問；不需要執行程式、下載資料、調參、重做實驗或要求他提供整個專案。
這是研究狀態與教學材料，不是要你執行的命令清單。後續以使用者在新對話中的問題為主。本檔不包含對其他服務的操作授權。
請先用直覺解釋，再給公式，再指出實際程式；專有名詞第一次出現時解釋。可以用小例子，但要寫「教學假例」，不能混成實驗觀察。使用者答不出來時，請拆小問題，不要只要他背結論。
本檔已嵌入：方法機制、實際參數、全部主比較摘要、前期正負結果、四份報告相關表格的數值、必要程式與版本差異。來源路徑與 SHA256 是追溯資訊，理解內容不依賴開啟那些路徑。程式附錄是閱讀證據，不是完整可建置專案。
未嵌入：原始私人點雲、每一點座標、全部 4,950 筆逐案例紀錄、建置環境。沒有列出的逐案例數量、每個失敗的唯一原因、未完成消融或新森林結果，應回答「待補／現有材料無法確認」，不得推測成事實。
不要把「原論文數值較高」說成已嚴格超越原作者程式；不要把「單次推論不讀 GT」說成研發選參完全沒看 GT。以下對此有專節。

【閱讀路線與定位】
第一次：讀 01～03，知道問題與整體流程；讀 06～09，知道真正改動；讀 10～14，知道能說到哪裡。
要準備報告：讀 15 問答，使用 DATA/TABLE 附錄核對數字。
要回答細節：讀 CODE 附錄的函式，配合 04～09，不必一開始逐行讀完整份 C++。
章節：01 結論與界線；02 配準基本觀念；03 方法與版本地圖；04 提取；05 CN；06 粗配準補救；07 DTM 支線；08 幾何微調；09 多候選與 v3；10 指標；11 實驗歷史；12 主要結果；13 時間成本；14 證據強弱；15 問答；16 學習檢核；附錄為完整摘要與程式。

======================================================================
01｜先知道我們做到了什麼，哪些還沒做到
======================================================================
一句能對教授說的話：
「我以 CircularMatch 的樹位 CN matching 為粗配準骨架，在本地參考實作上調整距離容差、加入弱配對共識補救、多尺度 6DOF 幾何微調，以及多候選的結構驗證。固定同一套最終設定，在 Tongji-trees 四個 Plot 與 ETH-tree 的既定兩站配對上均通過本專案成功門檻，平均 e_p 數值低於論文 Table I；但原作者可執行 baseline、完全獨立新森林與部分消融仍缺，不能宣稱已完成嚴格超越原作或未見森林一般化。」

主實驗兩個 dataset：Tongji-trees 是一組，Plot1～Plot4 是它的四個場景；ETH-tree 是另一組。不是五個獨立 dataset。
Tongji 每 Plot 五站，10 個唯一配對；ETH 六站，15 個唯一配對。合計 26 個掃描站、55 個配對。每配對 seed31～40 重跑十次，因此每方法 550 次；九方法共 4,950 筆。
最終 multi_v3_refined：Tongji 各 Plot 100/100，ETH 150/150。這是 55 個配對各十個隨機提取／擬合種子，並非 550 片不同森林，也不是成功機率永遠等於 1。
每次推論只用當前兩站。六站是 ETH dataset 的大小，並不是配準時要六站一起幫忙。過去另有六站 graph_mst 支線，資訊條件不同，未拿來充當兩站主結果。
最終方法沒有使用 DTM smooth 支線；原 step 距離保留法、原 DTM、提取 verticality>0.9 仍保留。不可把所有嘗試都畫成最終方法的組成。

最需要講清楚的 GT 限制：
GT = Ground Truth／資料提供的參考姿態變換。推論時的候選產生、候選評分、輸出選擇、ICP 接受判斷不讀 GT，沒有用 GT 為當前配對挑最佳輸出。
然而研發期間確實用 GT 成功率與誤差比較過設定：ETH matching/recovery 的設定選擇、Plot2 DTM profile 選擇，以及失敗案例診斷。DTM smooth 特別是按 Plot2 seed31～35 的 GT 成功數選出；這不是完全不看真值的研發。
因此只能稱「GT-free inference」，不能稱「整個研發過程完全無 GT 選參」「完全未見資料盲測」。不同 seed 也不會消除共享掃描站的資料接觸。
早期研究需求有「禁止使用真值決定參數、選候選或挑輸出」這個較嚴格說法。若它要求連離線開發選參都不能看 GT，現有紀錄不滿足這個強版本限制；不能靠重新命名把它說成已滿足。若只限定單次推論，程式的資料輸入邊界支持此限定。這個差別要主動承認。

======================================================================
02｜先懂輸入、輸出與「為何一點旋轉會差很多」
======================================================================
TLS 是 Terrestrial Laser Scanning，地面雷射掃描。同一棵樹在不同站掃描，點座標原點與朝向不同，又有遮擋、不同可見面與雜點，所以不能直接把兩份 XYZ 疊起來。
輸入 source A 與 target B；輸出 T_A_to_B=(R,t)：q = R p + t。
p 是 A 座標的三維點，q 在 B 座標；R 是 3×3 正交、det=+1 的旋轉，t 是三維平移，距離單位 m。齊次 4×4 寫法右下角是 1。反向 T_B_to_A=(R^T,-R^T t)，不能只把 t 加負號。
6DOF 是三個旋轉自由度（roll、pitch、yaw）與三個平移（x、y、z）。我們最後仍估完整 6DOF；以「大致直立」產生 yaw 初值，不代表把最終 roll/pitch 鎖成零。

教學假例：一個點距離旋轉中心 20m，旋轉誤差 1°≈0.01745rad，單此項可帶來約 20×0.01745=0.349m 位移。這只是小角度直覺，並非本實驗某案例實測。故平移誤差很小，不保證遠處點的 e_p 也很小。
Tree position 是圓柱軸線與估計地面的交點，不是任意取一個樹幹點，更不是每棵樹的真值中心。
Surface normal 是表面局部法向量，不是樹幹軸線：直立圓柱的軸近 z，側面的 normal 近水平；水平地面的 normal 近 z。後面兩種垂直篩選都利用這個差別。
Normal 正負號可能不一致，故幾何對應用 abs(dot)；n 與 -n 表示同一局部平面的方向，不能因符號相反就判定不一致。

流程圖（箭頭表示資料依賴，不表示每步都用相同點集合）：
每站原始點雲
  ├─ DTM／高度／1cm voxel／法向／距離保留／分群／圓柱 → 樹位集合
  │    → CN descriptor → matching → 弱配對補救 → 6DOF 粗估計
  │                                      └─ 弱時提出多個幾何初值
  └─ 0.5/0.2/0.1m 原始全表面 cache ─────────→ 多尺度微調與驗證 → 輸出 T
輸出 T 與 benchmark GT ─→ 事後 e_p/e_R/e_t/成功率；GT 不沿上面箭頭進入推論。
但開發人員讀事後結果選過設定，這是離線開發層，不能用上述流程圖掩蓋。

======================================================================
03｜原論文、本地參考與版本關係
======================================================================
論文：Sun 等，CircularMatch；DOI 10.1109/LGRS.2025.3650098。頁碼以下指上傳 PDF 實體第 1～5 頁：p2 方法，p3 Algorithm1，p4 Eq.(8)/Table I/II，p5 Table III。
論文明列：提取 voxel 0.01m、法向半徑 0.10m、verticality>0.9、離 DTM 高度 0～3m、按水平距離 5/10m 分界保留 0.5/0.75/1；CN N183、w10、Thscore40、距離容差 0.05m。
未在論文明確給足的項目：精確 DTM 參數、分群距離、剛體求解器、成功判定門檻、逐場景配對分母、e_R/e_t 報告值。這些缺口不能自行補上後稱為作者設定。
本地 5cm reference 使用自己實作的提取與 CN，剛體求解用 Kabsch SVD；還有本地 DTM 與圓柱過濾。它是 paper-inspired reference，並非已證明與原作者程式逐項等價。
2026-10-25 13:17:46 UTC 的作者程式庫稽核：MercerSU/CircularMatch 當時只有 README，HEAD 1e9623fd8893633770b689baf8bda66f7f36f307，無可重跑實作。這是當時查核，不宣稱未來永不提供。原作者未修改程式在本環境重跑的 baseline：待補。

九版本對照（所有主比較保持相同提取樹位與既定配對）：
reference5_raw：0.05m CN、SVD、無新增 recovery、無 surface refinement。
reference5_refined：上面 reference 加舊全表面驗證的 refinement。
eth10_raw：0.10m CN＋註冊 RANSAC＋weak evidence recovery，沒有 surface refinement。
eth10_refined：eth10_raw 加相同的舊 refinement。
multi_refined：v1 搜尋較小的多初值＋舊 refinement/驗證。
multi_v2_refined：v2 擴大候選＋yaw seeds＋全兩站樹位支持＋保留初值，仍是舊驗證。
reference5_refined_v3：reference 加 v3 相同 refinement backend，作公平控制。
eth10_refined_v3：eth10_raw 加 v3 backend，不增加多初值，作公平控制。
multi_v3_refined：v2 候選引擎＋v3 垂直結構／收緊 gate 的驗證；最終主方法。
eth10 這個名稱表示「最初在 ETH 開發後凍結的一組粗配準設定」，不是只可用 ETH，也不表示單純只改 5→10cm。真正隔離 threshold 與 solver 的 2×2 在早期 seed11～20，不能拿主表整組差異替代它。

設定檔注意：config/default.ini 是模板，不是統一主實驗的完整有效設定。模板 cluster.tolerance=0.1、seed_base=1；實驗 override 實際用分群 0.2m、step、scanner origin=(0,0,0)、seed31～40。多候選與 geometry 的不少常數寫在 C++，unified_config.json 是實驗快照，不能假設改 JSON 就會改 native 行為。
模板「synthetic only」與舊文件「no ICP」屬早期狀態／階段邊界，不代表最終版沒做森林試驗或沒加 refinement。本文與附錄列明是哪個版本，避免拿歷史註解解讀新結果。

======================================================================
04｜提取樹位：每一步在消除什麼問題
======================================================================
4.1 座標與 DTM
prepareCloud 先以 scanner position 減出局部座標，避免大世界座標轉 float 損失精度，輸出樹位再加回。統一實驗各站指定 scanner=(0,0,0)，假設檔案是以站點為原點、m 為單位且 z 向上。這是輸入假設，並非程式已估出每個掃描器的真實地理姿態。
DTM（Digital Terrain Model）估 z=f(x,y)，目的是依離地高度截出下段樹幹，並定義圓柱軸線的落地點。地形不是一個全場景固定高度。
原 DTM cell=0.5m；每格至少3點；supported_lowest 找低處且上方0.05m內有至少3點支持，percentile5 作後備；斜率過濾 max_slope=0.5m/m、tol=0.3m、radius=5m；高於鄰近平面0.5m的樣本可剔除；局部平面 fit_radius=2cells、fill_max_radius=5cells，QR rank 要足。
這些為本地實作選擇。較大窗口能減少孤立低樣本造成的高度起伏，也可能跨過真實坡折而偏差；它不是保證地面越平越準。
高度條件為 dz>=-0.05m 且 dz<=3m；原論文敘述0～3m，本地額外給5cm負容差。DTM 高度錯會改變保留的點，也會改變樹位 z。

4.2 1cm voxel 與 normals
提取 voxel 每1cm格保留「最靠近格內質心的原始點」，並非以質心產生一個新點；穩定 tie 規則利於重現。
法向量取0.10m半徑 PCA；至少5個鄰點，octree resolution0.05m。局部 covariance 最小特徵值的 eigenvector 是 normal；curvature=最小特徵值／特徵值總和。
提取 verticality = 1-|normal.z| >0.9，即 |normal.z|<0.1。保留近垂直表面，如樹幹側面；不是保留 normal.z>0.9，也不是「樹幹軸線水平」。

4.3 依距離的 downsampling，主方法到底用哪個
原始遠距掃描通常較稀，所以主方法用論文 step 保留機率：d<5m 保留0.5；5<=d<10m 保留0.75；d>=10m保留1。d 是相對當前站原點的 XY 水平距離，非離另一站距離、非離 GT 世界原點；use_3d_distance=false。
每點由 seed 與量化座標的 hash 產生 u，u<機率才留；hash_quantum=0.0001m。這是 Bernoulli 保留，有限點集不一定精確留下50%；CRN（common random numbers）使同一點同 seed 可重用同 u 比較保留規則。
早期試過 linear_a、linear_mid、physical，但主配準報告未證明它們在全場景更好，所以最終仍 step。不能回答成最終採用了 physical 或 voxel取代距離保留。

4.4 分群與第一種 RANSAC：擬合圓柱
Euclidean cluster tolerance 實際0.20m，最少30點、上限10M點。太小可能切碎一棵樹，太大可能混不同物體；未有完整此參數消融。
圓柱 RANSAC 在一個點群內找「點與法向量符合圓柱」的模型；max_iterations1000、probability0.99、distance0.02m、normal_weight0.1；幾何內點再做最多2次 LM refit。
另有本地過濾：內點比例>=0.5、半徑0.03～1m、軸傾角<=20°、周向覆蓋>=100°（5° bins）、>=75%點的法向符合模型且角差<=30°。不是論文明確逐項指定。
作用：排掉牆、枝葉、只看見極小圓弧等不可靠圓柱，但也可能漏掉細樹、斜樹與遮擋嚴重樹幹。較寬 verticality 或較多 RANSAC 迭代並非一定更好，負面實驗見後文。
active seed_mode=follow：由實驗 seed 與穩定群集 hash 派生圓柱 RNG；不因模板 fixed_seed12345 就誤認全都固定12345。註冊 RANSAC 的 seed=1 是另一回事。

4.5 軸線與地面交點
圓柱提供軸 a+τv。求 a_z+τv_z=f(a_x+τv_x,a_y+τv_y)；標準化且軸朝 z 正向。固定點迭代最多50次、殘差容差0.001m；必要時在軸上±5m、step0.1m搜尋 bracket 再二分。無可靠交點則不輸出那棵樹。
這裡的1mm是數值求根容差，不代表 DTM 對真實地面有1mm準確度。樹位 CSV 四位小數是0.1mm表示精度，也不是物理精度。

======================================================================
05｜CN descriptor 與 matching：不要把不同 threshold 混在一起
======================================================================
CN 用「一棵樹周圍其他樹的空間排列」辨認這棵樹，不用樹皮顏色。樹位來源是上一章的估計交點。
本地 CN 使用 XY 距離與 XY 角度，z 留給後面的3D剛體求解。需中心之外至少3個不同 XY 樹位，才有有效 descriptor；XY重合<1e-9不提供角度。
N=183 =3個最近鄰 gatekeeper 距離＋180個扇區維度；每扇區2°。不是183棵樹、不是183°，也不是每站只取183點。
分別以中心→第1、2、3最近鄰為0角度方向，各建180扇區，每格記最近樹的距離與索引。概念矩陣為3列distance＋3列index，即6×180；再按 encode1優先、其次2、其次3逐格填成一條183維descriptor。此規則是優先填入，不是將三個distance求平均。
空格 val=0、idx=-1；matching必須先判idx有效，空格對空格不算匹配。0不代表真的另一棵樹就在0m。
沒有另設固定 neighborhood radius，當站其餘抽出的樹位都可能參與扇區。相對角度設計期望對 yaw 有穩定性，但漏樹、最近鄰次序改變、2°量化、非直立旋轉會破壞完全一致；不能說對任意3D旋轉完全不變。

matching 的次序：
① Gatekeeper：三個對應最近鄰距離，只要任一對差<threshold就通過，三對全不符才拒絕。不是全部三對都要符。
② DC：只在180個扇區，兩邊都非空且距離差<threshold，計1分。前三gatekeeper不在DC加分。每來源先找DC最高target，再每target留最佳來源；DC>=min_dc3才保留。
③ FM：累計扇區證據反覆支持同一對樹的票數；MatchScore=w×FM+DC，w=10。DC是單個descriptor的一致格數；FM是跨候選支持，兩者不是同一概念。
④ 分數>Thscore40時取候選中心對。最佳分數<40則用它支持的樹集合做三角形fallback；三邊XY距離逐一<threshold，依排列比對、頂點投票，最後貪婪一對一。這是原先matching內的fallback，不是後來新增recovery。
本地對論文歧義的決定：DC採>=3；score剛好40不接受，因程式是>40，且只有最佳<40才fallback。論文文字與偽碼的邊界不完全清晰，不能聲稱唯一正確復刻。必要時直接指附錄 matchCn。

5→10cm 實際改變：同一個 threshold 由0.05到0.10m，會放寬gatekeeper、DC的距離一致性與triangle的邊長比較，因而改變候選與證據。沒有把 Thscore40改成10，也不是 ICP 對應半徑改成10cm。
教學假例：兩站對同一對樹估出距離5.00與5.07m，差0.07，5cm不通過、10cm通過；若兩棵錯樹也碰巧差小於10cm，一樣可能被放進來。容差擴大同時增加召回與誤配，後面需要幾何共識；不能只說「更寬更準」。

======================================================================
06｜粗配準：SVD、第二種 RANSAC，以及新增弱配對補救
======================================================================
Kabsch：給已配對的3D樹位 (a_i,b_i)，求 R,t 最小化 Σ||Ra_i+t-b_i||²。減去兩邊質心、做cross-covariance SVD、修正det避免鏡射，最後 t=mean(b)-R mean(a)。至少3個不共線對應；退化會回傳無變換。
註冊 RANSAC：在「已matching的樹位對」中抽3對求Kabsch，3D距離<0.30m算內點；最多1000次抽樣嘗試，重複索引跳過，seed固定1；取內點最多的一組，再用其內點重估。這不是在原始樹幹點擬合圓柱；兩種RANSAC輸入與門檻不同。
SVD會把誤配一起最小平方，RANSAC希望先找主要一致群；但正確內點少、位置退化或高度系統偏差仍會出錯。RANSAC1000是最多抽樣嘗試，不代表一定有1000個有效模型。
XY descriptor匹配看似正確，但DTM的z偏差會進入3D Kabsch，可能被solver解讀為roll/pitch偏差。水平匹配不保證3D姿態正確，也不能因此擅自把z全刪掉並聲稱等價。

新增 evidence recovery：只在原CN輸出對數<6時觸發；取前3個ranked descriptor候選與中心／sector證據，票數排序後最多60條對應邊。枚舉由3條邊組成的相容三元組：src/tgt不得重複，XY邊長差<2×CNthreshold（10cm時是20cm）。對三元組解完整6DOF，再用同一證據池的3D<0.30m與一對一共識評分。
先比支持數，並列比較median殘差、再比總殘差；支持>=6且多於原對數才採用，否則保留原matching。這是GT-free當次幾何判斷，不是用GT看哪組e_p比較好。
為何6？3是數學最少，6是增加冗餘的啟動／採用啟發式；不是「不到6棵樹完全不能求解」。recovery的輸出是更多可信樹位對，之後仍走同一RANSAC。
它和後面的surface multistart不同：recovery比較同一CN證據池的樹位共識；multistart會對多個pose做點雲微調、用兩站全部樹位及表面驗證。

已完成證據（ETH，其他條件固定）：
seed11～20的matching×solver 2×2：5cm+SVD118/150，5cm+RANSAC123/150，10cm+SVD135/150，10cm+RANSAC142/150。
固定SVD只改容差新增17次（+11.333pp）；固定RANSAC新增19次（+12.667pp）。固定5cm換RANSAC新增5次，固定10cm新增7次；不假設兩者效果必可線性相加。
選出weak recovery後，seed21～30固定10cm/RANSAC：142→150/150，新增8次、無退步；啟動／採用31次，不代表31次原本都失敗。
共同成功142次 mean e_p 0.08618106527669014→0.07515705036880281m；舊 paired_validation.json 的 common_ep欄其實取median，不可把它當mean。median約0.067380268495→0.06194308208m。
到Tongji的第一次固定設定遷移只有329/400，沒有全通過；所以「ETH粗配準100%」從來不代表「所有新dataset粗配準都100%」。

======================================================================
07｜DTM smooth 支線：有局部改善，但沒放進最終 v3
======================================================================
問題假說：局部地面模型造成樹位高度噪聲，讓XY較對但3D姿態受影響。因此比較 baseline cell0.5m/fit2/fill5 與 smooth cell0.5m/fit6/fill6。radius單位是格，不是m；fit2為方形半寬2格（約±1m範圍），fit6約±3m。
其他profile：coarse cell1.0m/fit3/fill6；quantile1 cell0.5/fit2/fill5、地面改第1百分位。
Plot2 seed31～35共50次開發：baseline18、smooth24、coarse21、quantile1 15次成功；按GT成功數選smooth。這是離線GT選參的明確例子。seed36～40雖是後續種子，仍是同五站，不能稱新森林。

兩種不同實驗務必分開：
A 固定舊圓柱軸，只用新DTM重新求交點。Tongji主要RANSAC分支總329→344/400；Plot2 34→48/100、Plot4 95→96/100，Plot1/3仍100。
Plot2配對分析新增21、退步7、共同成功27；總數只多14，但改變的不只14個案例。Plot3成功限定mean e_p 4.90176→5.89067cm，反而差約0.989cm。不能說四Plot都變準。
B 完整重提取，DTM同時影響height filter、cylinder population與交點。只有Plot2做完，34→48/100；新增23、退步9、共同成功25。A/B最終都是48，並不代表成功的是同48次。其餘三Plot全流程與ETH同組比較：待補。
若宣稱「新DTM全流程Tongji86%」，就是把A的定位測試錯當B。最終v3沒有套smooth，不能把550/550歸功這支線。

機制線索：固定原CN樹對，用GT檢查每案例高度殘差中位數，再對100案例取median，0.4540364→0.3781113m，79/100下降。這是GT高度診斷，不是e_p，也不是與人工標地面的DTM真實誤差。
baseline定位稽核200份CSV、22,375個樹位座標完全一致，最大軸／DTM殘差約0.00209m，支持重新求交點沒有巨大writer錯誤；不證明地面準確至2mm。
合理推測是較平滑的height估計減少部分不一致；但可能抹去真實坡折，亦會改變提取群集，所以有退步。缺少獨立地面GT與逐項消融，不能說原因已全部證實。

還有不能漏報的Plot4：定位測試成功95→96，但finite mean e_p為109.373→138.312cm，有限N也從99變100。新輸出與災難失敗會影響全有限平均；成功限定mean較低不等於全體誤差較低。因有限案例集合改變，不能把兩個mean直接解讀成同100案例的paired下降率，表格將該率留待補。
原始碼額外細節，與「做過實驗」分開：fit loop從fit_radius跑到max(fit_radius,fill_max_radius)，遇足夠rank的fit就停。smooth的fit6已令上限至少6，因此在這段loop裡fill5→6可能是功能冗餘；不是6比5獨立改善的證據。報告仍應說測試的是整組profile，因沒實際跑fill欄單獨消融，勿發明其效果。

======================================================================
08｜多尺度幾何 refinement：為何能更準，也為何救不了所有初值
======================================================================
8.1 第二條點雲路徑
粗配準只用稀疏樹位。refinement回到當前兩站原始全表面，沒有只用0～3m樹幹，也不讀其他站。
每站各建h=0.50/0.20/0.10m三層voxel centroid；這裡是真質心，與提取1cm保留原點不同。每層PCA半徑2.5h、至少6鄰點、curvature<=0.1。幾何cache用float儲存位置與normal，讀回normal正規化。
voxel lexicographic順序後，以floor(i*N/k)決定性均勻索引抽樣：optimizer每層最多20,000來源點；verification各方向最多8,000。是計算上限，不是評估e_p只抽8,000點。cache可跨seed重用，成本需另列。

8.2 一輪局部最小化做什麼
把source p用目前T變到target座標，找一個最近的target q，目標normal n；距離<3h且abs((R n_source)·n_target)>=0.7才考慮。若最近點normal不合，不再找第二近替代。
按距離排序，同一target最多保留一個source（較近者優先）。至少30對才求解；保留max(30,floor(0.7N))對。所以「70% trimming」帶有至少30對保護，不是所有N都精確70%。
point-to-plane residual r=n·(p-q)，單位m。它量測點到target局部切平面，與點對點歐氏距離、benchmark GT e_p不是一樣的量。
Huber δ=max(0.05h,2×upper_median(|r|))，權重 |r|<=δ時1，其餘δ/|r|。upper_median為nth_element的n/2元素，偶數時取較大中間值，不是兩個中間值平均。這是實作細節，不是已證明比標準median好。
採左乘小更新，δp≈δω×p+δt。J=[p×n,n]，H=ΣwJJ^T、b=-ΣwJr，解Hδx=b；這可由n·(δω×p)=(p×n)·δω推得。
若λ_min(H)<1e-8 λ_max(H)則當層停，保留目前姿態。若H退化，代表某些方向局部約束不足；H不退化也不代表全域姿態一定對。
一次更新旋轉限制0.15rad、平移限制h，超限對整個6維更新共同縮小。R_new=R_delta R_old，t_new=R_delta t_old+delta_t。不是僅把t相加而忘記左乘旋轉它。
最多30/25/20次，由粗到細；若角度<1e-6rad且平移<1e-5m則當層收斂。若表面點不足30也跳過該層。

8.3 驗證與回退
finest層比較起點與微調後的geometry quality=overlap/(0.03+NN_RMSE)。若after+1e-9<before，退回原初值；這個接受判斷不看GT。
這是幾何啟發式，並不保證e_p下降。若iteration計數>0但最後回退，不能說最後輸出一定做了有效更新；iterations可保留已嘗試的數目。
平面例子：只有水平地面，point-to-plane對XY平移與yaw約束很弱；加不同方向的垂直表面有助約束。森林中的重複排列或錯位仍可能產生局部低殘差，ICP不是全域定位保證。
三尺度、normal gate、一對一、trimming、Huber、退化檢查、接受回退是一組新增設計，沒有逐項ablation，不能說70%這一個數獨立造成全部進步。

主要結果eth10_raw→eth10_refined：
Plot1 100→100，finite mean e_p 17.423531→2.886601cm。
Plot2 34→100，50.900681→1.185416cm。
Plot3 100→100，4.901756→1.105659cm。
Plot4 95→96，109.373197→103.475574cm，兩邊有限N=99，仍1次缺變換。
ETH 150→150，8.189515→0.603814cm。
全體479→546/550；Plot4仍有4失敗，局部微調不是充分解。5cm reference加相同舊backend也423→507，故不能把所有增益叫作新CN優勢。

======================================================================
09｜多候選與 v3：多試幾個為何反而變差，如何修正驗證
======================================================================
9.1 v1到v2是整組搜尋改動
multistart僅在原始match.pairs.size()<6觸發，不是檢查recover之後數量；recovery找到6對以上也可能繼續觸發，因原CN仍弱。
v1：top5 descriptor candidates、60 evidence edges、最多10,000相容三元組、最多16個不同alternative poses、精化最多3。
v2/v3：top20、120 edges；加入upright雙點yaw初值，最多C(120,2)=7140種邊對可能（不代表每種有效、每種都refine）；也保留最多10,000相容3D三元組。最多256 alternative poses，粗篩後精化最多8個。
雙點yaw要求兩對樹位src/tgt不重複、source水平間距>=1m、邊長差<0.2m。依兩條XY向量角差給yaw，xy/t初值用兩對偏移平均，z用XY支持證據的upper median高度差。這依賴近直立掃描假設，但後面回到完整6DOF。
三元組走3D Kabsch。globalSupport把「當前兩站全部抽出樹位」變換後比XY距離<0.3m，greedy target一對一計數；global不是跨六站地圖，也不是真值樹對數。
pose差<5°且平移差<0.5m視同一群，依支持數／誤差保留、最多256；之後既有初值會插入。因此CSV proposal_count上限可為257，不應拿256去斷言紀錄錯誤（v1同理可17）。
v2保證既有粗初值進refinement beam；總beam最多8，所以是初值加至多7其他候選，不是8再加1。保留其參與不保證最後選回它，別的候選可拿更高分。
粗排序 q=bidirectional_overlap/(0.1+NN_RMSE)×log(1+support)；最終 q=overlap/(0.03+NN_RMSE)×log(1+globalSupport)。支持數與quality是啟發式，不是成功機率。
v1→v2同改了budget、yaw、globalSupport與保留初值等，不能單獨把差異歸因top-k5→20。

9.2 已完成的反例
Plot4 v1 96/100，v2反降78/100。所有finite mean e_p 1.72584→8.64762m，v2 median卻僅0.009033m；少數大錯會被median遮掉。
實例Plot4 S4-S2 seed31：v1正確的舊geometry overlap=0.1669778241；v2錯誤卻0.3377234446，e_p=45.9351400m。兩者同分數定義，因此支持「較大舊overlap不保證真值配準較好」。
大面積地面在錯移後仍可相似，是合理機制線索；現有材料沒有逐點人工地面標註或固定候選的方向消融，不能說每次失敗已證明全是地面造成。

9.3 v3只改verification，optimizer仍用全方向
每層選 |normal.z|<0.30 的點；source與target都至少30點才用此垂直表面子集，不足任一邊便兩邊退回全方向。分類使用每站本地z，未隨候選旋轉重新分類；大傾斜掃描是限制。
verification gate改min(h,0.2m)：h0.5/0.2/0.1時為0.2/0.2/0.1m；舊3h為1.5/0.6/0.3m。target最近鄰索引也用子集，並非全表面找到NN後才篩。
optimizer維持全方向表面、3h配對距離、原6DOF求解；地面並未從優化完全刪掉。這樣可能保留高度／姿態約束，同時讓辨識候選少受水平面歧義影響；「可能」不等於已逐項證明。
新的geometryScore用於候選粗排序、refinement前後接受／回退、最終選擇。不是只改最後一行sort。
directional overlap是抽樣來源中，NN在gate內且normal abs dot>0.7的比例；RMSE是這些合格對的NN距離均方根，無支持回1e9。雙向分數 overlap=sqrt(o_AtoB×o_BtoA)，rmse=(rmse_AtoB+rmse_BtoA)/2，support=min(雙向支持數)。這不是重疊真值標籤或已校準機率。
v3與v1/v2原始overlap不能直接按數字高低比較，因點集合與gate已改；比較輸出後共同定義的GT e_p與成功率才可。
重要：提取的verticality>0.9（|nz|<0.1）沒改；驗證的|nz|<0.3是另一段程式、另一點集合。不能說「把論文0.9改0.3」。

9.4 效果與trade-off
Plot4 v2→v3：78→100/100，新增22、退步0；finite mean 864.762201→3.392023cm。
但共同成功78例，平均e_p仍增加約0.733mm；Plot3 v2→v3也增加約0.127mm。穩健性提升不等於每一例精度更好。
同一v3backend公平控制：reference5_refined_v3=507/550，eth10_refined_v3=546/550，完整multi_v3_refined=550/550。多候選相對eth10+同backend額外救4次，全在Plot4；不能把全部127次提升歸給multistart。
缺少只改normal子集／只改gate的2×2實驗，故22次只能歸因驗證整組規則，不能分配成0.3門檻獨立貢獻多少。

======================================================================
10｜指標與統計：報告最容易被問倒的部分
======================================================================
e_R = acos(clamp((tr(R_GT^T R_est)-1)/2,-1,1))，轉成degree。
e_t = ||t_est-t_GT||，單位m。
e_p = (1/N)Σ_p ||(R_est-R_GT)p+(t_est-t_GT)||，p遍歷來源scan所有finite原始點。先對每案例取平均，再跨案例取mean；不是把所有scan點混一起依點數加權。
e_p和Eq.(8)對應，量測估計與benchmark GT變換施加到同一來源點的差；不是source到target最近鄰、不是point-to-plane loss，也不是樹位誤差。來源方向改變可能改e_p，因來源點集合不同。
本專案成功門檻：e_R<2° AND e_t<0.5m，嚴格小於。e_p並未直接進這個二元門檻。論文未確認完全相同門檻，因此不能把同為100%視為同難度標準。
e_p低至0.6cm可以小於0.1m finest voxel，因pose是連續估計、利用多點約束；它不是宣稱每個點量測、每棵樹或絕對世界位置都有6mm真實精度。GT不確定度亦未完整建模。

統計人口三種，不能混：
finite mean：有finite變換可算e_p的都納入，即使未成功。沒有變換仍算success失敗，e_p留缺，不填0；要顯示finite N。
success-only mean/median：僅通過門檻的子集，容易受前後成功集合變動影響。
common-success paired mean：同pair、同seed且兩邊成功才比較，回答原本可配準案例是否更精，但有存活者偏差，不能取代全體success/finite mean。
N不同時finite mean的改善可能也受缺輸出集合改變影響，務必與success與missing一起報。Plot4 reference5只有83個finite、不是100。
P90用線性插值quantile；一般統計median偶數採兩中值平均，optimizer/height seed用的upper median另有說明。

差異算法：成功率差用percentage points，例如34%→100%是+66pp，不是+66%相對增幅。
誤差／時間下降率=(baseline-ours)/baseline；正數下降、負數變差；baseline=0或缺值則待補。1m→0.2m為80%下降（教學假例）。
相較paper與相較local baseline要分別標基準。論文e_p與本地變換誤差定義可對照，但配對分母、門檻與baseline程式未齊，數值參考不等於嚴格公平排名。
Tongji四Plot每個100例，四場景macro mean與400例case-weighted mean此處恰相同；不要任意再與ETH150例混成五場景平均、兩dataset平均或550例平均而不交代權重。
十個seed敏感度不是十個獨立森林，配對還共享站點。沒有事先獨立新場景與適當獨立樣本設計，不應憑550/550直接給跨森林顯著性或泛化信賴區間。

======================================================================
11｜實驗接觸歷史、隨機性與「固定」真正代表什麼
======================================================================
先在ETH seed1～10探索CN設定；seed11～20做新種子對照及recovery開發；seed21～30做凍結recovery檢查；再把ETH選定粗設定首次固定用到Tongji。
之後看過Tongji表現，做Plot2 DTM開發與Plot4候選失敗診斷，再完成最終九版seed31～40統一測試。最終所有五場景都已接觸，不能回頭把首次跨來源遷移的獨立性套到最終v3。
固定的意思：最終所有場景同一套參數與方法，不替每個Plot選最好版本／反向／seed。不是說人類開發過程從未看GT，也不是所有結果都是第一次接觸dataset。
ETH主表用15個全配對；官方pairs.txt另外10個子集有輔助統計，不能把兩種分母混用。Tongji每Plot10對方向固定高號→低號（S2-S1等），ETH低號→高號（s1-s2等）。GT方向組合檢查55對誤差<1e-4。
完整配對清單：Tongji每Plot均為S2-S1、S3-S1、S3-S2、S4-S1、S4-S2、S4-S3、S5-S1、S5-S2、S5-S3、S5-S4；ETH為s1-s2、s1-s3、s1-s4、s1-s5、s1-s6、s2-s3、s2-s4、s2-s5、s2-s6、s3-s4、s3-s5、s3-s6、s4-s5、s4-s6、s5-s6。pair名稱左側是來源，右側是目標。
seed影響距離保留及圓柱擬合；註冊RANSAC另固定seed1；geometry取樣由排序後索引決定。不能把seed31解釋成第31個掃描站。
有資源中斷與同案例續跑；不能宣稱所有取得從未被打斷。未完成batch恢復同清單，保留首次完整輸出、不因GT失敗挑成功重跑。硬體／耗時口徑須另記。

其他支線不可拿來增加最終獨立dataset數：
WHU-Park不是論文WHU-FGI。舊固定ETH粗配準0/310、密度探索2/310；另稀疏保留站點0→2/120，分母不同，不可混。沒有最終v3 Park完整測試，所以不能說v3在Park已失敗或成功。
Park有樹不代表可提足夠穩定樹位：小圓弧、遮擋、密度、雜物和聚類可能破壞提取；這些是機制假說，缺逐項對照就不能指定唯一原因。
Liang／TreeTool處理後六樣區只做提取匹配，macro rate control38.34%、vertical08534.94%、ransac99938.52%；不是有GT分站變換的registration benchmark，不報為第三個驗證成功dataset。
graph_mst六站150/150使用跨站圖，不是本輪每次兩站的條件；距離retention其他函數支線也沒證明最終全配對更好。

======================================================================
12｜主要數字：目前到底哪裡更好、更壞
======================================================================
以下cm皆是每案例全來源點e_p再跨finite案例mean，非成功限定；附錄DATA_U有全部45列精確m數字。
場景    paper TableI  reference5_raw      eth10_raw          multi_v3_refined
Plot1   5.2558cm      100/100;16.812648   100/100;17.423531  100/100;2.886601
Plot2   11.4808cm     33/100;54.504972    34/100;50.900681   100/100;1.185416
Plot3   4.8756cm      100/100;5.376451    100/100;4.901756   100/100;1.118316
Plot4   4.3560cm      76/100;123.234679   95/100;109.373197  100/100;3.392023
ETH     15.3427cm     114/150;12.620930   150/150;8.189515   150/150;0.603814
finite N：reference5為100/100/100/83/127；eth10_raw為100/100/100/99/150；最終全部等於總N。
論文各場景100%，但各分母待補；不把本地分母填給論文。
相較paper e_p數值下降依序45.0778%、89.6748%、77.0630%、22.1299%、96.0645%；這是報告值對照，非原作者程式同環境重跑證據。
最終Tongji400/400 mean e_p約2.145589cm；ETH150/150 mean約0.603814cm，兩來源分開報。

全550成功數：reference5_raw423；reference5+舊refinement507；reference5+v3backend507；eth10_raw479；eth10+舊refinement546；eth10+v3backend546；multi_v1 546；multi_v2 528；multi_v3 550。
「v2更多搜尋但退步」以及「DTM Plot3變差」是成果的一部分，不能從報告拿掉。v3不為Plot3退步換回舊版，維持統一設定。
若只回答『我們比論文好幾倍』會漏掉基準不等價、資料接觸與時間代價；較準確是報各場景下降百分比和上述限制。

======================================================================
13｜時間與資源：不能把不同時間欄硬相減
======================================================================
論文TableII matching ms，Plot1/2/3/4/ETH=20.718/30.699/4.5596/19.522/3.0405；TableIII total s=11.657/16.712/8.3584/9.4516/8.6992。
paper硬體i7-9700F 3.00GHz，C++ PCL OpenMP；實際thread數與RAM未明。不能拿本機數百毫秒refinement去比paper數毫秒matching稱快或慢幾倍。
新epoch20261025T162801Z：Xeon Platinum8573C、CPU quota4、32GiB、每native process OMP4。online欄不含每站樹位提取、geometry cache建立與GT評估，且geometry cache載入在per-case Timer外；新版包含樹位CSV讀取與descriptor/matching/recovery/solve/proposal/refine。
舊cm_register的time_ms主要descriptor+matching+recovery+solver，樹位CSV讀取在外。不同程式的time_ms/online_ms不是天生相同範圍。
v2/v3跨epoch與不同取得並行度，因此主表online均值可保留，但無公平速度下降率；待補同環境重測。

另外量了同新epoch、native串行、固定seed31代表pair的cold case（每法每dataset僅1次）：
ETH s1-s2：reference加同v3backend344.3s，完整v3 369.1s，約慢7.2%；native子程序峰值約1422MiB。
Tongji Plot1 S2-S1：703.9s對651.8s，約快7.4%；峰值約2009MiB。
包含重新提取兩站與建立geometry cache；OS page cache未刻意清，故cold是application cache重建，不保證儲存設備全冷。樣本僅1次，不能當穩定速度排名或顯著差異。
峰值是最大native child RSS，非Python父程序＋所有子程序＋OS page cache整機總和。
26站feature成本另列：4站重用同epoch同設定首次cold實測、22站新量測；後者最多4worker各OMP4共享quota4，不能把並行wall當單程序獨占時間。舊v1 Plot1 cold曾與workbook rendering同時，亦不可拿作嚴格排名。
中斷未完整取得wait4的CPU/RSS成本是未知，不填0。GT評估也不是演算法推論的一部分，若講整個實驗總成本需另算。
共用站點特徵可以攤銷，但不能把一次性成本消失。例如同一站供多pair使用時，說清楚報的是cache後pair latency或含預處理end-to-end，不用一個數字混講。

======================================================================
14｜證據等級與誠實的研究主張
======================================================================
已有實驗支持：
• ETH新seed2×2支持容差與solver的固定條件差異；不是新森林。
• recovery凍結後另一組seed支持同ETH的穩定性，首次Tongji遷移顯示局限。
• 相同樹位與粗方法加整組refinement，提高多場景精度與成功率；5cm相同backend控制也改善。
• v1/v2同定義overlap反例，證明較高舊geometry分數仍可錯得更遠。
• v2→v3整組驗證規則改善Plot4；相同backend控制隔出完整候選版額外4次成功。
• 固定同一最終方法在兩來源五已知場景達550/550，不靠每Plot挑版本。

合理機制，但尚非獨立因果證明：
• 容差容納樹位誤差；地面支配舊分數；較平滑DTM減少高度擾動；多方向surface提供更多約束；v3選到不同微小姿態可使少數e_p變差。
這些有程式與觀察相容，不等於已用單因素實驗排除所有其他解釋。只能講到整組改動效果。

仍待補：
作者程式同硬體同配對baseline；paper成功門檻與分母；真正未見森林；固定候選的vertical-only/gate-only2×2；DTM完整跨場景與ETH；大傾斜／低重疊／樹種分布變化；多次cold runtime統計與GT不確定度。
新對話不要替我們跑這些。使用者要的是理解，未來實作是另一任務。

驗證範圍：此前93個CTest通過，249.48s，是程式測試、不是93個森林；四Excel重開驗證了734項檢核、175公式、10圖表、63組原始數據重聚合。
另用獨立Python在一個v1 multi_refined ETH s1-s2 seed31、19,633,280來源點核對Eq8：native0.006989566414892994m，Python0.006989566414892969m，差約2.515e-17m。不能擴大說4950例全獨立Python重算。

======================================================================
15｜教授常見追問：先給短答，再交代證據或限制
======================================================================
Q01 你到底改descriptor了嗎？
A 最終仍N183/w10/Thscore40；改CN距離容差與後端recovery、surface refinement和候選驗證。早期N93試過但沒放最終，不能宣稱設計全新descriptor。
Q02 為什麼5cm改10cm？
A 允許兩站估計樹間距有更多不一致，可能救回正配對，也可能多誤配。ETH2×2在固定RANSAC時123→142/150支持此條件下有效，不能無限制再放大。
Q03 10cm是不是樹位可差10cm都成功？
A 它是descriptor距離與triangle邊長比較，不是GT樹位允許誤差或最終成功門檻。
Q04 Thscore40與threshold0.1是一樣嗎？
A 前者是wFM+DC的無量綱分數門檻；後者是距離一致性的m單位門檻。
Q05 N183代表什麼？
A 三最近鄰距離＋180扇區值，每區2°；不是183棵樹。三參考方向的矩陣再優先合成一條descriptor。
Q06 為什麼三個最近鄰只要一個通過？
A 這是本地對gatekeeper的實作，容許部分最近鄰受漏樹影響；仍須DC/FM與幾何驗證，不能只靠一距離就接受pose。
Q07 有幾個RANSAC？
A 至少須分圓柱RANSAC（點＋normal，2cm）與註冊RANSAC（樹位對，30cm）。recovery枚舉三元組不等同這兩者的隨機抽樣。
Q08 不足6棵一定不能配準？
A 不對，3個非共線3D對已能解Kabsch；6是弱配對觸發／共識接受冗餘規則。
Q09 是不是把六站一起用才100%？
A 最終每次只當前兩站。ETH共有六站，15個全配對。舊graph_mst用六站圖是另個任務。
Q10 recovery與multistart差在哪？
A 前者輸出樹位對、只評CN證據池的3D一致性；後者保留多pose、做dense微調，再用全當前兩站樹位及geometry選pose。
Q11 為何recovery成功還會跑multistart？
A 觸發檢查原match.pairs<6，不是recovered.pairs數量，兩者刻意不同。
Q12 100%後為什麼還要改？
A 成功門檻是粗二元標準。ETH粗已150/150，e_p8.19cm仍可由surface降至0.604cm；另Tongji不全成功。
Q13 為何DTM會影響旋轉？
A 高度進3D樹位，Kabsch試圖同時擬合xyz；有系統z偏差可能以roll/pitch補償。這是機制解釋，非每例都已確認。
Q14 smooth為什麼沒用最終？
A Plot2有改善，但Plot3誤差退步；全流程只驗Plot2，不能說跨dataset有效。最終v3另解refinement/驗證問題且維持原DTM。
Q15 DTM48/100就是一般化了？
A 沒有，Plot2參與GT選profile；同樣48在固定軸與全提取也不是相同成功案例，其他全流程／ETH待補。
Q16 downsampling最後選哪種？
A step 0.5/0.75/1，水平距離5/10m分界。其他函數只有支線，不拿未完成比較宣稱最佳。
Q17 1cm改成50cm會不會丟精度？
A 沒有直接替換；提取仍1cm，surface另建50/20/10cm層，連續pose可細於voxel尺寸。輸出誤差不等於每點物理測量精度。
Q18 ICP到底最小化什麼？
A point-to-plane r=n·(Tp-q)的robust局部殘差；e_p是事後對GT變換的全來源點差，二者不同。
Q19 ICP用了幾個點？
A 各層最多20,000來源點，驗證雙向各8,000；e_p全來源finite點。不能只回答『全部』或『八千』而不說階段。
Q20 Huber與trimming各幹嘛？
A trimming先排大距離尾部，Huber再降低大平面殘差權重；屬整組設計，沒有單獨證明哪個貢獻最大。
Q21 退化檢查能確保全域正確？
A 不能，只檢查局部Hessian是否缺約束；另一個錯誤但局部約束充分的盆地仍可通過。
Q22 保留原初值為何v2還退步？
A 保留它進beam不等於最後一定選它。錯誤候選的舊分數更高仍會勝出，Plot4真例支持這點。
Q23 你證明地面是唯一元兇？
A 沒有。舊分數可被錯位騙過已證實，地面支配是合理機制線索，還缺固定候選方向／gate消融。
Q24 v3把地面刪掉會不會丟z約束？
A 只在verification偏好垂直表面；optimizer仍用全方向。直立假設和不足點fallback的限制仍存在。
Q25 verticality0.9降到0.3？
A 不是同一參數。提取1-|nz|>0.9沒變；另新增surface verification |nz|<0.3。
Q26 哪個v3參數救了22次？
A 目前只能說整組驗證修改；沒有vertical-only和gate-only2×2，不能指派單參數功勞。
Q27 v3有沒有更差的地方？
A Plot3平均多約0.127mm；Plot4共同成功78例平均多約0.733mm；總體成功率救回災難錯配不保證每例更精。
Q28 多候選帶來全部127次提升？
A 沒有。同v3backend的eth10已有546/550，完整候選550/550只額外4次；reference同backend507也顯示ICP本身很重要。
Q29 mean很差但median很好是不是log壞了？
A 不一定。v2 Plot4多數誤差小但少數數十m，median約0.903cm而mean864.762cm；要保留失敗並看分布，不能只報median。
Q30 沒有pose的e_p是多少？
A 缺值，不是0。仍算success失敗，finite平均要標N；不能把缺值刪掉後稱全成功。
Q31 e_p和e_t怎麼不同？
A e_t只比較t；e_p還受R差與來源點距離影響。來源越遠，小角差可造成大點位移，見20m教學例。
Q32 100%和paper相同，怎麼談超越？
A 在這些固定資料上e_p報告值更低；但作者實作與門檻／分母未齊，不能嚴格稱超越原作者程式。
Q33 你用GT嗎？
A 單次推論不讀GT；離線開發看GT評估並選過設定，尤其DTM。若要求任何階段都不能GT選參，目前不滿足強限制，必須承認。
Q34 不同seed能當test set？
A 可測同資料隨機敏感度，不能當新森林；全部站點都相同且配對共享站。最終也接觸過兩dataset。
Q35 為什麼有九版本卻只有四次報告？
A 報告按有意義改動分組：容差／recovery、DTM支線、surface refinement、多候選與驗證。不是每版本或每Plot硬拆一篇。
Q36 速度有超越嗎？
A 沒有公平paper速度結論。計時範圍硬體不同；同epochcold只各一次，ETH變慢、Plot1變快，不代表穩定排名。
Q37 資料cache成本算了嗎？
A 分別記錄預處理與cache後online，也有含重建cold case。不能以共用cache為由說它免費，還須交代攤銷方式。
Q38 你確定不是GT方向反了？
A 已檢查55對GT方向組合<1e-4，source-target固定；另一例全來源e_p有Python獨立核對。不是每4950例都獨立重算。
Q39 Park有樹為什麼不行？
A 有樹不保證相同提取假設成立。舊粗方法結果確實差，但遮擋／密度／群集等尚未隔離；最終v3 Park未完成，不對未測方法下結論。
Q40 這樣能做論文嗎，下一步怎麼證明？
A 目前有可追溯的benchmark方法改進與反例，研究潛力不等於已足夠投稿。最關鍵缺的是新森林盲測、作者公平baseline與候選固定的消融；只能列待完成，別寫已驗證。
Q41 你能保證新森林100%？
A 不能；模型假設、掃描姿態、樹排列、遮擋與重疊會變。550/550只描述既定有限測試。
Q42 v3是不是純粹參數調整？
A 包含新後端流程與候選驗證方法，也有啟發式常數。可以清楚列實作改動，但缺文獻全面對照不能僅據本檔宣稱每個元件具學術新穎性。

======================================================================
16｜你應該練到能自己解釋的事
======================================================================
先不看答案，試著畫兩條資料路徑：樹位提取與surface cache；說出GT在哪層出現。
用自己的話解釋：normal不是axis；N183不是183樹；6支持不是6站；threshold0.1m不是Thscore；ICP loss不是e_p。
選Plot2說明raw→refined，選Plot4說明v1→v2→v3，選Plot3說明保留退步；每次指出哪個對照固定了什麼、哪組參數仍無法拆因果。
手算34/100→100/100是+66pp；5.2558→2.8866cm的下降率約45.1%；檢查公式分母用哪一版。
最後練30秒回答：「GT-free inference，但GT參與離線開發；最終已見資料benchmark改善，未完成新森林驗證；原作者程式baseline待補。」這段比背100%更重要。

可直接貼給新對話的開場：
「請用我附的實作交接檔與CircularMatch論文教我，暫時不要做實作。先確認你讀懂GT-free inference與GT參與開發的差別。請從兩站配準流程開始，一次只教一個主題，先用直覺與小例子，再給對應公式／程式。每主題問我兩題確認理解，再模擬教授追問；沒有實驗證據的地方請明講。」

【常見數字速查：同數值可能是完全不同的參數】
0.01m：樹位提取voxel；0.10m：提取normal半徑，或改後CN距離容差，依步驟區分。
0.02m：圓柱幾何inlier距離；0.20m：本輪實際cluster tolerance，也碰巧是10cm CN下recovery邊長相容界限。
0.30m：樹位註冊RANSAC的3D距離，或多候選globalSupport的XY距離，兩者維度不同。
0.3無單位：v3 verification的|normal.z|界限，不是0.3m，也不是提取verticality。
0.7無單位：normal abs dot門檻，亦碰巧是trimming保留比例；不是同一參數。
0.5m：DTM cell，亦是最粗surface voxel，亦碰巧是成功平移門檻與pose聚群平移界限；不可混用。
6：弱配對啟動／接受支持數，不是每次要6站。3：最小Kabsch對應數，不是全部方法只用3點。
40：CN MatchScore門檻；183：descriptor長度；20：v2/v3 descriptor候選數；256：alternative pose上限；8：總refine beam上限。
31～40：十個experiment seed，非樹號、站號、thread數。seed1：另固定註冊RANSAC隨機種子。

以下 DATA/TABLE/CODE 為本檔內建證據，不必再上傳其他資料。
"""


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    manifest = json.loads((REPORTS / "source_manifest.json").read_text())
    parts = [NARRATIVE]
    provenance = []

    def heading(title: str) -> None:
        parts.append("\n" + "=" * 70 + "\n" + title + "\n" + "=" * 70 + "\n")

    def include(path: Path, title: str, numbered: bool = False) -> None:
        data = path.read_bytes()
        rel = str(path.relative_to(ROOT))
        provenance.append({"path": rel, "sha256": sha(data), "bytes": len(data)})
        heading(title)
        parts.append(f"來源：{rel}\nSHA256：{sha(data)}\n")
        content = data.decode("utf-8-sig")
        if numbered:
            parts.append("行號僅供閱讀定位；不得當成可直接編譯內容。\n")
            content = "\n".join(f"{i:04d} | {line}" for i, line in enumerate(content.splitlines(), 1))
        parts.append(content + "\n")

    # Independently reaggregate the 4,950 cases before embedding existing summaries.
    rows_path = ROOT / "results/unified_20261025/registration_rows.csv"
    raw_bytes = rows_path.read_bytes()
    assert sha(raw_bytes) == manifest["sources"]["U"]["sha256"]
    raw = list(csv.DictReader(io.StringIO(raw_bytes.decode())))
    assert len(raw) == 4950
    grouped = defaultdict(list)
    for row in raw:
        grouped[(row["scene"], row["method"])].append(row)
    summary_path = ROOT / "results/unified_20261025/scene_summary.csv"
    summary = list(csv.DictReader(summary_path.open()))
    assert len(summary) == 45 and len(grouped) == 45
    # Field names are checked below against the actual source schema.
    checks = []
    import math
    import statistics

    for row in summary:
        key = (row["scene"], row["method"])
        group = grouped[key]
        success_key = "success" if "success" in group[0] else "ok"
        ep_key = "e_p_m"
        ok = sum(g[success_key].lower() in {"1", "true"} for g in group)
        values = [float(g[ep_key]) for g in group if g[ep_key] and math.isfinite(float(g[ep_key]))]
        assert len(group) == int(row["trials"])
        assert ok == int(row["success_trials"])
        assert len(values) == int(row["finite_error_cases"])
        assert math.isclose(statistics.mean(values), float(row["mean_finite_ep_m"]), rel_tol=1e-12, abs_tol=1e-12)
        checks.append({"scene": key[0], "method": key[1], "trials": len(group), "ok": ok, "finite": len(values)})

    heading("DATA_P｜論文表值與比較限制")
    parts.append(json.dumps(manifest["paper_values"], ensure_ascii=False, indent=2) + "\n")
    parts.append("Table I：e_p以m；Table II：matching以ms；Table III：total以s。成功率分母／門檻與e_R/e_t待補；時間不算跨硬體下降率。\n")
    include(summary_path, "DATA_U｜主實驗45列完整場景摘要（原始CSV精度）")
    parts.append("上述rate為0～1；ep/trans單位m；rotation為deg；online與各stage為ms。pairs_all_seeds_success是所有十seed皆成功的唯一pair數，不是trial數。所有mean_finite包含有限失敗；time不含提取與cache，跨epoch條件限制見13章。\n")

    heading("DATA_CASE｜同一困難案例的逐版本輸出摘要")
    cases = [r for r in raw if r["scene"] == "Plot4" and r["seed"] == "31" and "S4-S2" in str(r)]
    # Select by the pair field only, to avoid matching an incidental path.
    cases = [r for r in cases if r.get("pair", r.get("pair_id", "")) == "S4-S2"]
    assert len(cases) == 9, (len(cases), list(raw[0]))
    desired = ["method", "success", "pose_ok", "e_p_m", "rot_err_deg", "trans_err_m", "proposal_count", "refined_candidates", "geo_overlap", "geo_rmse", "iterations"]
    fields = [k for k in desired if k in raw[0]]
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(cases)
    parts.append(buf.getvalue())
    parts.append("這是Plot4 S4-S2 seed31。v3 geometry分數定義已改，僅v1/v2可直接比raw overlap；GT e_p定義九版一致。pose_ok只表示有變換，不等於成功。raw版本未做geometry驗證，geo_overlap=0、geo_rmse=1e9為未計算預設值，不能當實際重疊率／誤差；該預設不影響GT e_p。\n")

    include(REPORTS / "實驗盤點與數據口徑.txt", "DATA_HISTORY｜完整支線盤點與負面結果")
    parts.append("補充澄清：此歷史盤點所說不靠GT選輸出，指單次推論；離線開發GT選參限制以本檔01/11/14章為準。其他支線率的定義見各段，非主registration成功率。\n")

    # Embed cached spreadsheet values and short notes; do not resave the workbooks.
    for index in range(1, 5):
        book = next(REPORTS.glob(f"report_{index:02d}_*.xlsx"))
        data = book.read_bytes()
        provenance.append({"path": str(book.relative_to(ROOT)), "sha256": sha(data), "bytes": len(data)})
        cached = load_workbook(book, data_only=True)
        formulas = load_workbook(book, data_only=False)
        heading(f"TABLE_R{index:02d}｜第{index}次報告Excel的全部比較表與註記（內嵌值）")
        parts.append(f"來源：{book.name}；SHA256：{sha(data)}\n表格列以TAB分隔；數值採Excel已重算cache；下方保留原表註解，圖表資料已在表中，未嵌入圖像。\n單位：本附錄error表e_p全為cm；pose表e_R為degree、e_t為cm；time依欄名ms或s；rate為百分率，pp欄為百分點，下降率為百分比。其他表依欄名，DATA_U的e_p則是m。\n")
        formula_count = 0
        for tab, spec in manifest["table_ranges"][str(index)].items():
            ws = cached[spec["sheet"]]
            fw = formulas[spec["sheet"]]
            parts.append(f"\n[{spec['sheet']} / {tab} / {spec['ref']}]\n")
            from openpyxl.utils.cell import range_boundaries
            cmin, rmin, cmax, rmax = range_boundaries(spec["ref"])
            for cells in ws.iter_rows(min_row=rmin, max_row=rmax, min_col=cmin, max_col=cmax):
                values = []
                for cell in cells:
                    value = cell.value
                    original = fw[cell.coordinate]
                    if original.data_type == "f":
                        assert value is not None and cell.data_type != "e", (book.name, cell.coordinate)
                        formula_count += 1
                    if value is None:
                        text = "—（空欄）"
                    elif isinstance(value, (int, float)) and not isinstance(value, bool):
                        text = f"{value * 100:.6f}%" if "%" in cell.number_format else f"{value:.12g}"
                    else:
                        text = str(value).replace("\n", " / ").replace("\t", " ")
                    values.append(text)
                parts.append("\t".join(values) + "\n")
        parts.append("\n[各頁原註解與文字欄，數表以以上TAB內容為準]\n")
        seen = set()
        for ws in cached:
            for row in ws:
                texts = [str(c.value).replace("\n", " / ") for c in row if isinstance(c.value, str)]
                # Keep long notes, explanations and source labels; skip ordinary headers.
                for text in texts:
                    if len(text) >= 35 and text not in seen:
                        parts.append(f"{ws.title}：{text}\n")
                        seen.add(text)
        parts.append(f"本次抽取公式cache檢查：{formula_count}個非空、無Excel錯誤。\n")
        cached.close()
        formulas.close()

    include(ROOT / "results/unified_20261025/unified_config.json", "CONFIG_ACTIVE｜統一主實驗有效設定快照")
    parts.append("這是記錄，不是native動態載入的唯一設定檔。v3 protocol的pre-test status屬建立時狀態，不能拿來否認後續已完成結果。\n")

    code_paths = [
        "include/cm/cn_descriptor.hpp", "include/cm/cn_matching.hpp",
        "src/cn_descriptor.cpp", "src/cn_matching.cpp", "docs/cn_ambiguities.md",
        "include/cm/registration.hpp", "src/registration.cpp",
        "include/cm/cn_recovery.hpp", "src/cn_recovery.cpp",
        "src/pipeline.cpp", "src/voxel.cpp", "src/normals.cpp", "src/retention.cpp",
        "src/dtm.cpp", "src/cylinder.cpp", "src/tree_position.cpp",
        "config/default.ini", "apps/cm_dtm_reposition.cpp",
        "include/cm/geometry_refinement.hpp", "apps/cm_geometry_register_v3.cpp",
        "apps/geometry_refinement_structural.cpp", "apps/cm_evaluate_transforms.cpp",
    ]
    heading("CODE_INDEX｜程式閱讀地圖與適用邊界")
    parts.append("CN：buildCnDescriptors→cnFeatureMatrix；matching：matchCn→cnGatekeeperPass/cnDcScore/cnTriangleMatch。粗姿態：kabsch/kabschRansac；補救：recoverCn。提取：pipeline→DTM/normal/retention/cylinder/intersection。最終推論：infer→proposals→geometryScore/refineGeometry；評估：cm_evaluate_transforms。\n")
    parts.append("以下是目前已核對的關鍵程式全文；其include依賴不全，故不是可執行套件。展示統一LF並略去行尾空白，SHA256則指原始檔bytes。docs/cn_ambiguities與default.ini含歷史註解，階段範圍以03章說明為準；cn_recovery.hpp的no raw-cloud ICP只指該recovery函式，不是後續整個方法。\n")
    for path in code_paths:
        include(ROOT / path, f"CODE｜{path}", numbered=True)

    for old, new in [
        ("apps/cm_geometry_register.cpp", "apps/cm_geometry_register_v2.cpp"),
        ("apps/cm_geometry_register_v2.cpp", "apps/cm_geometry_register_v3.cpp"),
        ("apps/geometry_refinement.cpp", "apps/geometry_refinement_structural.cpp"),
    ]:
        heading(f"DIFF｜{old} → {new}")
        parts.append("以下是實際文字差異；未變更的上下文見前述完整v3程式，非額外新實驗。\n")
        a, b = (ROOT / old).read_bytes(), (ROOT / new).read_bytes()
        for path, data in [(old, a), (new, b)]:
            parts.append(f"{path} SHA256 {sha(data)}\n")
            if not any(p["path"] == path for p in provenance):
                provenance.append({"path": path, "sha256": sha(data), "bytes": len(data)})
        parts.extend(difflib.unified_diff(a.decode().splitlines(True), b.decode().splitlines(True), fromfile=old, tofile=new, n=3))

    heading("PROVENANCE｜原始數據來源清單（理解本檔無須下載）")
    parts.append("實驗snapshot的SHA如下。SHA能辨識檔案內容，不能單靠雜湊證明科學結論；關鍵數值和程式已內嵌。本檔不含私人Drive識別碼、連結或原始點雲。\n")
    for key, source in manifest["sources"].items():
        parts.append(f"{key} | {source['path']} | SHA256 {source['sha256']} | {source.get('description', '')}\n")
    parts.append("\n本次嵌入檔案快照：\n" + json.dumps(provenance, ensure_ascii=False, indent=2) + "\n")
    heading("END｜給接手教學對話的最後提醒")
    parts.append("你已具備教學所需的實作、設定與比較摘要；先回答使用者眼前的問題，不必重述整份文件。請協助他能自己說明因果鏈、固定條件與限制。不得補造尚未列出的配對資料，不得把合理推測改寫成已驗證結論；推論不讀GT與研發看GT選參務必分開。\n")
    text = "\n".join(line.rstrip() for line in "".join(parts).replace("\r\n", "\n").replace("\r", "\n").splitlines()) + "\n"
    # Prevent accidental disclosure of private Drive identifiers or access URLs.
    forbidden = ["drive.google.com", "1F9v5CyDlGkN1PRMipLdJzpuA4x1pAO_l", "access_token", "refresh_token"]
    assert not any(term in text for term in forbidden)
    for marker in ["01｜", "16｜", "DATA_U", "TABLE_R04", "CODE_INDEX", "GT-free inference", "4,950"]:
        assert marker in text
    TARGET.write_text(text, encoding="utf-8")
    reopened = TARGET.read_text(encoding="utf-8")
    assert reopened == text
    result = {
        "generated_at": datetime.now(ZoneInfo("Asia/Taipei")).isoformat(),
        "file": TARGET.name,
        "bytes": TARGET.stat().st_size,
        "sha256": sha(TARGET.read_bytes()),
        "lines": len(text.splitlines()),
        "main_case_rows_reaggregated": len(raw),
        "main_scene_method_groups_verified": len(checks),
        "case_versions_embedded": len(cases),
        "spreadsheets_read_only": 4,
        "key_source_files_embedded": len(code_paths),
        "private_drive_identifiers_present": False,
        "source_files": provenance,
    }
    (OUT / "validation.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({k: v for k, v in result.items() if k != "source_files"}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
