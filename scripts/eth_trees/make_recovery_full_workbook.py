#!/usr/bin/env python3
"""Cumulative detailed workbook: prior study plus complete two-scan recovery runs."""
import csv
import json
import math
from pathlib import Path
import statistics as stats

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter as letter
from make_study_workbook import table
from recovery_study import ROOT, OUT, RUNS, RAW, load_frozen
from tune_registration import read_csv


def cell(value):
    if value is None or value=='':return None
    if isinstance(value,(int,float,bool)):return value
    try:
        n=float(value)
        if not math.isfinite(n):return None
        return int(n) if n.is_integer() else n
    except ValueError:return value


def csv_table(wb,name,path):
    rows=read_csv(path);assert rows
    headers=list(rows[0])
    widths=[32 if h in ['method','solver','diagnosis','label'] else 22 for h in headers]
    return table(wb,name,headers,[[cell(r[h]) for h in headers] for r in rows],widths)


def json_rows(data,prefix=''):
    rows=[]
    for key,value in data.items():
        name=f'{prefix}.{key}' if prefix else key
        if isinstance(value,dict):rows.extend(json_rows(value,name))
        else:rows.append([name,json.dumps(value,ensure_ascii=False) if isinstance(value,list) else value])
    return rows


def main():
    frozen=load_frozen()
    prior=ROOT/'results/cross_dataset_20261025/comparison.xlsx'
    wb=load_workbook(prior);previous_names=wb.sheetnames.copy()
    overview=[['內容範圍','保留前輪ETH參數與六站研究、Liang外部提取，另加入本輪完整兩站補救比較。'],
              ['最新驗證','seed21–30：600列＝兩種方法×兩種solver×150個配對案例；主要比較為svd+ransac。'],
              ['本輪主要結果','原推薦組142/150，加補救後150/150；官方列表98/100→100/100，補回8次、退步0次。'],
              ['改動','少於6對匹配時，收集前三個描述子候選的證據，使用完整3D一致性補救；推論不讀GT。'],
              ['保持固定','step距離降採樣、clustering0.20m、垂直度0.90、匹配閾值0.10m、配準RANSAC。'],
              ['新功能預設','補救預設關閉，須指定--recovery=evidence --recovery_trigger=6。'],
              ['資料分工','seed11–20用於本輪診斷與選方案；之後固定程式及參數，才產生seed21–30驗證。'],
              ['舊seed候選','1200列＝四種候選（含原推薦組）×兩種solver×150案例；已用於選方案，不能再視為盲測。'],
              ['舊seed點雲誤差','候選篩選省略原始點雲讀取與e_p；空白為未計算，不是零誤差。新驗證則計算完整source點雲e_p。'],
              ['工作表閱讀','本輪工作表排在前面；「說明」起的原17張工作表保留前輪資料及其原本seed範圍。'],
              ['成功定義','旋轉誤差<2°且平移誤差<0.5m；與前輪相同。'],
              ['精度比較','成功集合可能不同，配對比較另列兩邊共同成功142案例的e_p中位數。'],
              ['計時','僅描述子、匹配與求解；不含提取及完整點雲誤差，主機並行工作，非獨占效能基準。'],
              ['論文比較限制','論文未明確列成功門檻，ETH配對列表仍未完全確認；不能稱原論文100%的嚴格重現。'],
              ['外部驗證限制','新seed仍來自同一片ETH森林；Liang處理版只能做提取驗證，缺少分站與配準真值。'],
              ['固定時間UTC',frozen['frozen_at_utc']],
              ['精簡版','同目錄comparison_simple.xlsx；使用同一批seed21–30驗證數據。']]
    table(wb,'最新說明',['項目','內容'],overview,[30,125])
    rows=read_csv(RUNS/'validation_rows.csv');assert len(rows)==600
    official={tuple(s.split()) for s in (RAW/'pairs.txt').read_text().splitlines() if s.strip()}
    headers=['method','solver','seed','pair','official']+[k for k in rows[0] if k not in ['method','solver','seed','pair']]
    raw=[]
    for r in rows:
        r=dict(r,official=int(tuple(r['pair'].split('-')) in official))
        raw.append([cell(r[h]) for h in headers])
    table(wb,'新驗證逐次',headers,raw,[32 if h in ['method','solver','label','setting'] else 20 for h in headers])
    poses=[]
    posekeys=[f'r{i}{j}' for i in range(3) for j in range(3)]+[f't{i}' for i in range(3)]
    for r in rows:
        poses.append([r['method'],r['solver'],int(r['seed']),r['pair'],int(r['transform_ok'])]+
                     [cell(r[k]) if r['transform_ok']=='1' else None for k in posekeys])
    table(wb,'新估計變換',['method','solver','seed','pair','transform_ok']+posekeys,poses,[30,20,12,16,18]+[18]*12)
    summary=read_csv(OUT/'validation_summary.csv');summaries=[]
    def region(key):
        col=letter(headers.index(key)+1);return f"'新驗證逐次'!${col}$2:${col}$601"
    for i,r in enumerate(summary,2):
        criteria=f'{region("method")},A{i},{region("solver")},B{i}'
        summaries.append([r['method'],r['solver'],f'=COUNTIFS({criteria})',
                          f'=COUNTIFS({criteria},{region("success")},1)',f'=D{i}/C{i}',
                          f'=COUNTIFS({criteria},{region("official")},1)',
                          f'=COUNTIFS({criteria},{region("official")},1,{region("success")},1)',f'=G{i}/F{i}']+
                         [cell(r[k]) for k in ['median_success_ep_m','mean_match_solve_ms','median_match_solve_ms',
                                               'p95_match_solve_ms','recovery_triggered','recovery_adopted']])
    ws=table(wb,'新驗證摘要',['方法','solver','案例數','成功數','成功率','官方案例數','官方成功數','官方成功率',
                            '成功e_p中位數m','平均匹配求解ms','匹配求解中位數ms','匹配求解p95 ms','觸發補救數','採用補救數'],summaries,
             [26,20,15,15,15,18,18,18,24,24,25,25,18,18])
    for i in range(2,6):
        for j in [5,8]:ws.cell(i,j).number_format='0.0%'
    seeds=[];pairs=[]
    for method in ['baseline','recovery']:
        for solver in ['svd','svd+ransac']:
            rr=[r for r in rows if r['method']==method and r['solver']==solver]
            for seed in range(21,31):
                sr=[r for r in rr if int(r['seed'])==seed];n=sum(int(r['success']) for r in sr)
                seeds.append([method,solver,seed,len(sr),n,n/len(sr)])
            for pair in sorted({r['pair'] for r in rr}):
                pr=[r for r in rr if r['pair']==pair];n=sum(int(r['success']) for r in pr)
                pairs.append([method,solver,pair,int(tuple(pair.split('-')) in official),len(pr),n,n/len(pr)])
    for name,h,data,widths in [('新每seed',['方法','solver','seed','案例數','成功數','成功率'],seeds,[26,20,12,15,15,18]),
                               ('新各配對',['方法','solver','pair','官方列表','案例數','成功數','成功率'],pairs,[26,20,16,16,15,15,18])]:
        ws=table(wb,name,h,data,widths)
        for i in range(2,ws.max_row+1):ws.cell(i,len(h)).number_format='0.0%'
    table(wb,'新配對比較',['項目','值'],json_rows(json.loads((OUT/'paired_validation.json').read_text())),[45,125])
    csv_table(wb,'舊seed候選摘要',OUT/'discovery_summary.csv')
    discovery=[]
    for method in ['baseline','triangle_scarce','evidence_scarce','evidence_weak']:
        for path in sorted((RUNS/method).glob('*.csv')):
            seed,pair=path.stem.split('_',1)
            discovery.extend(dict(method=method,seed=int(seed),pair=pair,**r) for r in read_csv(path))
    assert len(discovery)==1200
    dh=list(dict.fromkeys(k for r in discovery for k in r))
    table(wb,'舊seed候選逐次',dh,[[cell(r.get(k)) for k in dh] for r in discovery],
          [32 if k in ['method','solver','label','setting'] else 20 for k in dh])
    csv_table(wb,'舊seed失敗診斷',OUT/'failure_audit.csv')
    params=json_rows(json.loads((OUT/'plan.json').read_text()))+json_rows(frozen,'frozen')
    table(wb,'補救參數與凍結',['項目','值'],params,[64,125])
    csv_table(wb,'新提取摘要',OUT/'validation_extraction_summary.csv')
    extraction=[]
    for scan in [f's{i}' for i in range(1,7)]:
        for seed in range(21,31):
            path=RUNS/'validation_extraction'/scan/f'step_seed{seed}_trees.csv'
            tr=read_csv(path)
            extraction.append([scan,seed,len(tr),str(path.relative_to(ROOT))])
    table(wb,'新提取逐次',['站點','seed','偵測樹數','輸出檔案'],extraction,[16,14,20,115])
    checks=[]
    for name in ['audit_validation','gt_independence','validation','workbook_validation']:
        checks+=json_rows(json.loads((OUT/f'{name}.json').read_text()),name)
    checks+=[['新驗證CSV列數',600],['舊候選CSV列數',1200],['原始欄位單位','*_deg為度；*_m為公尺；time_ms為毫秒；*_ok、success、used_triangle、official為0/1旗標。'],
             ['GT欄位','n_correct_matches及失敗診斷的gt_*僅供評分；補救API不接收GT。']]
    table(wb,'新核驗紀錄',['項目','值'],checks,[64,125])
    recent=['最新說明','新驗證摘要','新驗證逐次','新估計變換','新每seed','新各配對','新配對比較',
            '舊seed候選摘要','舊seed候選逐次','舊seed失敗診斷','補救參數與凍結','新提取摘要','新提取逐次','新核驗紀錄']
    wb._sheets=[wb[n] for n in recent+previous_names]
    for name in ['最新說明','新配對比較','補救參數與凍結','新核驗紀錄']:
        ws=wb[name]
        for row in ws.iter_rows(min_row=2):row[1].alignment=Alignment(wrap_text=True,vertical='top')
    target=RUNS/'input/comparison_full.xlsx';target.parent.mkdir(parents=True,exist_ok=True);wb.save(target)
    (RUNS/'full_workbook_manifest.json').write_text(json.dumps(dict(sheets=wb.sheetnames,validation_rows=600,
                                                                discovery_rows=1200,preserved_sheets=previous_names),ensure_ascii=False,indent=2)+'\n')
    print(target);print('sheets:',len(wb.sheetnames))


if __name__=='__main__':main()
