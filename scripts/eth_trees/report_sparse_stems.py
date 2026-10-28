#!/usr/bin/env python3
"""Full and simple reports; preserve the distinction between development and holdout."""
import hashlib
import itertools
import json
import math
import re
import statistics as stats

from openpyxl import Workbook
from make_study_workbook import table
from sparse_stem_study import ROOT,OUT,RUNS,frozen
from tune_registration import read_csv,write_csv


def numeric(x):
    if x in ['',None]:return None
    try:
        n=float(x);return n if math.isfinite(n) else None
    except (ValueError,TypeError):return x


def flatten(data,prefix=''):
    rows=[]
    for key,value in data.items():
        label=f'{prefix}.{key}' if prefix else key
        if isinstance(value,dict):rows+=flatten(value,label)
        elif isinstance(value,list) and value and isinstance(value[0],dict):
            for i,item in enumerate(value):rows+=flatten(item,f'{label}[{i}]')
        else:rows.append([label,json.dumps(value,ensure_ascii=False) if isinstance(value,list) else value])
    return rows


def main():
    f=frozen();plan=json.loads((OUT/'plan.json').read_text())
    development=read_csv(RUNS/'development_rows.csv');validation=read_csv(RUNS/'validation_rows.csv');eth=read_csv(RUNS/'eth_rows.csv')
    assert len(development)==1050 and len(validation)==480 and len(eth)==300
    primary={k:{(int(r['seed']),r['pair']):r for r in validation if r['profile']==k and r['solver']=='svd+ransac'} for k in ['baseline','selected']}
    before,after=primary['baseline'],primary['selected'];assert len(before)==len(after)==120
    gained=[k for k in before if before[k]['success']=='0' and after[k]['success']=='1']
    lost=[k for k in before if before[k]['success']=='1' and after[k]['success']=='0']
    summaries=[]
    for group,rows in [('Park保留站點',validation),('ETH回歸',eth)]:
        for profile in sorted({r['profile'] for r in rows}):
            for solver in ['svd','svd+ransac']:
                rr=[r for r in rows if r['profile']==profile and r['solver']==solver]
                for r in rr:
                    ok=r['transform_ok']=='1' and float(r['rot_err_deg'])<2 and float(r['trans_err_m'])<.5
                    assert int(r['success'])==int(ok)
                    if r['transform_ok']=='1':assert numeric(r['e_p_m']) is not None
                success=sum(int(r['success']) for r in rr)
                ep=[float(r['e_p_m']) for r in rr if r['success']=='1']
                summaries.append(dict(group=group,profile=profile,solver=solver,cases=len(rr),success=success,
                                      rate=success/len(rr),transforms_found=sum(int(r['transform_ok']) for r in rr),
                                      mean_ep_success_m=stats.mean(ep) if ep else '',
                                      mean_match_solve_ms=stats.mean(float(r['time_ms']) for r in rr)))
    write_csv(OUT/'validation_summary.csv',summaries)
    counts=[];policies=[]
    for profile in ['baseline','selected']:
        for scan in range(16,32):
            directory=RUNS/'validation/extraction'/profile/str(scan)
            policy=json.loads((directory/'arc_policy.json').read_text())
            policies.append(dict(profile=profile,scan=scan,**policy))
            log=(RUNS/'logs'/f'validation_extract_{profile}_{scan}.log').read_text()
            actual=[]
            for seed in range(31,41):
                n=len(read_csv(directory/f'step_seed{seed}_trees.csv'));actual.append(n)
                counts.append(dict(profile=profile,scan=scan,seed=seed,candidates=n))
            assert actual==list(map(int,re.findall(r'trees: (\d+) ;',log)))
    write_csv(OUT/'holdout_extraction.csv',counts);write_csv(OUT/'holdout_density_policies.csv',policies)
    original=json.loads((ROOT/'results/whu_park/transfer_20261025/plan.json').read_text())
    from transfer_whu_park import DATA
    assert all(hashlib.sha256((DATA/name).read_bytes()).hexdigest()==expected for name,expected in original['input_sha256'].items())
    regression=json.loads((OUT/'eth_regression.json').read_text())
    assert regression['outputs_identical'] and regression['primary_success']==150
    old={(int(r['seed']),r['pair'][1:].replace('-s','-'),r['solver']):r for r in read_csv(ROOT/'results/eth_trees/failure_followup_20261025/runs/validation_rows.csv') if r['method']=='recovery'}
    compare=['n_matches','n_correct_matches','ransac_inliers','transform_ok','success','rot_err_deg','trans_err_m','e_p_m']
    compare+=[f'r{i}{j}' for i in range(3) for j in range(3)]+[f't{i}' for i in range(3)]
    for r in eth:assert all(r[key]==old[int(r['seed']),r['pair'],r['solver']][key] for key in compare)
    eth_n=sum(int(r['success']) for r in eth if r['solver']=='svd+ransac');baseline_n=sum(int(r['success']) for r in before.values());selected_n=sum(int(r['success']) for r in after.values())
    pairrows=[]
    for a,b in plan['validation_pairs']:
        p=f'{a}-{b}';a0=sum(int(before[s,p]['success']) for s in range(31,41));a1=sum(int(after[s,p]['success']) for s in range(31,41))
        pairrows.append([p,a0,a1,a1-a0,'改善' if a1>a0 else ('退步' if a1<a0 else '相同')])
    development_summary=read_csv(OUT/'development_summary.csv')
    chosen=next(r for r in development_summary if r['profile']==f['profile']['name'])
    settings=[['法向量半徑','Park維持20cm；ETH維持10cm','沿用前輪密度適配設定，本輪比較不動這項。'],
              ['法向量權重','0.10 → 0.02','降低稀疏資料中不穩定法向量對圓柱擬合的影響。'],
              ['圓柱貼合容差','2cm → 5cm','允許較大的點到圓柱殘差；配準成功門檻保持原樣。'],
              ['圓弧覆蓋門檻','100° → 60°','保留較窄、較稀疏的圓柱片段；本方案仍用原本佔格量法。'],
              ['法向一致性比例','75% → 50%','接受較多法向量有噪聲的候選；也可能引入錯誤候選。'],
              ['何時啟用','點間距中位數 >5cm','只在稀疏資料啟用上述四項修改；密集資料保留原圓柱設定。'],
              ['未改的設定','step、分群20cm、垂直度0.90、匹配10cm、三維補救','成功標準仍為旋轉<2°、平移<0.5m；只使用兩站，無ICP或4DOF。']]
    headline=[['主要結論','在本輪保留站點只補回2次，仍有118次失敗；不足以採用為可靠公園配準方案。'],
              ['開發組結果',f'0/75 → {chosen["success"]}/75；用來選參數，不能當作保留資料的驗證。'],
              ['Park保留站點',f'{baseline_n}/120 → {selected_n}/120（{selected_n/120:.1%}）；補回{len(gained)}次、退步{len(lost)}次。'],
              ['改善分布','只在19→18配對的兩個seed成功，其餘11個配對仍沒有成功。'],
              ['ETH回歸',f'{eth_n}/150成功；180份CSV逐位元組相同，300列估計變換與誤差保持原值。'],
              ['本輪資料分工','前16站開發、後16站保留驗證；跨兩組的4個配對排除。保留組使用seed31–40。'],
              ['保留資料的限制','Park全部站點先前已做過整體測試；本輪不再用保留站點選參數，但它不是全新森林。'],
              ['角度新量法是否採用','沒有。90%角度跨度量法在開發組0/75；選出的是放寬稀疏擬合條件的組合。'],
              ['候選不是樹木真值','圖上可見直立弧形候選，但XYZ幾何不足以確認樹、柱或其他物件；本資料沒有樹木標註。'],
              ['建議','保留原ETH推薦方法；公園版本繼續視為研究方案，下一步需改進稀疏柱狀地標提取或取得較密集原始森林資料。']]
    simple=Workbook();simple.remove(simple.active)
    table(simple,'先看結果',['項目','白話說明'],headline,[30,145])
    table(simple,'改了哪些參數',['項目','調整','理由與代價'],settings,[29,62,100])
    table(simple,'保留站點各配對',['配對','對照成功／10','新方案成功／10','差值','觀察'],pairrows,[23,27,27,22,25])
    simple.save(OUT/'comparison_simple.xlsx')
    full=Workbook();full.remove(full.active)
    table(full,'總覽',['項目','內容'],headline,[30,145]);table(full,'參數說明',['項目','調整','理由'],settings,[30,65,100])
    csvtables=[('開發組候選比較',development_summary),('驗證與ETH摘要',summaries),('保留站點提取',counts),('密度判定',policies)]
    for name,rr in csvtables:
        keys=list(rr[0]);ws=table(full,name,keys,[[numeric(r[k]) for k in keys] for r in rr],[30]*len(keys))
        if 'rate' in keys:
            for i in range(2,ws.max_row+1):ws.cell(i,keys.index('rate')+1).number_format='0.0%'
    table(full,'保留各配對',['pair','baseline_success','selected_success','gain','result'],pairrows,[22,28,28,22,25])
    raw=[]
    for dataset,rr in [('WHU-Park development',development),('WHU-Park holdout',validation),('ETH regression',eth)]:raw.extend(dict(dataset=dataset,**r) for r in rr)
    keys=list(raw[0]);table(full,'完整逐次結果',keys,[[numeric(r[k]) for k in keys] for r in raw],[32 if k in ['dataset','profile','phase','label','solver','setting'] else 21 for k in keys])
    table(full,'ETH逐站回歸',list(regression['scans'][0]),[[numeric(r[k]) for k in regression['scans'][0]] for r in regression['scans']],[28]*5)
    meta=flatten(plan,'plan')+flatten(f,'frozen')+[
          ['GT使用','只在開發組選參數及事後評分；提取API與資料密度分流不讀GT。'],
          ['物理原點限制','Park仍採座標零原點，未獨立確認掃描器光心。'],['未執行e_p','開發組省略原始點雲誤差，空白不是0。保留驗證與ETH回歸使用全部source點雲。'],
          ['成功門檻','旋轉<2°且平移<0.5m，與前輪相同。'],['測試','新增4項角度跨度測試通過；原模式2seed的6份CSV與舊提取逐位元組一致；候選匯出已核對全部分群點數及重心。'],
          ['density政策欄位','sparse_span_used只標示是否換成角度跨度量法。本輪選定方案仍用佔格量法；稀疏擬合參數是否啟用，由median_nearest_m>0.05判定。']]
    table(full,'計畫與核驗',['項目','內容'],meta,[35,145]);full.save(OUT/'comparison_full.xlsx')
    checks=dict(development_cases_per_profile=75,profiles=7,holdout_cases_per_method=120,holdout_baseline_success=baseline_n,holdout_selected_success=selected_n,
                rescued_cases=[dict(seed=s,pair=p) for s,p in gained],regressed_cases=[dict(seed=s,pair=p) for s,p in lost],
                eth_success=eth_n,eth_extraction_files_identical=180,eth_registration_rows_unchanged=300,
                park_input_sha256_matches_previous_transfer=True,sparse_arc_unit_tests=4,
                files={n:dict(sha256=hashlib.sha256((OUT/n).read_bytes()).hexdigest(),bytes=(OUT/n).stat().st_size) for n in ['comparison_full.xlsx','comparison_simple.xlsx']})
    (OUT/'validation.json').write_text(json.dumps(checks,indent=2)+'\n')
    lines=['# 公園稀疏提取：分站開發與保留驗證','',
           '開發站0–15，共15個配對、seed21–25；保留站16–31，共12個配對、seed31–40；跨組4個配對排除。Park全場先前曾做整體測試，本輪分組只能視為本輪不參與調參的保留資料，不是全新森林。','',
           '## 候選檢視與方案','',
           '第0與第7站的候選群匯出後，全部分群點數及重心與舊紀錄一致。部分被拒絕群的俯視圖有圓弧、側視圖近直立；原佔格角度量法可能受稀疏度影響，但沒有語義樹木真值，不能斷言它們全是真樹。','',
           '比較7個事前列出的方案。只降低法向權重0/75、加5cm幾何容差1/75、只放寬覆蓋門檻0/75、90%角度跨度0/75；四項放寬的組合12/75最佳。新跨度量法沒有被選為最終方案。完整數字見development_summary.csv。','',
           '選定稀疏組合：法向量權重0.1→0.02；圓柱殘差容差2→5cm；佔格覆蓋100→60°；法向一致性75→50%。僅在離地高度過濾及voxel後的點間距中位數>5cm時啟用。密集輸入保留原擬合設定。Park法向半徑仍20cm、ETH仍10cm，兩者分群均20cm，未改配準求解與成功門檻。','',
           f'方案於{f["frozen_at_utc"]}固定，之後才產生保留組seed31–40；驗證結果出來後沒有重調參數。','',
           '## 結果與判斷','',
           f'保留組：對照{baseline_n}/120，新方案{selected_n}/120；補回{len(gained)}、退步{len(lost)}。只在19→18配對兩個seed成功；開發組12/75未穩定延伸到保留站點，不能把本方案視為可靠泛化。','',
           f'ETH回歸：原21–30的180份樹位／分群／保留率CSV逐位元組一致；重新執行配準後300列的變換、匹配與誤差仍與原值一致，主要solver維持{eth_n}/150。這確認密度分流沒有破壞原ETH行為，並不表示Park已解決。','',
           '目前不建議以新組合取代既有ETH推薦方案。公園提取仍缺乏穩定且足夠的共同柱狀候選；下一步需換提取策略或取得原始密度的森林分站資料，並建立新的獨立驗證。','',
           '## 重現與檔案','',
           'sparse_stem_study.py screen只讀開發組並保存frozen.json；validate讀保留站點；eth核對既有ETH21–30；最後執行report_sparse_stems.py。新功能在cm_sparse_extract中，原cm_extract_batch、核心預設及cm_register均未改動。','',
           '[精簡版](comparison_simple.xlsx)含結果、調參說明與保留配對；[完整版](comparison_full.xlsx)另含1050筆開發、480筆保留、300筆ETH結果與核驗。','',
           '![候選幾何檢視](candidate_audit.png)']
    (OUT/'README.md').write_text('\n'.join(lines)+'\n')
    print(json.dumps(checks,indent=2))


if __name__=='__main__':main()
