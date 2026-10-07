#!/usr/bin/env python3
"""Excel workbook comparing the four P(d) (distance-dependent down-sampling) versions for one fixed
configuration (improved + cluster.tolerance 0.20 m), official 10 pairs and 15 pairs, 10 seeds.

All statistics are Excel formulas over the raw sheet. Inputs:
  results/eth_trees/dtm_slope/retention_registration_ctol20.csv   (compare_retention.py)
  results/eth_trees/dtm_slope/seeds/ctol20/runs/<scan>/improved/summary_{raw,matched}.csv  (cm_compare)
  python3 scripts/eth_trees/make_retention_xlsx.py [--out FILE]
"""
import argparse
import csv
import itertools
import math
import os

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.formula import ArrayFormula

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
RAW = os.path.join(ROOT, "data/eth_trees/raw/trees")
RUNS = os.path.join(ROOT, "results/eth_trees/dtm_slope/seeds/ctol20/runs")
SCANS = ["s1", "s2", "s3", "s4", "s5", "s6"]
SEEDS = range(1, 11)
# key in data, display name, formula, origin, (summary file, column prefix) for kept points
VERS = [
    ("step_raw", "step", "d < 5 m: 0.5；5 ≤ d < 10 m: 0.75；d ≥ 10 m: 1", "論文式 (5)", ("raw", "step")),
    ("linear_a_raw", "linear_a", "min(1, 0.5 + 0.05·d)", "我們的修改", ("raw", "linear_a")),
    ("linear_mid_raw", "linear_mid", "min(1, 0.375 + 0.05·d)", "我們的修改", ("raw", "linear_mid")),
    ("physical_raw", "physical", "clip((d / 10 m)², 0.1, 1)", "我們的修改", ("raw", "physical")),
    ("linear_matched", "linear（同點數）", "linear 形狀，截距調整到保留點數 = step", "我們的修改（點數對齊）",
     ("matched", "linear*")),
    ("physical_matched", "physical（同點數）", "physical 形狀，d0 調整到保留點數 = step", "我們的修改（點數對齊）",
     ("matched", "physical*")),
]

FONT = "Arial"
F = Font(name=FONT, size=10)
FB = Font(name=FONT, size=10, bold=True)
FH = Font(name=FONT, size=10, bold=True, color="FFFFFF")
FT = Font(name=FONT, size=14, bold=True)
FBLUE = Font(name=FONT, size=10, color="0000FF")
FILL_H = PatternFill("solid", fgColor="1F4E78")
FILL_SUB = PatternFill("solid", fgColor="D9E1F2")
FILL_BEST = PatternFill("solid", fgColor="E2EFDA")
FILL_PAPER = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def header(ws, row, values, col=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.font, c.fill, c.alignment, c.border = FH, FILL_H, CENTER, BOX


def cell(ws, row, col, value, fmt=None, font=F, align=CENTER, fill=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font, c.alignment, c.border = font, align, BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(ROOT, "results/eth_trees/dtm_slope/retention_comparison.xlsx"))
    a = ap.parse_args()
    official = {tuple(l.split()) for l in open(os.path.join(RAW, "pairs.txt")) if l.strip()}
    pairs = list(itertools.combinations(SCANS, 2))
    rows = [r for r in csv.DictReader(open(os.path.join(ROOT, "results/eth_trees/dtm_slope/retention_registration_ctol20.csv")))
            if r["solver"] == "svd"]
    order = {v[0]: i for i, v in enumerate(VERS)}
    rows.sort(key=lambda r: (order[r["version"]], int(r["seed"]), r["pair"]))

    wb = Workbook()
    # ---------------- 原始資料 ----------------
    wr = wb.active
    wr.title = "原始資料"
    cols = ["P(d) 版本", "seed", "pair", "官方 10 對", "src 樹位數", "tgt 樹位數", "CN 配對數", "正確配對數",
            "求得轉換", "旋轉誤差 [°]", "平移誤差 [m]", "e_p [m]", "成功"]
    header(wr, 1, cols)
    for i, r in enumerate(rows, start=2):
        vals = [r["version"], int(r["seed"]), r["pair"], int(r["official"]), int(r["n_src_kp"]), int(r["n_tgt_kp"]),
                int(r["n_matches"]), int(r["n_correct"]), int(r["transform_ok"]),
                float(r["rot_err_deg"]) if r["transform_ok"] == "1" else None,
                float(r["trans_err_m"]) if r["transform_ok"] == "1" else None,
                float(r["e_p_m"]) if r["transform_ok"] == "1" else None, int(r["success"])]
        for j, v in enumerate(vals, start=1):
            cell(wr, i, j, v, fmt="0.000" if j in (10, 11, 12) else None)
    n = len(rows) + 1
    rng = lambda col: f"'原始資料'!${col}$2:${col}${n}"
    V, S, P, O, NM, NC, EP, OK = rng("A"), rng("B"), rng("C"), rng("D"), rng("G"), rng("H"), rng("L"), rng("M")
    for j, w in enumerate([18, 7, 9, 10, 10, 10, 10, 10, 9, 12, 12, 10, 7], start=1):
        wr.column_dimensions[L(j)].width = w
    wr.freeze_panes = "A2"
    wr.auto_filter.ref = f"A1:M{n}"

    # ---------------- 保留點數 ----------------
    wk = wb.create_sheet("保留點數")
    wk["A1"] = "各站保留點數與樹位數（10 個 seed 的平均，cm_compare summary_raw.csv / summary_matched.csv）"
    wk["A1"].font = FB
    header(wk, 3, ["站"] + [f"{v[1]} 保留點數" for v in VERS] + [f"{v[1]} 樹位數" for v in VERS])
    for i, s in enumerate(SCANS, start=4):
        cell(wk, i, 1, s)
        summ = {}
        for kind in ("raw", "matched"):
            for rr in csv.reader(open(os.path.join(RUNS, s, "improved", f"summary_{kind}.csv"))):
                summ[(kind, rr[0])] = rr
            hdr = next(csv.reader(open(os.path.join(RUNS, s, "improved", f"summary_{kind}.csv"))))
            summ[(kind, "_hdr")] = hdr
        for j, (_, _, _, _, (kind, pre)) in enumerate(VERS):
            hdr = summ[(kind, "_hdr")]
            k = hdr.index(f"{pre}_mean")
            cell(wk, i, 2 + j, float(summ[(kind, "kept points")][k]), fmt="#,##0", font=FBLUE)
            cell(wk, i, 2 + len(VERS) + j, float(summ[(kind, "trees")][k]), fmt="0.0", font=FBLUE)
    r_avg = 4 + len(SCANS)
    cell(wk, r_avg, 1, "平均", font=FB)
    for j in range(2 * len(VERS)):
        c = L(2 + j)
        cell(wk, r_avg, 2 + j, f"=AVERAGE({c}4:{c}{r_avg - 1})", fmt="#,##0" if j < len(VERS) else "0.0", font=FB)
    wk.cell(row=r_avg + 2, column=1, value="藍字：cm_compare 輸出的數值（10 個 seed 的平均）。樹位數是送進配準前的數量。").font = F
    wk.column_dimensions["A"].width = 8
    for j in range(2, 2 + 2 * len(VERS)):
        wk.column_dimensions[L(j)].width = 15

    # ---------------- 每 seed ----------------
    wsd = wb.create_sheet("每seed")
    wsd["A1"] = "每個 seed 的成功對數（COUNTIFS 自「原始資料」）"
    wsd["A1"].font = FB
    cell(wsd, 2, 1, "", fill=FILL_SUB)
    wsd.merge_cells(start_row=2, start_column=2, end_row=2, end_column=1 + len(VERS))
    cell(wsd, 2, 2, "官方 10 對", font=FB, fill=FILL_SUB)
    wsd.merge_cells(start_row=2, start_column=2 + len(VERS), end_row=2, end_column=1 + 2 * len(VERS))
    cell(wsd, 2, 2 + len(VERS), "15 對", font=FB, fill=FILL_SUB)
    header(wsd, 3, ["seed"] + [v[1] for v in VERS] * 2)
    for i, sd in enumerate(SEEDS, start=4):
        cell(wsd, i, 1, sd)
        for j, v in enumerate(VERS):
            cell(wsd, i, 2 + j, f'=COUNTIFS({V},"{v[0]}",{S},$A{i},{O},1,{OK},1)')
            cell(wsd, i, 2 + len(VERS) + j, f'=COUNTIFS({V},"{v[0]}",{S},$A{i},{OK},1)')
    seed_first, seed_last = 4, 3 + len(SEEDS)
    wsd.column_dimensions["A"].width = 7
    for j in range(2, 2 + 2 * len(VERS)):
        wsd.column_dimensions[L(j)].width = 14

    # ---------------- 成對 ----------------
    wp = wb.create_sheet("成對")
    wp["A1"] = "同一 seed、同一對的成功與否（1 = 成功）。用於「勝 / 輸 step」與符號檢定。"
    wp["A1"].font = FB
    header(wp, 2, ["seed", "pair", "官方 10 對"] + [v[1] for v in VERS])
    pr = 3
    for sd in SEEDS:
        for sa, sb in pairs:
            cell(wp, pr, 1, sd)
            cell(wp, pr, 2, f"{sa}-{sb}")
            cell(wp, pr, 3, int((sa, sb) in official))
            for j, v in enumerate(VERS):
                cell(wp, pr, 4 + j, f'=SUMIFS({OK},{V},"{v[0]}",{S},$A{pr},{P},$B{pr})')
            pr += 1
    p_last = pr - 1
    for j, w in enumerate([7, 9, 10] + [14] * len(VERS), start=1):
        wp.column_dimensions[L(j)].width = w
    wp.freeze_panes = "D3"

    # ---------------- 每一對 ----------------
    wq = wb.create_sheet("每一對")
    wq["A1"] = "每一對在 10 個 seed 中成功的次數"
    wq["A1"].font = FB
    header(wq, 2, ["pair", "官方 10 對", "兩站距離 [m]"] + [v[1] for v in VERS])
    for i, (sa, sb) in enumerate(pairs, start=3):
        T = [[float(x) for x in l.split()] for l in open(os.path.join(RAW, "groundtruth", f"{sa}-{sb}.tfm")) if l.strip()]
        cell(wq, i, 1, f"{sa}-{sb}")
        cell(wq, i, 2, "是" if (sa, sb) in official else "否")
        cell(wq, i, 3, round(math.hypot(T[0][3], T[1][3]), 1), fmt="0.0", font=FBLUE)
        for j, v in enumerate(VERS):
            cell(wq, i, 4 + j, f'=COUNTIFS({V},"{v[0]}",{P},$A{i},{OK},1)')
    last_q = 2 + len(pairs)
    cell(wq, last_q + 1, 1, "合計", font=FB)
    for j in range(len(VERS)):
        c = L(4 + j)
        cell(wq, last_q + 1, 4 + j, f"=SUM({c}3:{c}{last_q})", font=FB)
    wq.cell(row=last_q + 3, column=1, value="兩站距離：GT 轉換的水平平移量（藍字，自 groundtruth/*.tfm 計算）。").font = F
    for j, w in enumerate([9, 10, 12] + [14] * len(VERS), start=1):
        wq.column_dimensions[L(j)].width = w

    # ---------------- 摘要 ----------------
    ws = wb.create_sheet("摘要", 0)
    ws["A1"] = "距離降採樣 P(d) 四種版本的配準比較（ETH Trees）"
    ws["A1"].font = FT
    ws["A2"] = ("固定組合：improved（5 個擬合後檢查全開）+ clustering 距離 0.20 m；新 DTM（坡度濾波）；"
                "配準為論文設定（閾值 5 cm、N = 183、w = 10、Th = 40），SVD 求轉換；10 個 seed。")
    ws["A2"].font = F
    ws.merge_cells("A2:V2")
    hdr1 = [("", 1, 5), ("官方 10 對（pairs.txt）", 6, 10), ("15 對（6 站兩兩）", 11, 15), ("精度（15 對，成功的對）", 16, 18),
            ("與 step 成對比較（15 對 × 10 seed）", 19, 22)]
    for text, c1, c2 in hdr1:
        if c2 > c1:
            ws.merge_cells(start_row=4, start_column=c1, end_row=4, end_column=c2)
        for c in range(c1, c2 + 1):
            cell(ws, 4, c, text if c == c1 else None, font=FB, fill=FILL_SUB)
    header(ws, 5, ["P(d) 版本", "公式（d = 水平距離）", "來源", "平均保留點數 / 站", "平均樹位數 / 站",
                   "成功對數 平均", "標準差", "最少", "最多", "成功率",
                   "成功對數 平均", "標準差", "最少", "最多", "成功率",
                   "e_p 平均 [m]", "e_p 中位數 [m]", "正確配對比例",
                   "勝 step", "輸 step", "符號檢定 p", "判讀"])
    for i, v in enumerate(VERS):
        r = 6 + i
        sc10, sc15 = L(2 + i), L(2 + len(VERS) + i)
        r10 = f"'每seed'!{sc10}{seed_first}:{sc10}{seed_last}"
        r15 = f"'每seed'!{sc15}{seed_first}:{sc15}{seed_last}"
        kp, tr = L(2 + i), L(2 + len(VERS) + i)
        pc = L(4 + i)
        cell(ws, r, 1, v[1], font=FB, align=LEFT)
        cell(ws, r, 2, v[2], align=LEFT)
        cell(ws, r, 3, v[3])
        cell(ws, r, 4, f"='保留點數'!{kp}{r_avg}", fmt="#,##0")
        cell(ws, r, 5, f"='保留點數'!{tr}{r_avg}", fmt="0.0")
        for base, rr_, nn in ((6, r10, 10), (11, r15, 15)):
            cell(ws, r, base, f"=AVERAGE({rr_})", fmt="0.0")
            cell(ws, r, base + 1, f"=STDEV({rr_})", fmt="0.0")
            cell(ws, r, base + 2, f"=MIN({rr_})", fmt="0")
            cell(ws, r, base + 3, f"=MAX({rr_})", fmt="0")
            cell(ws, r, base + 4, f"={L(base)}{r}/{nn}", fmt="0%")
        cell(ws, r, 16, f'=AVERAGEIFS({EP},{V},"{v[0]}",{OK},1)', fmt="0.000")
        ws[f"Q{r}"] = ArrayFormula(f"Q{r}", f'=MEDIAN(IF(({V}="{v[0]}")*({OK}=1),{EP}))')
        cell(ws, r, 17, ws[f"Q{r}"].value, fmt="0.000")
        cell(ws, r, 18, f'=SUMIFS({NC},{V},"{v[0]}")/SUMIFS({NM},{V},"{v[0]}")', fmt="0.00")
        if i == 0:
            for c in (19, 20, 21):
                cell(ws, r, c, "–")
            cell(ws, r, 22, "基準（論文）", align=LEFT)
        else:
            cell(ws, r, 19, f"=SUMPRODUCT(('成對'!{pc}3:{pc}{p_last}=1)*('成對'!D3:D{p_last}=0))")
            cell(ws, r, 20, f"=SUMPRODUCT(('成對'!{pc}3:{pc}{p_last}=0)*('成對'!D3:D{p_last}=1))")
            cell(ws, r, 21, f'=IF(S{r}+T{r}=0,1,MIN(1,2*BINOMDIST(MIN(S{r},T{r}),S{r}+T{r},0.5,TRUE)))', fmt="0.000")
            cell(ws, r, 22, f'=IF(U{r}<0.05/5,IF(S{r}>T{r},"顯著較好","顯著較差"),"無顯著差異")', align=LEFT)
    last = 5 + len(VERS)
    # paper reference row
    rp = last + 1
    cell(ws, rp, 1, "論文 CircularMatch", font=FB, align=LEFT, fill=FILL_PAPER)
    cell(ws, rp, 2, "step（論文式 5）；論文自己的樹位擷取", align=LEFT, fill=FILL_PAPER)
    cell(ws, rp, 3, "論文 Table I", fill=FILL_PAPER)
    for c in list(range(4, 15)) + [17, 18, 19, 20, 21, 22]:
        cell(ws, rp, c, "–", fill=FILL_PAPER)
    c = cell(ws, rp, 15, 1.0, fmt="0%", font=FBLUE, fill=FILL_PAPER)
    c.comment = Comment("論文 Table I，ETH-tree，Ours：rate 100%。論文沒寫 ETH 用了幾對（推測 15 或 30 對）。", "Claude")
    c = cell(ws, rp, 16, 0.153427, fmt="0.000", font=FBLUE, fill=FILL_PAPER)
    c.comment = Comment("論文 Table I，ETH-tree，Ours：e_p = 0.153427 m。論文沒說是平均或中位數、是否只算成功的對。", "Claude")
    # highlight best success rate per pair set
    from openpyxl.formatting.rule import FormulaRule
    for col in ("J", "O"):
        ws.conditional_formatting.add(f"{col}6:{col}{last}",
                                      FormulaRule(formula=[f"{col}6=MAX(${col}$6:${col}${last})"], fill=FILL_BEST))
    notes = [
        "說明",
        "• 成功：旋轉誤差 < 2° 且平移誤差 < 0.5 m（論文沒有給成功標準，這是我們定的）。",
        "• 成功對數：每個 seed 分別計算，再取 10 個 seed 的平均、標準差、最少、最多（見「每seed」）。",
        "• e_p（論文式 8）：對 source 掃描的所有點計算，這裡只平均成功的對；論文的 e_p 統計方式未說明，不能直接比較。",
        "• 勝 / 輸 step：同一 seed、同一對（共用隨機數 CRN），該版本成功而 step 失敗為「勝」，反之為「輸」（見「成對」）。",
        "• 符號檢定 p：雙尾二項檢定；5 個版本各與 step 比一次，判讀用 Bonferroni 門檻 0.05 / 5 = 0.01。"
        "同一對在不同 seed 之間不獨立，p 值偏樂觀。",
        "• step 是論文式 (5)；其他版本是我們的修改。「同點數」版本把參數調到保留點數與 step 相同，用來分開「點數」與「曲線形狀」的影響。",
        "• 固定組合中「論文沒有」的部分：半徑、傾角、inlier 比例、法向一致性、圓弧覆蓋角 5 個擬合後檢查；"
        "clustering 距離 0.20 m（論文未給值）與 DTM 方法（論文未指定）是實作選擇。此組合是在 ETH 上選出的，沒有獨立驗證資料。",
        "• 綠底：該欄成功率最高。黃底：論文數字（藍字為直接抄自論文）。",
    ]
    for k, t in enumerate(notes):
        c = ws.cell(row=rp + 2 + k, column=1, value=t)
        c.font = FB if k == 0 else F
        ws.merge_cells(start_row=rp + 2 + k, start_column=1, end_row=rp + 2 + k, end_column=22)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[rp + 2 + k].height = 28 if len(t) > 90 else 16
    widths = [16, 34, 14, 13, 11, 10, 8, 6, 6, 8, 10, 8, 6, 6, 8, 10, 11, 10, 8, 8, 10, 12]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[L(j)].width = w
    ws.row_dimensions[5].height = 30
    ws.freeze_panes = "B6"
    for sh in wb.worksheets:  # print: landscape, one page wide
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth, sh.page_setup.fitToHeight = 1, 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(a.out)
    print("wrote", a.out)


if __name__ == "__main__":
    main()
