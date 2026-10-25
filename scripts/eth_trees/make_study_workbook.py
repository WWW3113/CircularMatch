#!/usr/bin/env python3
"""Plain, auditable workbook; unavailable datasets never become zero scores.

Dataset registry: results/cross_dataset_20261025/datasets.json. Add normalized
fresh_rows/graph_rows paths only after data provenance and execution are checked.
Output starts in ignored runs/input; recalculate it with Calc into the report dir.
"""
import csv
import json
import math
from pathlib import Path
import statistics as st
import subprocess
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter as letter

from tuning_graph import sign_p

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'results/cross_dataset_20261025'
ETH=ROOT/'results/eth_trees/tuning_20261025'
LIANG=ROOT/'results/liang_trees/transfer_20261025'


def load(path):
    with Path(path).open() as f:return list(csv.DictReader(f))


def number(value):
    try:
        n=float(value);return n if math.isfinite(n) else None
    except (TypeError,ValueError):return None


def table(wb,name,headers,rows,widths=None):
    ws=wb.create_sheet(name);ws.append(headers)
    for row in rows:ws.append(row)
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False
    for c in ws[1]:
        c.font=Font(name='Arial',size=10,bold=True)
        c.fill=PatternFill('solid',fgColor='E7E6E6')
        c.border=Border(bottom=Side(style='thin',color='808080'))
        c.alignment=Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[1].height=32
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font=Font(name='Arial',size=10);c.alignment=Alignment(vertical='top')
            if isinstance(c.value,float):c.number_format='0.000'
    for i in range(1,len(headers)+1):ws.column_dimensions[letter(i)].width=(widths[i-1] if widths else 18)
    for row in ws.iter_rows(min_row=2):
        lines=1
        for c in row:
            if isinstance(c.value,str) and not c.value.startswith('='):
                visual=sum(2 if unicodedata.east_asian_width(ch) in ('W','F') else 1 for ch in c.value)
                count=max(1,math.ceil(visual/max(1,ws.column_dimensions[c.column_letter].width-2)))
                if count>1:c.alignment=Alignment(vertical='top',wrap_text=True)
                lines=max(lines,count)
        if lines>1:ws.row_dimensions[row[0].row].height=min(120,15*lines+3)
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows='1:1'
    return ws


def main():
    datasets=json.loads((OUT/'datasets.json').read_text());frozen=json.loads((ETH/'frozen_matching.json').read_text())
    entries=[];official={};data_by_id={d['id']:d for d in datasets}
    for d in datasets:
        if not d['available']:continue
        official[d['id']]={tuple(x.split()) for x in (ROOT/d['official_pairs']).read_text().splitlines() if x.strip()}
        for path in d['rows']:
            for r in load(ROOT/path):
                r=dict(r);r['dataset']=d['id'];r['scene']=d.get('scene','');r['source']=path
                r['seed']=int(r['seed']);r['ok']=int(r['transform_ok']);r['success']=int(r['success'])
                r['mode']='多站' if r['method']=='graph_mst' else '兩站'
                r['official']=int(tuple(r['pair'].split('-')) in official[d['id']])
                entries.append(r)
    assert entries,'No completed dataset results'
    identity=lambda r:(r['dataset'],r['method'],r['solver'],r['seed'],r['pair'])
    assert len({identity(r) for r in entries})==len(entries),'duplicate cases'
    methods=list(dict.fromkeys((r['method'],r['solver']) for r in entries))
    wb=Workbook();wb.remove(wb.active)
    completed=[d for d in datasets if d['available']]
    families={d['family'] for d in completed}
    external=load(LIANG/'extraction_rows.csv') if (LIANG/'extraction_rows.csv').exists() else []
    overview=[['已完成配準資料集數',len(families)],['已完成配準資料集','、'.join(d['id'] for d in completed)],
              ['跨資料集配準狀態','已完成' if len(families)>1 else '尚未完成：其他原始站點與配準真值待取得'],
              ['外部樹木提取驗證','Liang2018 / TreeTool處理後六樣區、三設定、seed11–20；見外部提取工作表' if external else '尚未執行'],
              ['本輪ETH的seed','11–20；篩選使用舊seed1–5，6–10另列'],
              ['成功定義','旋轉誤差 < 2° 且平移誤差 < 0.5 m'],
              ['兩站方法','每次只用兩站；控制組5 cm＋SVD，候選為固定10 cm＋配準RANSAC'],
              ['多站方法','最大生成樹使用同場景全部站；與獨立兩站任務分開解讀'],
              ['e_p定義','source全部點的平均變換差；表內中位數只對成功案例'],
              ['新資料策略','使用ETH固定參數直接套用，不先針對新資料重調'],
              ['ETH建議','獨立兩站優先10cm＋配準RANSAC：142/150 (94.7%)；control SVD為118/150 (78.7%)'],
              ['多站路徑建議','六站最大生成樹150/150；需取得全部站點，不與獨立兩站作相同條件宣稱'],
              ['資料缺漏','未下載／未執行列明確標示；不是0%成功率'],
              ['Git基底',subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()],
              ['Git分支',subprocess.check_output(['git','branch','--show-current'],cwd=ROOT,text=True).strip()]]
    if external:
        overview += [['外部提取結論','六樣區等權參考匹配比例：control 38.34%、vertical085 34.94%、ransac999 38.52%。保留垂直度0.90；增加圓柱擬合資源的改善有限。'],
                     ['後續優先順序','取得原始分站資料與配準真值，驗證固定10cm＋配準RANSAC；提取先檢查漏樹原因、感測器原點與標註範圍。']]
    ws=table(wb,'說明',['項目','內容'],overview,[26,110])
    for row in ws.iter_rows(min_row=2):row[1].alignment=Alignment(wrap_text=True,vertical='top')
    raw=[];poses=[]
    for r in entries:
        a,b=r['pair'].split('-');m=r['method']
        kind='vertical085' if 'vertical085' in m else ('ransac999' if 'ransac999' in m else 'control')
        raw.append([r['dataset'],r['scene'],r['mode'],m,r['solver'],r['seed'],r['pair'],r['official'],r['ok'],r['success'],
                    number(r.get('rot_err_deg')),number(r.get('trans_err_m')),number(r.get('e_p_m')),
                    number(r.get('n_matches')),number(r.get('n_correct_matches')),number(r.get('ransac_inliers')),
                    r.get('path',''),number(r.get('time_ms')),kind,r['source']])
        poses.append([r['dataset'],m,r['solver'],r['seed'],a,b,r['ok']]+[number(r.get(k)) if r['ok'] else None for k in [f'r{i}{j}' for i in range(3) for j in range(3)]+[f't{i}' for i in range(3)]])
    wr=table(wb,'逐對結果',['資料集','場景','任務','方法代碼','solver','seed','pair','官方配對','求得變換','成功','旋轉誤差°','平移誤差m','e_p m','CN匹配數','GT正確匹配數(僅評分)','RANSAC inliers','多站路徑','匹配＋求解ms','擷取設定','結果來源'],raw,[18,18,8,35,17,8,13,11,11,9,15,15,15,14,22,17,30,17,18,70])
    table(wb,'估計變換',['資料集','方法代碼','solver','seed','source','target','求得變換']+[f'r{i}{j}' for i in range(3) for j in range(3)]+[f't{i}' for i in range(3)],poses,[18,35,17,8,10,10,11]+[15]*12)
    last=len(raw)+1
    rg=lambda col:f"'逐對結果'!${col}$2:${col}${last}"
    summaries=[];paired=[];seed_rows=[];group_stats={}
    for d in completed:
        dr=[r for r in entries if r['dataset']==d['id']]
        base={(r['seed'],r['pair']):r for r in dr if r['method']=='fresh_control' and r['solver']=='svd'}
        for method,solver in methods:
            rr=[r for r in dr if r['method']==method and r['solver']==solver]
            if not rr:continue
            seeds=sorted({r['seed'] for r in rr});counts=[];baseline=[]
            for seed in seeds:
                sr=[r for r in rr if r['seed']==seed];counts.append(sum(r['success'] for r in sr));baseline.append(sum(v['success'] for k,v in base.items() if k[0]==seed))
                seed_rows.append([d['id'],method,solver,seed,len(sr),counts[-1],counts[-1]/len(sr)])
            ep=[float(r['e_p_m']) for r in rr if r['success']]
            mode=rr[0]['mode'];row=len(summaries)+2
            crit=f'{rg("A")},$A{row},{rg("D")},$C{row},{rg("E")},$D{row}'
            summaries.append([d['id'],mode,method,solver,f'=COUNTIFS({crit})',f'=COUNTIFS({crit},{rg("J")},1)',f'=IF(E{row}=0,"",F{row}/E{row})',
                              f'=COUNTIFS({crit},{rg("H")},1)',f'=COUNTIFS({crit},{rg("H")},1,{rg("J")},1)',f'=IF(H{row}=0,"",I{row}/H{row})',
                              st.median(ep) if ep else None,st.stdev(counts) if len(counts)>1 else None,'完成',
                              st.median(float(r['time_ms']) for r in rr if r.get('time_ms')) if any(r.get('time_ms') for r in rr) else None])
            wins=sum(a>b for a,b in zip(counts,baseline));losses=sum(a<b for a,b in zip(counts,baseline))
            common=[(r,base[r['seed'],r['pair']]) for r in rr if r['success'] and base[r['seed'],r['pair']]['success']]
            comparisons=len(methods)-1;p=sign_p(wins,losses)
            paired.append([d['id'],method,solver,st.mean(a-b for a,b in zip(counts,baseline)),wins,losses,len(seeds)-wins-losses,p,min(1.,p*comparisons),len(common),
                           st.median(float(b['e_p_m']) for _,b in common) if common else None,st.median(float(a['e_p_m']) for a,_ in common) if common else None])
            group_stats[d['id'],method,solver]=dict(n=len(rr),success=sum(r['success'] for r in rr),row=row)
    ws=table(wb,'資料集摘要',['資料集','任務','方法代碼','solver','案例數','成功數','成功率','官方案例數','官方成功數','官方成功率','成功e_p中位數m','每seed成功對數SD','狀態','匹配＋求解ms中位數'],summaries,[18,8,35,17,12,12,13,14,14,14,19,21,10,26])
    for row in range(2,ws.max_row+1):
        for col in [7,10]:ws.cell(row,col).number_format='0.0%'
    ws=table(wb,'每seed',['資料集','方法代碼','solver','seed','案例數','成功數','成功率'],seed_rows,[18,35,17,10,14,14,14])
    for row in range(2,ws.max_row+1):ws.cell(row,7).number_format='0.0%'
    ws=table(wb,'配對比較',['資料集','方法代碼','solver','Δ成功對數/seed','勝seed數','負seed數','相同seed數','符號檢定p','Bonferroni p','共同成功案例','control共同成功e_p','此方法共同成功e_p'],paired,[18,35,17,20,14,14,16,17,18,18,24,24])
    for row in range(2,ws.max_row+1):
        for col in [8,9]:ws.cell(row,col).number_format='0.0000'
    cross=[]
    for method,solver in methods:
        mode='多站' if method=='graph_mst' else '兩站';row=[mode,method,solver]
        for d in datasets:
            stat=group_stats.get((d['id'],method,solver))
            row.append(f"='資料集摘要'!G{stat['row']}" if stat else '未執行／待取得資料')
        cross.append(row)
    ws=table(wb,'跨資料集',['任務','方法代碼','solver']+[d['id'] for d in datasets],cross,[8,35,17]+[25]*len(datasets))
    for row in range(2,ws.max_row+1):
        for col in range(4,ws.max_column+1):ws.cell(row,col).number_format='0.0%'
    params=[]
    for method,solver in methods:
        vert=.85 if 'vertical085' in method else .90;prob=.999 if 'ransac999' in method else .99;iters=3000 if 'ransac999' in method else 1000
        threshold=frozen['profile']['threshold'] if method=='fresh_matching' or method.endswith('_matching') else .05
        params.append([method,solver,'多站' if method=='graph_mst' else '兩站',threshold,183,10,40,3,.20,vert,prob,iters,.30 if solver=='svd+ransac' else None,
                       '至少4個CN配對、6個3D支持；互為最近鄰距離<0.30m，中位殘差<=0.15m；不讀GT選邊' if method=='graph_mst' else '完整6DOF；不做ICP'])
    table(wb,'方法設定',['方法代碼','solver','任務','匹配閾值m','N','w','Th_score','min_dc','cluster m','verticality','圓柱prob','圓柱迭代上限','配準RANSAC inlier m','補充'],params,[35,17,8,15,9,9,13,12,14,14,14,18,25,95])
    sources=[]
    names={'available':'已可執行配準比較','acquired':'已取得森林點雲','extraction_completed':'外部提取驗證完成','status':'狀態','scene':'場景／子集','scans':'已核驗站數','points':'已核驗點數',
           'source':'來源頁','download':'原始下載候選','processed_download':'已取得處理版下載','sha256':'本地核驗SHA256','reference':'引用','candidate_source':'補充來源','note':'注意'}
    for d in datasets:
        for key,label in names.items():
            if key in d:sources.append([d['id'],label,d[key]])
    ws=table(wb,'資料來源',['資料集','項目','內容'],sources,[18,25,115])
    for row in ws.iter_rows(min_row=2):
        row[2].alignment=Alignment(wrap_text=True,vertical='top')
        if len(str(row[2].value))>100:ws.row_dimensions[row[0].row].height=36
    screen=load(ETH/'screen_summary.csv');profiles={p['name']:p for p in json.loads((ETH/'plan.json').read_text())['matching_profiles']}
    table(wb,'ETH篩選',['設定','solver','seed1–5成功','seed1–5案例','seed6–10成功','seed6–10案例','參數'],[[r['method'],r['solver'],int(r['discovery_success']),75,int(r['existing_seed_check_success']),75,json.dumps(profiles[r['method']],ensure_ascii=False)] for r in screen],[22,17,17,17,17,17,95])
    commands=[json.loads(line) for line in (ETH/'runs/commands.jsonl').read_text().splitlines()]
    elapsed={r['label']:r['seconds'] for r in commands if r['exit_code']==0}
    extraction=load(ETH/'extraction_summary.csv')
    table(wb,'ETH提取',['設定','站點','seed數','平均樹數','樹數SD','10seed批次壁鐘秒','計時條件'],
          [[r['variant'],r['scan'],int(r['seeds']),float(r['mean_trees']),float(r['sd_trees']),
            elapsed.get(f'extract_{r["variant"]}_{r["scan"]}'),'共享前處理、同時有其他工作；非獨占主機效能基準'] for r in extraction],
          [22,12,13,18,18,25,90])
    edges=load(ETH/'runs/graph_edges.csv')
    table(wb,'ETH圖選邊',['seed','pair','採用','品質score','3D支持點','overlap','中位殘差m'],[[int(r['seed']),r['pair'],int(r['selected']),number(r['score']),int(r['support']),number(r['overlap']),number(r['median_residual'])] for r in edges],[12,16,12,17,17,16,19])
    validation=json.loads((ETH/'graph_validation.json').read_text())
    if external:
        assert len(external)==180
        erows=[[r['dataset'],int(r['plot']),int(r['plot']) in [1,3,5],r['variant'],int(r['seed']),
                int(r['n_reference']),int(r['n_detected']),int(r['n_matched']),float(r['reference_match_rate']),
                number(r['matched_xy_rmse_m']),number(r['matched_xy_median_m'])] for r in external]
        ws=table(wb,'外部提取逐次',['資料版本','Plot','論文使用樣區','提取設定','seed','參考樹數','偵測數','匹配數','參考匹配比例','匹配XY RMSE m','匹配XY中位數m'],erows,[25,10,18,22,10,15,15,15,20,22,22])
        for i in range(2,ws.max_row+1):ws.cell(i,9).number_format='0.0%'
        last_ext=len(erows)+1
        xr=lambda col:f"'外部提取逐次'!${col}$2:${col}${last_ext}"
        esummary=[]
        for kind in ['control','vertical085','ransac999']:
            for plot in range(1,7):
                row=len(esummary)+2
                crit=f'{xr("B")},$B{row},{xr("D")},$A{row}'
                rr=[r for r in external if r['variant']==kind and int(r['plot'])==plot]
                esummary.append([kind,plot,plot in [1,3,5],int(rr[0]['n_reference']),f'=COUNTIFS({crit})',
                                 f'=AVERAGEIFS({xr("G")},{crit})',f'=AVERAGEIFS({xr("H")},{crit})',
                                 f'=AVERAGEIFS({xr("I")},{crit})',st.stdev(float(r['reference_match_rate']) for r in rr),
                                 st.median(float(r['matched_xy_rmse_m']) for r in rr if r['matched_xy_rmse_m']) if any(r['matched_xy_rmse_m'] for r in rr) else None])
        ws=table(wb,'外部提取摘要',['設定','Plot','論文使用樣區','參考樹數','seed數','平均偵測數','平均匹配數','平均參考匹配比例','seed比例SD','每seed匹配RMSE中位數m'],esummary,[22,10,18,16,13,18,18,23,20,30])
        for i in range(2,ws.max_row+1):
            for col in [8,9]:ws.cell(i,col).number_format='0.0%'
        # Each plot gets equal weight; seed repeats are not independent forests.
        compare=[]
        for subset,plots in [('全部六樣區',list(range(1,7))),('論文Plot1/3/5',[1,3,5])]:
            baseline=st.mean(float(r['reference_match_rate']) for r in external if r['variant']=='control' and int(r['plot']) in plots)
            for kind in ['control','vertical085','ransac999']:
                rr=[r for r in external if r['variant']==kind and int(r['plot']) in plots]
                rate=st.mean(float(r['reference_match_rate']) for r in rr)
                compare.append([subset,kind,len(plots),10,rate,rate-baseline,'各樣區等權；固定0.5m一對一XY匹配，未調新資料參數'])
        ws=table(wb,'外部提取比較',['樣區集合','設定','樣區數','每樣區seed數','平均參考匹配比例','相對control差值','解讀'],compare,[24,22,13,19,25,23,90])
        for i in range(2,ws.max_row+1):
            for col in [5,6]:ws.cell(i,col).number_format='0.0%'
        plan=json.loads((LIANG/'plan.json').read_text())
        protocol=[['版本','TreeTool公開處理後PCD，每樣區一份；不是原始分站掃描'],
                  ['原點','沿用PCD VIEWPOINT零原點；角度格網檢查不能獨立確認真實掃描中心'],
                  ['匹配','參考檔欄位TreeID/X/Y/height/diameter；地面樹位與參考XY，一對一最大匹配數，閾值0.5m'],
                  ['主要指標','匹配參考樹 / 全部提供的參考樹；六樣區等權平均，10個seed是重複擬合'],
                  ['未報precision','標註範圍不明，未匹配偵測可能位於標註範圍之外，不能直接當誤報'],
                  ['位置誤差限制','參考XY錨點高度不明；只列匹配XY誤差，未評DBH/高度精度'],
                  ['跨資料集限制','只驗證提取階段；沒有分站配準真值，不能評10cm匹配參數或多站路徑的外部配準效果']]
        for name,digest in plan['input_sha256'].items():protocol.append([name,digest])
        ws=table(wb,'外部資料核驗',['項目','內容'],protocol,[52,125])
        for row in ws.iter_rows(min_row=2):row[1].alignment=Alignment(wrap_text=True,vertical='top')
    definitions=[['可作的結論','ETH新seed的配準改善；外部提取工作表列Liang處理版結果，跨資料集配準尚待原始站點及真值。'],
                 ['不能混用的基準','本輪control是seed11–20的78.7%；不是舊seed1–10的84%。'],
                 ['資料集與seed','多個seed重複同片森林，不等於多個獨立資料集。'],
                 ['多站比較','圖方法存取同場景全部站；兩站方法只存取單一配對。不可混稱同一輸入條件。'],
                 ['參數篩選','只用舊seed1–5決定10cm＋SVD/RANSAC；擷取與組合設定在新seed評分前指定。'],
                 ['p值','以每seed成功對數比較，非把150個配對視為獨立森林樣本；Bonferroni校正所有非control比較。'],
                 ['精度比較','成功e_p中位數可能受成功集合不同影響；配對比較另列雙方共同成功的案例。'],
                 ['真值使用','GT只用評分；圖路徑選邊只用估計變換後的3D支持點。'],
                 ['e_p核對',f'NumPy完整點雲與C++差異={validation["cpp_ep_reference_error"]:.3g}m'],
                 ['GT擾動核對',str(validation.get('cpp_transform_independent_of_evaluation_gt'))+'；評分GT平移100m，估計變換不變'],
                 ['批次擷取測試','2個seed、6份CSV逐位元組一致；180份新seed樹位數均與執行日誌一致'],
                 ['新增測試','4項Python測試通過，包含批次擷取一致性、空seed範圍、圖路徑正反向與3D支持'],
                 ['方法碼','svd+ransac指配準階段RANSAC；ransac999方法碼指擷取階段圓柱擬合，兩者不同'],
                 ['時間欄位','逐對結果的time_ms只含描述子、匹配與求解；不含完整點誤差計算和擷取時間']]
    ws=table(wb,'定義與核驗',['項目','說明'],definitions,[25,125])
    for row in ws.iter_rows(min_row=2):row[1].alignment=Alignment(wrap_text=True,vertical='top')
    target=OUT/'runs/input/comparison.xlsx';target.parent.mkdir(parents=True,exist_ok=True);wb.save(target)
    manifest=dict(rows=len(entries),methods=len(methods),completed_datasets=len(families),external_extraction_rows=len(external),sheets=wb.sheetnames,
                  expected_summary=[dict(dataset=k[0],method=k[1],solver=k[2],**v) for k,v in group_stats.items()])
    (OUT/'runs/workbook_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print(target);print(json.dumps({k:manifest[k] for k in ['rows','methods','completed_datasets','sheets']},ensure_ascii=False))


if __name__=='__main__':main()
