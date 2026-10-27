#!/usr/bin/env python3
"""Six compact professor reports. Only completed scalar experiments are read."""
from __future__ import annotations
import argparse,csv,hashlib,json,math,statistics as st,re,unicodedata
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import xlsxwriter
from openpyxl import load_workbook
from make_parameter_detail_reports import stats,finite,key,scene
from six_week_talks import later_talk

ROOT=Path(__file__).resolve().parents[2]
OLD=ROOT/'results/professor_reports_20261026'
OUT=ROOT/'results/six_week_reports_20261028'
TITLES={1:'CN距離容差與Solver',2:'弱配對補救',3:'DTM平滑與樹位高度',4:'加入幾何微調',5:'擴大候選搜尋的退步',6:'結構驗證與最終比較'}
SCENES=['Plot1','Plot2','Plot3','Plot4','ETH']
PAPER=dict(zip(SCENES,[.052558,.114808,.048756,.04356,.153427]))
NA='待補'
COMMON='成功門檻：e_R<2°且e_t<0.5m。e_p為Eq8全來源點變換誤差；finite mean包含可計算失敗，無變換仍算失敗、誤差不填0。'
HISTORY='GT不參與當次配準輸出選擇；開發期間曾以GT評分選設定。相同掃描站換seed只測隨機敏感度，不是未見森林。'

def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def cm(x):return None if x is None else x*100
def fmt(x,n=3):return NA if x is None else f'{x:.{n}f}'
def label(s):return 'ETH-tree' if s=='ETH' else 'Tongji／'+s
def pp(a,b):return 100*(b['rate']-a['rate'])

class Evidence:
 def __init__(self):self.tables={};self.sources={};self.groups={};self.paired={}
 def load(self,sid,file):
  p=OLD/'evidence'/file
  self.sources[sid]=dict(path=str(p.relative_to(ROOT)),sha256=sha(p))
  self.tables[sid]=list(csv.DictReader(p.open()))
 def load_path(self,sid,path):
  p=ROOT/path;archive=OUT/'evidence'/(sid+'_'+p.name);archive.parent.mkdir(parents=True,exist_ok=True)
  if not archive.exists():archive.write_bytes(p.read_bytes())
  self.sources[sid]=dict(path=str(archive.relative_to(ROOT)),sha256=sha(archive));self.tables[sid]=list(csv.DictReader(archive.open()))
 def group(self,gid,sid,**sel):
  if gid in self.groups:return self.groups[gid]
  inds=[n for n,r in enumerate(self.tables[sid]) if all((scene(r) if k=='scene' else r[k])==str(v) for k,v in sel.items())]
  rows=[self.tables[sid][n] for n in inds];g=dict(id=gid,source=sid,selector=sel,csv_lines=[n+2 for n in inds],rows=rows,**stats(rows))
  self.groups[gid]=g;return g
 def u(self,s,m):return self.group('U:'+s+':'+m,'U',scene=s,method=m)
 def pair(self,a,b):
  gid=a['id']+' -> '+b['id']
  if gid in self.paired:return self.paired[gid]
  aa={key(r):r for r in a['rows']};bb={key(r):r for r in b['rows']};assert aa.keys()==bb.keys()
  both=[k for k in aa if aa[k]['success']==bb[k]['success']=='1']
  def avg(side):
   v=[float(side[k]['e_p_m']) for k in both if finite(side[k]['e_p_m']) is not None];return st.mean(v) if v else None
  r=dict(gained=sum(aa[k]['success']=='0' and bb[k]['success']=='1' for k in aa),lost=sum(aa[k]['success']=='1' and bb[k]['success']=='0' for k in aa),common=len(both),before=avg(aa),after=avg(bb),
         same_finite={k for k in aa if finite(aa[k]['e_p_m']) is not None}=={k for k in bb if finite(bb[k]['e_p_m']) is not None})
  self.paired[gid]=r;return r

class Book:
 def __init__(self,i,e,sheets):
  self.i=i;self.e=e;self.path=OUT/f'report_{i:02d}_{TITLES[i]}.xlsx';self.wb=xlsxwriter.Workbook(self.path);self.checks=[];self.ranges={};self.chart_count=0
  self.formats={}
  for name,extra in {'title':{'bold':True,'font_size':18,'font_color':'243E52'},'head':{'bold':True,'bg_color':'E8EDF1'},'text':{},'num':{'num_format':'0.000'},'int':{'num_format':'0'},'pct':{'num_format':'0.0%'},'pp':{'num_format':'0.0" pp"'},'note':{'font_size':10,'font_color':'444444'}}.items():
   if name in ['num','int','pct','pp']:extra['align']='center'
   self.formats[name]=self.wb.add_format(dict(font_name='Microsoft JhengHei',font_size=13,text_wrap=True,valign='vcenter')|extra)
  self.sheets={};self.ends={}
  for title in sheets:
   ws=self.wb.add_worksheet(title);self.sheets[title]=ws;ws.hide_gridlines(2);ws.set_zoom(95);ws.freeze_panes(5,1)
   widths=[26,28,35,72,5,5,5,5] if i>1 and title=='設定與說明' else [27,21,20,20,18,20,20,30]
   for c,w in enumerate(widths):ws.set_column(c,c,w)
   self.ranges['_widths_'+title]=widths
   ws.set_landscape();ws.set_paper(8);ws.fit_to_pages(1,0);ws.set_margins(.25,.25,.4,.4);ws.repeat_rows(0,2)
   ws.set_header('&C第'+str(i)+'次｜'+TITLES[i]);ws.set_footer('&L單位依欄名；pp為百分點&R&P / &N')
   self.note(title,0,f'第{i}次｜{TITLES[i]}','title',34)
   self.note(title,1,COMMON,height=40);self.note(title,2,HISTORY,height=38)
 def write(self,s,row,col,val,kind='text'):
  value=NA if val is None else val;self.sheets[s].write(row,col,value,self.formats[kind]);self.checks.append(dict(sheet=s,row=row,col=col,expected=value));self.ends[s]=max(self.ends.get(s,0),row)
 def formula(self,s,r,c,expr,value,kind='num'):
  value=NA if value is None else value;self.sheets[s].write_formula(r,c,expr,self.formats[kind],value);self.checks.append(dict(sheet=s,row=r,col=c,expected=value,formula=expr));self.ends[s]=max(self.ends.get(s,0),r)
 def note(self,s,row,text,kind='note',height=42):
  self.sheets[s].merge_range(row,0,row,7,text,self.formats[kind]);self.sheets[s].set_row(row,height);self.ends[s]=max(self.ends.get(s,0),row)
 def table(self,s,row,title,headers,data,keyname,types=None):
  self.note(s,row,title,'head',28);row+=1
  for c,h in enumerate(headers):self.write(s,row,c,h,'head')
  self.sheets[s].set_row(row,40);start=row+1
  for r,record in enumerate(data,start):
   for c,v in enumerate(record):self.write(s,r,c,v,(types or {}).get(c,'num' if isinstance(v,float) else 'text'))
   widths=self.ranges['_widths_'+s];lines=1
   for col,v in enumerate(record):
    if isinstance(v,str):lines=max(lines,sum(max(1,math.ceil(sum(2 if unicodedata.east_asian_width(ch) in 'WF' else 1 for ch in segment)/max(5,widths[col]-2)*13/11)) for segment in v.split('\n')))
   self.sheets[s].set_row(r,min(300,max(45,lines*18+6)))
  self.ranges[keyname]=dict(sheet=s,header=row,start=start,end=start+len(data)-1,columns=len(headers))
  return start+len(data)+2
 def chart(self,s,keyname,cols,title,unit,anchor,percent=False):
  t=self.ranges[keyname];c=self.wb.add_chart({'type':'column'});c.set_title({'name':title});c.set_y_axis({'name':unit,'min':0,**({'max':1,'num_format':'0%'} if percent else {})});c.set_legend({'position':'bottom'} if len(cols)>1 else {'none':True})
  for k,col in enumerate(cols):c.add_series({'name':[t['sheet'],t['header'],col],'categories':[t['sheet'],t['start'],0,t['end'],0],'values':[t['sheet'],t['start'],col,t['end'],col],'fill':{'color':['9BA7B0','5B7790','708FA3','254E70'][k%4]},'data_labels':{'value':True,'num_format':'0%' if percent else '0.00'}})
  c.set_size({'width':570,'height':310});self.sheets[s].insert_chart(anchor,c);self.chart_count+=1
 def sources(self,s,row):
  self.note(s,row,'數據來源與比較限制','head',28);row+=1
  for sid,p in self.e.sources.items():self.note(s,row,f"{sid}：{p['path']}；SHA256 {p['sha256']}",height=38);row+=1
  self.note(s,row,'論文：Sun等，DOI 10.1109/LGRS.2025.3650098。PDF p2方法、p3 Algorithm1、p4 TableI/II與Eq8、p5 TableIII。作者程式同環境重跑、論文成功門檻與逐場景分母待補；本地5cm參考不是作者可執行baseline。',height=52)
 def save(self,talk):
  self.checks=list({(x['sheet'],x['row'],x['col']):x for x in self.checks}.values())
  for s,ws in self.sheets.items():ws.print_area(0,0,self.ends.get(s,0),7)
  self.wb.close();txt=self.path.with_suffix('.txt');txt.write_text(talk.strip()+'\n',encoding='utf-8')
  m=dict(report=self.i,title=TITLES[self.i],generated_at=datetime.now(ZoneInfo('Asia/Taipei')).isoformat(),xlsx=self.path.name,txt=txt.name,xlsx_sha256=sha(self.path),txt_sha256=sha(txt),sources=self.e.sources,groups={k:{a:b for a,b in g.items() if a!='rows'} for k,g in self.e.groups.items()},paired=self.e.paired,checks=self.checks,tables=self.ranges,charts=self.chart_count)
  (OUT/f'report_{self.i:02d}_manifest.json').write_text(json.dumps(m,ensure_ascii=False,indent=2)+'\n')
  validate(self.path,m);print(json.dumps(dict(report=self.i,file=self.path.name,txt=txt.name,sheets=list(self.sheets),checks=len(self.checks),charts=self.chart_count),ensure_ascii=False))
  return self.path

def validate(path,m):
 wb=load_workbook(path,data_only=False);v=load_workbook(path,data_only=True);assert 1<=len(wb.sheetnames)<=3 and not wb._external_links
 for x in m['checks']:
  z=v[x['sheet']].cell(x['row']+1,x['col']+1).value;ex=x['expected']
  assert (math.isclose(z,ex,rel_tol=1e-10,abs_tol=1e-10) if isinstance(ex,(int,float)) else z==ex or ex=='' and z is None),(path.name,x,z)
  if 'formula' in x:assert wb[x['sheet']].cell(x['row']+1,x['col']+1).data_type=='f'
 for ws in v:
  assert not any(c.data_type=='e' for row in ws for c in row)
 assert sum(len(s._charts) for s in wb)==m['charts'];wb.close();v.close()

def format_pair_table(b,e,s,row,entries,title='相同配對與seed：新增成功、退步與共同成功精度',keyname='paired'):
 data=[]
 for name,a,z in entries:
  j=e.pair(a,z);data.append([name,j['gained'],j['lost'],j['common'],cm(j['before']),cm(j['after']),None,'僅共同成功；不取代全體成功率'])
 end=b.table(s,row,title,['比較','新增成功','退步失敗','共同成功N','前e_p cm','後e_p cm','下降率','限制'],data,keyname)
 t=b.ranges[keyname]
 for r,(_,a,z) in enumerate(entries,t['start']):
  j=e.pair(a,z);rr=r+1;expected=(j['before']-j['after'])/j['before'] if j['before'] not in [0,None] and j['after'] is not None else None
  b.formula(s,r,6,f'=IF(AND(ISNUMBER(E{rr}),ISNUMBER(F{rr}),E{rr}<>0),(E{rr}-F{rr})/E{rr},"待補")',expected,'pct')
 return end

def report1():
 e=Evidence();e.load_path('F','results/report_experiments_20261028/eth_coarse_rows.csv');e.load_path('T','results/tongji_trees/fixed_eth_20261025/registration_rows.csv');g=[]
 for threshold in ['0.05','0.1']:
  for solver in ['svd','svd+ransac']:g.append(e.group('F:'+threshold+':'+solver,'F',threshold=threshold,solver=solver))
 assert [x['ok'] for x in g]==[114,121,138,144]
 names=['5cm／SVD','5cm／RANSAC','10cm／SVD','10cm／RANSAC'];by_scene={'ETH':g}
 for ss in SCENES[:4]:by_scene[ss]=[e.group('T:'+ss+':'+m+':'+so,'T',scene=ss,method=m,solver=so) for m in ['reference5cm','baseline10cm'] for so in ['svd','svd+ransac']]
 b=Book(1,e,['四組結果比較','設定與說明','精度時間與成對比較']);s='四組結果比較'
 data=[[label(ss),'10配對×10seed' if ss!='ETH' else '15配對×10seed',*[x['rate'] for x in by_scene[ss]],1,PAPER[ss]*100] for ss in SCENES]
 row=b.table(s,4,'兩組dataset；每個場景都比較相同四組，補救全部關閉',['dataset／場景','案例組成',*names,'論文成功率\n分母待補','論文e_p cm'],data,'main',{2:'pct',3:'pct',4:'pct',5:'pct',6:'pct'})
 for r,ss in enumerate(SCENES,b.ranges['main']['start']):
  for col,x in enumerate(by_scene[ss],2):b.formula(s,r,col,f"={x['ok']}/{x['n']}",x['rate'],'pct');b.sheets[s].write_formula(r,col,f"={x['ok']}/{x['n']}",b.wb.add_format({'font_name':'Microsoft JhengHei','font_size':13,'num_format':f'"{x["ok"]}/{x["n"]} ("0.0%")"','valign':'vcenter'}),x['rate'])
 b.note(s,row,'統一seed31–40：Tongji每Plot五站、10對×10seed；ETH六站、15對×10seed。ETH四組為本次補跑，先前118/123/135/142屬seed11–20，保留為歷史資料。Tongji四Plot是同一dataset。',height=48);row+=2
 data=[[label(ss),*[cm(x['ep']) for x in by_scene[ss]],PAPER[ss]*100] for ss in SCENES]
 row=b.table(s,row,'有限mean e_p（cm），包含可計算失敗；樣本數另列',['dataset／場景',*names,'論文e_p cm'],data,'errors')
 row=b.table(s,row,'有限誤差樣本數：缺變換是失敗，e_p不填0',['dataset／場景',*names],[[label(ss),*[x['finite'] for x in by_scene[ss]]] for ss in SCENES],'finite')
 b.note(s,row,'ETH改善不能直接外推Tongji。Plot2部分改善、Plot4退步；有限案例集合不同時不直接算誤差下降率。完整e_R/e_t、時間及共同成功見第三頁。',height=42)
 b.chart(s,'main',[2,3,4,5],'四組成功率：各場景分開','成功率','A34',True);b.chart(s,'errors',[1,2,3,4],'四組有限mean e_p','cm','A53');b.ends[s]=72
 comparisons=[('SVD固定','5→10cm',g[0],g[2]),('RANSAC固定','5→10cm',g[1],g[3]),('5cm固定','SVD→RANSAC',g[0],g[1]),('10cm固定','SVD→RANSAC',g[2],g[3])]
 s='設定與說明'
 data=[['CN距離容差','0.05m','0.10m','允許descriptor距離差更大；也可能增加錯配'],['Solver','Kabsch／SVD','Kabsch＋RANSAC','先尋找3D一致群，再估剛體變換'],['弱配對補救','關閉','關閉','第一週固定，不介紹或歸因此模組'],['提取與descriptor','同一批樹位','同前','step、分群0.20m、verticality>0.9；N183/w10/Thscore40固定']]
 row=b.table(s,4,'本週只看兩個變數',['項目','前／設定A','後／設定B','作用與固定條件'],data,'settings')
 all_comparisons=[]
 for ss in SCENES:
  z=by_scene[ss]
  all_comparisons.extend([(ss,'SVD固定','5→10cm',z[0],z[2]),(ss,'RANSAC固定','5→10cm',z[1],z[3]),(ss,'5cm固定','SVD→RANSAC',z[0],z[1]),(ss,'10cm固定','SVD→RANSAC',z[2],z[3])])
 data=[[label(ss),fixed,change,f"{a['ok']}/{a['n']}",f"{z['ok']}/{z['n']}",None,z['ok']-a['ok'],'只在同場景、同seed內比較'] for ss,fixed,change,a,z in all_comparisons]
 row=b.table(s,row,'固定另一項，逐場景看單項改動',['場景','固定條件','唯一改動','前成功數','後成功數','後−前 pp','成功數淨增加','比較限制'],data,'effects')
 for rr0,(_,_,_,a,z) in enumerate(all_comparisons,b.ranges['effects']['start']):b.formula(s,rr0,5,f'=({z["ok"]}/{z["n"]}-{a["ok"]}/{a["n"]})*100',pp(a,z),'pp')
 b.note(s,row,'RANSAC固定設定為最多1000次抽樣嘗試、3D inlier<0.30m、seed1；它們是solver定義，這週沒有各常數消融。GT正確matching欄只檢查XY<0.30m，因此XY正確也可能存在z不一致。',height=46);row+=2
 b.note(s,row,'時間為descriptor＋matching＋solver，不含樹位提取、原始點雲讀取及GT誤差評估。舊批次OpenMP4，CPU型號與並行干擾未完整控制，只作本批描述性代價；不能與論文matching／total時間直接算速度提升。',height=46);row+=2
 b.note(s,row,'論文TableI的ETH e_p=15.3427cm，Tongji Plot1～4=5.2558/11.4808/4.8756/4.3560cm，各報100%；分母與成功門檻待補。四組本地控制實驗不是作者程式重跑；目前結果不支持「這兩項改動在所有場景都改善」。',height=46);row+=2
 b.sources(s,row)
 s='精度時間與成對比較';data=[]
 for ss in SCENES:
  for name,x in zip(names,by_scene[ss]):data.append([label(ss),name,cm(x['ep']),x['finite'],x['rot'],cm(x['trans']),x['time'],f"{x['ok']}/{x['n']}"])
 row=b.table(s,4,'四組×五場景：精度與時間，未混算兩來源平均',['場景','組合','finite mean e_p cm','有限N','mean e_R °','mean e_t cm','平均time_ms','成功／總數'],data,'detail')
 b.note(s,row,'e_R/e_t亦平均可計算值、包含有限失敗。time_ms只含descriptor＋matching＋solver；不同來源批次不作跨dataset速度排名。沒有變換的案例仍算成功率失敗。',height=42);row+=2
 format_pair_table(b,e,s,row,[(label(ss)+'；'+fixed+'；'+change,a,z) for ss,fixed,change,a,z in all_comparisons])
 talk=talk1(g,e,comparisons,by_scene)
 return b.save(talk)

def talk1(g,e,comparisons,by_scene):
 results='\n'.join(f"{name}：成功{x['ok']}/{x['n']}，成功率{x['rate']*100:.2f}%；有限mean e_p={fmt(cm(x['ep']))}cm，有限N={x['finite']}；平均e_R={fmt(x['rot'])}°、e_t={fmt(cm(x['trans']))}cm；平均time_ms={fmt(x['time'])}ms。" for name,x in zip(['5cm／SVD','5cm／RANSAC','10cm／SVD','10cm／RANSAC'],g))
 paired='\n'.join(f"{fixed}、{change}：新增{e.pair(a,z)['gained']}、退步{e.pair(a,z)['lost']}，共同成功{e.pair(a,z)['common']}次；共同成功mean e_p {fmt(cm(e.pair(a,z)['before']))}→{fmt(cm(e.pair(a,z)['after']))}cm。" for fixed,change,a,z in comparisons)
 return f'''第1次報告：CN距離容差與Solver
建議主稿約10～12分鐘，加圖表停頓與問答可至15分鐘。末尾備答不必全部念。
搭配 report_01_CN距離容差與Solver.xlsx；先看「四組結果比較」，再看「設定與說明」，必要時查「精度時間與成對比較」。本版同時包含Tongji-trees與ETH-tree。

【0:00～1:30｜今天只回答兩個簡單問題】
老師好，今天先從兩個容易控制的變數開始：第一，CN距離容差從5公分放寬到10公分，能不能找回更多可用的匹配？第二，樹位已經配好後，直接用SVD求姿態，和加入RANSAC有什麼差別？
這週先不加入弱配對補救，也不加點雲幾何微調。四組補救全部關閉，這樣比較的差異才不會同時混進另一個模組。弱配對補救留到第二次獨立說明。
我做的是2×2對照：5cm搭SVD、5cm搭RANSAC、10cm搭SVD、10cm搭RANSAC。兩組dataset都列相同四種設定，各來源內使用同一批樹位，統一seed31到40。ETH六站、15配對，每組150次；Tongji-trees每個Plot五站、10配對，每組100次。四Plot是同一dataset不同場景，不是四個獨立來源。ETH本次補跑600筆；舊版118/123/135/142是seed11～20，沒有把它們混進這批。新seed仍是既有森林，不是盲測。

【1:30～3:30｜CN距離容差改在哪裡】
CircularMatch用一棵樹周圍其他樹的排列建立descriptor。兩站看到同一棵樹時，因為提取交點、遮擋和漏樹，描述子的距離不一定完全一致。CN threshold是在比較這些距離時，允許它們相差多少。
原論文使用0.05m，我比較0.10m。這不是把成功門檻改寬，也不是把樹幹提取voxel改成10cm；提取仍是原來設定。更不是把Thscore40改成10，Thscore是分數，threshold是以m表示的距離。
這個距離容差同時用在gatekeeper、descriptor一致格數DC，以及triangle fallback的邊長比較。因此放寬它可能讓更多原本稍微不一致的真實關係通過，也可能讓不正確的關係混進來。目的不是假設越大越好，而是檢查5到10cm這一個範圍的效果。
舉一個教學假例：兩站估同一對樹的距離為5.00m與5.07m，差7cm，5cm門檻不過，10cm門檻能過。可是另一組錯樹的距離若也碰巧接近，它也可能通過，所以容差不等於正確性的保證。

【3:30～5:00｜Solver為什麼可能有影響】
matching提供的是樹位對應，Solver把它們轉成旋轉R與平移t。Kabsch／SVD會最小化這些3D對應的平方殘差。只要對應可靠，這個方法可以直接求解；若裡面有不一致的樹位，平方誤差也會試著遷就它們。
RANSAC會從現有對應中隨機取三對，估一個姿態，再看有多少對應在這個姿態下足夠一致，最後用較一致的一群重估。這裡的RANSAC是在樹位配準階段，不是前面點群擬合圓柱時的另一個RANSAC。
即使事後GT的XY檢查說一對是正確的，它的z仍可能不一致。因為GT正確對應欄只檢查XY，而Solver使用XYZ，所以RANSAC仍有可能排掉3D高殘差。這是機制解釋，不能在沒有逐案診斷下說每次改善都一定是z造成。

【5:00～8:00｜直接讀四組與單項對照】
先看第一頁ETH那一列，四組的完整數字如下；第三頁可查同一組的旋轉、平移與時間。
{results}
成功率最清楚的讀法是固定另一項。固定SVD只改5到10cm，{g[0]['ok']}增加到{g[2]['ok']}，多{g[2]['ok']-g[0]['ok']}次，成功率增加{pp(g[0],g[2]):.2f}個百分點。固定RANSAC只改容差，{g[1]['ok']}增加到{g[3]['ok']}，多{g[3]['ok']-g[1]['ok']}次，增加{pp(g[1],g[3]):.2f}個百分點。
反過來，固定5cm改Solver，{g[0]['ok']}到{g[1]['ok']}，多{g[1]['ok']-g[0]['ok']}次；固定10cm改Solver，{g[2]['ok']}到{g[3]['ok']}，多{g[3]['ok']-g[2]['ok']}次。這支持兩種改動在本輪各有作用，但作用大小與另一項設定有關，不能直接假設完全獨立、線性相加。
兩項一起改，第一組{g[0]['ok']}到第四組{g[3]['ok']}，淨增加{g[3]['ok']-g[0]['ok']}次。不能全部稱為放寬容差的貢獻，也不能全部算給RANSAC。下方四條固定條件的比較就是為了避免這種歸因錯誤。
成功數的淨增加也不代表沒有任何退步。第三頁列同pair與seed的新增、退步，以及兩組都成功時的平均誤差；以下先讀ETH：
{paired}

再看Tongji，這是本週不能省略的跨來源檢查。四組順序仍是5cm/SVD、5cm/RANSAC、10cm/SVD、10cm/RANSAC：
{chr(10).join(label(ss)+'：'+ '、'.join(str(x['ok'])+'/100' for x in by_scene[ss])+'。' for ss in SCENES[:4])}
Plot1與Plot3四組都100%，成功率沒有再提升，但誤差仍需分開看。Plot2放寬容差時，SVD由33到37，RANSAC由28到34；然而同容差改成RANSAC，反而分別33到28、37到34。Plot4則更直接：5cm兩solver都是76/100，10cm變成69或70/100。因此「ETH較好的組合」不能直接稱為「所有森林較好的組合」。
若只對Tongji四Plot累計成功數，四組是309/400、304/400、306/400、304/400。這只是同dataset試驗數相加，不是把四Plot誤當四dataset；我們沒有將它再與ETH混成一個不明權重的平均。5cm/SVD到10cm/RANSAC在Tongji是309到304，並未改善。主表完整列出這些退步，沒有為各Plot挑不同最佳參數。

【8:00～10:00｜精度與時間怎麼解讀】
第一頁同時保留e_p、旋轉和平移誤差。e_p是把全部來源點分別用估計變換與GT變換搬動後，比較兩個位置的距離，再取平均；不是兩片點雲的最近鄰fitness。
有限mean包含有輸出的失敗，不會只留下成功案例。但5cm兩組只有{g[0]['finite']}個finite誤差，10cm兩組有{g[2]['finite']}個，所以不同容差的平均值還受輸出案例集合改變影響。圖可以呈現這些平均值，但我不把它們直接算成同一批案例的精度下降率。缺的輸出仍是成功率的失敗，誤差不填0。
若要回答原本就能成功的案例是否更準，第三頁的共同成功平均比較有幫助；它仍有存活者偏差，因此不能取代全體成功率。這兩個統計回答的問題不同。Tongji也須按各組有限N檢查，不能只挑成功案例的平均值聲稱全體改善。
時間方面，ETH四組大約是{'、'.join(fmt(x['time']) for x in g)}ms。新增計算有代價，不能只報成功率。這些是descriptor、matching與Solver的內部時間，不包含提取或GT評分；本次ETH補跑採兩worker、每程序OMP4，共享CPU quota4。Tongji時間來自舊批次，不做跨批次速度排名，也不宣稱與論文公平速度排名。

【10:00～11:30｜本次結論與下一個問題】
本次已支持的結論是：相同ETH樹位與配對下，10cm比5cm有較多成功，RANSAC比直接SVD也有較多成功；但Tongji的控制結果證明效果不一致，Plot4放寬容差退步、Plot2換Solver也有退步。這不只是一句抽象的「不能外推」，而是已有另一來源的反例。現有材料不足以把Tongji所有退步指定為同一個高度或錯配原因，後續需分別診斷。
原論文只說求Euclidean transform，沒有足夠細節讓我們確認其Solver與本地SVD完全相同。原作者程式的同環境重跑仍待補；論文成功分母與門檻也未齊，所以這週不宣稱嚴格超越原作。
ETH本地較好的這組是{g[3]['ok']}/150，仍有{150-g[3]['ok']}次失敗；同組無補救設定在Tongji只有304/400。下一次不繼續堆更多小參數，而是看最後接受的對應是否不足，以及能否回收CN已有的證據來補救。那是一個獨立方法主題，本次不把它的改善加進這四組。

【教授追問備答】
問：為什麼挑10cm，不試到更大？
答：這是既有探索後的候選，今天只主張5到10cm的對照結果。容差更大可能誤配更多，不能外推單調改善。
問：5cm到10cm讓成功率提高多少？
答：ETH本批固定SVD是+{pp(g[0],g[2]):.2f}pp；固定RANSAC是+{pp(g[1],g[3]):.2f}pp；Tongji各Plot另列，不能只說全體都改善。
問：這裡有用弱配對補救嗎？
答：四組都關閉，否則不能只比較這兩個變數。
問：成功數為什麼與舊版不同？
答：這份使用統一seed31～40，舊ETH四組為11～20。兩批完整保留，沒有挑最好seed；同配置不同隨機提取可得到不同成功數。
問：為何GT正確對應數一樣，Solver還有差？
答：該診斷只看GT下XY<0.30m；3D高度與估計姿態仍可能不一致，且門檻只是判準，不等於零誤差。
問：平均e_p比較小就代表每一例都變好？
答：不代表；要看有限N、配對差異及共同成功。沒有變換不能以零誤差填入。
問：是否完全不用GT？
答：推論不靠GT選輸出，但離線開發用GT評分比較過設定，不能說完全無真值選參或盲測。

【資料追溯】
F：{e.sources['F']['path']}
SHA256：{e.sources['F']['sha256']}
T：{e.sources['T']['path']}
SHA256：{e.sources['T']['sha256']}
各組精確CSV行號、配對比較與公式期望值另存report_01_manifest.json；理解本稿不需先讀manifest。
原論文：DOI 10.1109/LGRS.2025.3650098；PDF p2、p3 Algorithm1、p4 TableI/II與Eq8。
'''

def setup_later(i):
 e=Evidence();e.load('U','U_registration_rows.csv');e.load_path('C','results/report_experiments_20261028/eth_coarse_rows.csv');e.load_path('T','results/tongji_trees/fixed_eth_20261025/registration_rows.csv')
 entries=[];extra=[]
 if i==2:
  for s in SCENES:
   a=e.group('C:10R','C',threshold='0.1',solver='svd+ransac') if s=='ETH' else e.group('T:'+s+':10R','T',scene=s,method='baseline10cm',solver='svd+ransac')
   entries.append((s,a,e.u(s,'eth10_raw')))
  settings=[['CN容差／Solver','0.10m／RANSAC','相同','固定第一週選定的共用配置，不為各Plot挑不同版本'],['弱配對補救','關閉','evidence模式；原CN<6對啟動','本週唯一新增方法，內部常數作為整組'],['補救證據與採用','無新增模組','前三候選證據、最多60邊、至少6個3D支持','只列整組定義，不將改善分配給各常數'],['地形／提取／精化','原DTM、step、無ICP','相同','固定其他處理，推論僅當前兩站']]
 elif i==3:
  e.load_path('FULL','results/report_experiments_20261028/full_dtm_smooth_rows.csv');e.load_path('EP','results/report_experiments_20261028/eth_dtm_position_rows.csv');e.load('D','D_registration_rows.csv');e.load_path('FC','results/report_experiments_20261028/feature_count_rows.csv')
  for s in SCENES:
   entries.append((s,e.u(s,'eth10_raw'),e.group('FULL:'+s,'FULL',scene=s,solver='svd+ransac')))
   pos=e.group('EP:R','EP',solver='svd+ransac') if s=='ETH' else e.group('D:'+s+':smooth','D',scene=s,method='selected',solver='svd+ransac')
   extra.append((s,e.u(s,'eth10_raw'),pos))
  settings=[['DTM cell','0.50m','0.50m','網格大小固定'],['DTM fit_radius／fill_max_radius','2格／5格','6格／6格','同時修改；fit_radius亦影響鄰近平面剔除，不拆單參數歸因'],['前次coarse方法','10cm＋RANSAC＋weak recovery','相同','只改DTM profile；不為ETH重新選參'],['主比較輸入','原DTM完整提取','smooth完整重提取','每站重新height/voxel/normal/cluster/cylinder/intersection'],['定位階段控制','固定既有圓柱軸線','僅重算axis×DTM交點','另列兩dataset定位測試；不與完整提取混稱']]
 elif i==4:
  e.load_path('FULL','results/report_experiments_20261028/full_dtm_smooth_rows.csv')
  for s in SCENES:
   entries.append((s,e.group('FULL:'+s,'FULL',scene=s,solver='svd+ransac'),e.u(s,'eth10_refined')))
   extra.append((s,e.u(s,'eth10_raw'),e.u(s,'eth10_refined')))
  settings=[['跨週DTM設定','前次smooth完整提取','回到原DTM','跨週同時回復DTM與新增精化，只能說整組效果'],['surface refinement','無','新增完整6DOF point-to-plane','另列原DTM raw→refined控制，隔離新增後端整組'],['表面尺度／迭代預算','無新階段','0.50→0.20→0.10m；30/25/20次','新增階段，提取voxel仍1cm'],['robust與回退','無新階段','normal gate、唯一target、trimming、Huber、幾何接受','整組啟用，沒有每個內部常數單獨消融'],['前次coarse方法','10cm＋RANSAC＋weak recovery','相同','固定CN與solver；不使用多站圖或GT選輸出']]
 elif i==5:
  for s in SCENES:
   entries.append((s,e.u(s,'eth10_refined'),e.u(s,'multi_v2_refined')));extra.append((s,e.u(s,'multi_refined'),e.u(s,'multi_v2_refined')))
  settings=[['前次方法','單初值＋surface refinement','多初值v2＋相同舊refinement','主表直接比較第4次，不跳過前一版本'],['候選證據範圍 v1→v2','top5／最多60邊','top20／最多120邊','與其他搜尋／評分同改，不稱top-k獨立效果'],['姿態提案與精化預算 v1→v2','最多16替代姿態／精化3個','最多256替代姿態／精化8個','加上既有初值後proposal_count可為257；beam8含初值'],['初值與支持 v1→v2','3D三元組、證據池支持','另加水平雙點yaw、當前兩站全部樹位支持','最終仍完整6DOF；global不是跨站地圖'],['既有初值 v1→v2','不保證進精化beam','保留進beam','保留參與不保證最後選回它'],['驗證規則','全方向、gate=3h','相同','本週先保留錯誤分數造成退步的反例']]
 else:
  for s in SCENES:
   entries.append((s,e.u(s,'multi_v2_refined'),e.u(s,'multi_v3_refined')))
  settings=[['多候選引擎','v2','相同v2引擎','候選產生規則固定'],['驗證表面','全方向','優先|normal.z|<0.30','兩站各至少30點，不足退回全方向'],['驗證NN gate','3h：1.50/0.60/0.30m','min(h,0.20)：0.20/0.20/0.10m','與表面篩選一起改；缺單項消融'],['optimizer與使用位置','全方向6DOF point-to-plane','optimizer相同；新分數用於排序、接受與最終選擇','不是只換最後sort；地面仍可參與優化'],['原DTM與提取verticality','原DTM；verticality>0.9','相同','不是把提取0.9改成驗證0.3']]
 return e,entries,extra,settings

def later(i):
 e,entries,extra,settings=setup_later(i);b=Book(i,e,['結果比較','設定與說明','必要對照']);s='結果比較'
 versions={2:('第1次：10cm/RANSAC，補救關閉','第2次：相同coarse＋weak recovery'),3:('第2次：原DTM＋weak recovery','第3次：smooth完整重提取＋相同coarse'),4:('第3次：smooth完整提取、無refinement','第4次：原DTM＋surface refinement'),5:('第4次：原DTM單初值refinement','第5次：多初值v2＋舊驗證'),6:('第5次：多初值v2＋舊驗證','第6次：多初值v2＋v3驗證')}
 before,after=versions[i];b.note(s,3,'修改前＝'+before+'；修改後＝'+after,height=42)
 data=[[label(ss),1,f"{e.u(ss,'reference5_raw')['ok']}/{a['n']}",f"{a['ok']}/{a['n']}",f"{z['ok']}/{z['n']}",None,None,f"{len({r['pair'] for r in a['rows']})}對×10seed；31–40"] for ss,a,z in entries]
 row=b.table(s,5,'每次均列原論文、本地5cm參考、前一次、本次',['場景','論文成功率\n分母待補','本地5cm/SVD','前一次成功／總數','本次成功／總數','vs前一次 pp','vs論文 pp','案例組成'],data,'success',{1:'pct'})
 for r,(_,a,z) in enumerate(entries,b.ranges['success']['start']):
  b.formula(s,r,5,f'=({z["ok"]}/{z["n"]}-{a["ok"]}/{a["n"]})*100',pp(a,z),'pp');rr=r+1;b.formula(s,r,6,f'=({z["ok"]}/{z["n"]}-B{rr})*100',(z['rate']-1)*100,'pp')
 b.note(s,row,'Tongji四Plot屬同一dataset，ETH屬另一dataset；兩者不混算誤差平均。前後均同配對與seed，失敗與缺輸出都保留。原論文門檻與分母未確認一致，100%只能作報告值參考。',height=44);row+=2
 data=[]
 for ss,a,z in entries:
  ref=e.u(ss,'reference5_raw');data.append([label(ss),cm(PAPER[ss]),cm(ref['ep']),cm(a['ep']),cm(z['ep']),None,None,f"{ref['finite']}／{a['finite']}／{z['finite']}（總N{a['n']}）"])
 row=b.table(s,row,'全來源點Eq8後再跨有限案例mean；單位cm',['場景','論文e_p','本地5cm e_p','前一次e_p','本次e_p','vs前一次下降率','vs論文下降率','finite N：參考／前／後'],data,'errors')
 for r,(ss,a,z) in enumerate(entries,b.ranges['errors']['start']):
  rr=r+1;j=e.pair(a,z)
  if j['same_finite']:b.formula(s,r,5,f'=IF(AND(ISNUMBER(D{rr}),ISNUMBER(E{rr}),D{rr}<>0),(D{rr}-E{rr})/D{rr},"待補")',(a['ep']-z['ep'])/a['ep'] if a['ep'] not in [None,0] and z['ep'] is not None else None,'pct')
  if z['finite']==z['n']:b.formula(s,r,6,f'=IF(AND(ISNUMBER(B{rr}),ISNUMBER(E{rr}),B{rr}<>0),(B{rr}-E{rr})/B{rr},"待補")',(PAPER[ss]-z['ep'])/PAPER[ss],'pct')
 b.note(s,row,'下降率=(基準−本次)/基準。有限案例集合不同不算vs前下降率；本次仍缺變換不算完整vs論文下降率。這是不可比的統計，不能靠補跑把缺輸出填0；共同成功另列。',height=44);row+=2
 # Clean chart sources: rates and finite means, all populated on every scene.
 row=b.table(s,row,'圖表使用相同場景的前後值',['場景','前成功率','後成功率','前finite mean cm','後finite mean cm'],[[label(ss),a['rate'],z['rate'],cm(a['ep']),cm(z['ep'])] for ss,a,z in entries],'plot',{1:'pct',2:'pct'})
 b.chart(s,'plot',[1,2],'成功率：前一次與本次','成功率','A35',True);b.chart(s,'plot',[3,4],'有限mean e_p：含有限失敗','cm','A54');b.ends[s]=73
 s='設定與說明';row=b.table(s,4,'當週重要變數與組合，不列全部內部常數',['項目','前一次／修改前','本次／修改後','作用與歸因限制'],settings,'settings')
 b.note(s,row,'主表比較的是上一週方法與本次方法。第3週分開列定位與完整重提取；第4週明列回復DTM＋加refinement兩項變更，另有相同DTM的精化控制；第5週另保留v1→v2負面結果。',height=46);row+=2
 times=dict(zip(SCENES,[(20.718,11.657),(30.699,16.712),(4.5596,8.3584),(19.522,9.4516),(3.0405,8.6992)]))
 row=b.table(s,row,'原論文時間只作範圍參考；不計跨硬體速度提升',['場景','論文matching ms','論文total s','硬體／範圍限制'],[[label(ss),*times[ss],'i7-9700F 3GHz；thread數待補。本地online不含提取/cache/GT，範圍不同。'] for ss in SCENES],'paper_time')
 b.sources(s,row)
 s='必要對照';row=format_pair_table(b,e,s,4,[(label(ss),a,z) for ss,a,z in entries])
 row=b.table(s,row,'旋轉、平移與本地逐配對時間',['場景','前e_R °','後e_R °','前e_t cm','後e_t cm','前online ms','後online ms','計時限制'],[[label(ss),a['rot'],z['rot'],cm(a['trans']),cm(z['trans']),a['time'],z['time'],'跨批次／epoch或範圍差；不算速度下降率'] for ss,a,z in entries],'pose_time')
 if extra:
  caption={3:'固定舊軸線：只重算DTM交點，與完整重提取分開',4:'相同原DTM：raw→refined，隔離新增幾何後端整組',5:'v1→v2：搜尋與評分整組改動的負面結果'}[i]
  row=b.table(s,row,caption,['場景','修改前成功','修改後成功','前finite e_p cm','後finite e_p cm','前finite N','後finite N','限制'],[[label(ss),f"{a['ok']}/{a['n']}",f"{z['ok']}/{z['n']}",cm(a['ep']),cm(z['ep']),a['finite'],z['finite'],'同pair/seed；整組對照，非每個內部參數消融'] for ss,a,z in extra],'extra')
 if i==5:
  rr=e.tables['U'];case=[r for r in rr if r['scene']=='Plot4' and r['pair']=='S4-S2' and r['seed']=='31' and r['method'] in ['multi_refined','multi_v2_refined']]
  row=b.table(s,row,'同一失敗反例：Plot4 S4-S2 seed31，舊分數定義相同',['版本','候選pose數','精化數','成功?','geometry overlap','e_p m','解讀'],[[r['method'],int(r['proposal_count']),int(r['refined_candidates']),int(r['success']),float(r['geo_overlap']),float(r['e_p_m']),'分數較大不代表GT下更準'] for r in case],'case')
 if i==3:
  data=[]
  for ss in SCENES:
   rr=[r for r in e.tables['FC'] if r['scene']==ss];a=sum(int(r['before_trees']) for r in rr)/len(rr);z=sum(int(r['after_trees']) for r in rr)/len(rr)
   data.append([label(ss),len({r['scan'] for r in rr}),len(rr),a,z,z-a,'每站每seed平均；不是registration成功率'])
  row=b.table(s,row,'中間證據：完整重提取實際改變的樹位數',['場景','站數','站×seed檔數','原平均樹位數','smooth平均樹位數','後−前','解讀'],data,'features')
  b.note(s,row,'FC保存每站每seed的樹數與前後CSV雜湊，不公開樹位座標。樹數改變支持DTM影響提取集合；沒有單獨實驗證明「樹數更多」就是誤差改善原因。',height=42)
 if i==6:
  row=b.table(s,row,'公平控制：三者使用同一v3 backend',['場景','5cm/SVD＋v3','10cm/recovery＋v3','完整multi_v3','5cm e_p cm','10cm e_p cm','multi e_p cm','限制'],[[label(ss),*[f"{e.u(ss,m)['ok']}/{e.u(ss,m)['n']}" for m in ['reference5_refined_v3','eth10_refined_v3','multi_v3_refined']],*[cm(e.u(ss,m)['ep']) for m in ['reference5_refined_v3','eth10_refined_v3','multi_v3_refined']],'不把全部增益歸給multistart'] for ss in SCENES],'fair')
  b.note(s,row,'全體成功數507/550、546/550、550/550，只累計成功數、不混算兩來源誤差。完整候選相對同backend的10cm方法額外救4次，均Plot4。Plot3較v2誤差增加約0.127mm；Plot4共同成功78例增加約0.733mm。',height=48);row+=2
  coldpath=OLD/'evidence/C_cold_benchmark.json';cold=json.loads(coldpath.read_text());data=[]
  for ss in ['ETH','Plot1']:
   a=next(x for x in cold if x['scene']==ss and x['method']=='reference5_refined_v3');z=next(x for x in cold if x['scene']==ss and x['method']=='multi_v3_refined')
   data.append([label(ss)+' '+a['pair'],a['wall_s'],z['wall_s'],None,a['peak_child_rss_kib']/1024,z['peak_child_rss_kib']/1024,1,'固定seed31單次；只作當次觀測'])
  row=b.table(s,row,'額外cold代表配對：同epoch與同backend，含提取＋cache',['場景／配對','5cm參考＋v3 秒','完整multi_v3 秒','時間下降率','參考child RSS MiB','本次child RSS MiB','每法N','限制'],data,'cold')
  for rr0,record in enumerate(data,b.ranges['cold']['start']):rr=rr0+1;b.formula(s,rr0,3,f'=IF(AND(ISNUMBER(B{rr}),ISNUMBER(C{rr}),B{rr}<>0),(B{rr}-C{rr})/B{rr},"待補")',(record[1]-record[2])/record[1],'pct')
  b.note(s,row,'來源 '+str(coldpath.relative_to(ROOT))+'；SHA256 '+sha(coldpath)+'。epoch20261025T162801Z，Xeon8573C、quota4、OMP4、native串行；OS cache未清。RSS為最大native子程序，不含Python與page cache。每法每dataset僅1次，不能宣稱穩定速度排名。',height=64)
 return b.save(later_talk(i,e,entries,extra,settings,versions[i]))

def main():
 OUT.mkdir(parents=True,exist_ok=True)
 ap=argparse.ArgumentParser();ap.add_argument('--report',type=int,default=1);a=ap.parse_args()
 if a.report==1:report1()
 elif 2<=a.report<=6:later(a.report)
 else:raise ValueError('Report index must be 1–6')
if __name__=='__main__':main()
