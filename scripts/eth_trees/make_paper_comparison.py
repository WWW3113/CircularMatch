#!/usr/bin/env python3
"""Plain-language comparison of paper-reported ETH results and our modified method."""
import hashlib
import json
import math
from pathlib import Path
import statistics as stats

from openpyxl import Workbook
from make_study_workbook import table
from tune_registration import ROOT, read_csv

OUT=ROOT/'results/paper_comparison_20261025'
FOLLOW=ROOT/'results/eth_trees/failure_followup_20261025'


def main():
    rows=read_csv(FOLLOW/'runs/validation_rows.csv')
    primary=[r for r in rows if r['method']=='recovery' and r['solver']=='svd+ransac']
    base=[r for r in rows if r['method']=='baseline' and r['solver']=='svd+ransac']
    assert len(primary)==len(base)==150 and all(r['success']=='1' for r in primary)
    mean_ep=stats.mean(float(r['e_p_m']) for r in primary)
    mean_ms=stats.mean(float(r['time_ms']) for r in primary)
    paired=json.loads((FOLLOW/'paired_validation.json').read_text())
    wb=Workbook();wb.remove(wb.active)
    overview=[
        ['本表在比什麼','論文中的ETH結果','我們的兩站改進版','原論文欄是文章報告值，並非把作者程式拿來同機重跑。'],
        ['每次使用的資料','兩站點雲','兩站點雲','這輪沒有借助另外四站。'],
        ['配準成功率','100%','100%（150/150）','成功率追平。論文未明確列成功門檻，不能當作完全相同條件。'],
        ['點雲誤差 e_p', '15.34 cm（Table I原值0.153427 m）',f'{mean_ep*100:.2f} cm（15個配對×10個seed的平均）',
         '我們的數字較低；作者ETH配對與資料集內彙總細節尚未完全對齊，不能據此宣稱嚴格超越。'],
        ['處理時間', '3.0405 ms（Table II：matching strategy）',f'{mean_ms:.2f} ms（描述子＋匹配＋求解平均）',
         '硬體不同，計時範圍也未確認一致；目前不能判定誰比較快。'],
        ['用了哪些案例','ETH精確配對列表／重複次數未明確列出','全部15個配對，各跑seed21–30，共150次','同一片森林重跑10個seed，不是10份獨立資料。'],
        ['成功怎麼判定','未明確定義','旋轉誤差<2°，且平移誤差<0.5m','我們的成功定義沒有在本輪放寬。'],
        ['目前能說的結論','—','成功率追平，誤差數字更低','可以說結果有競爭力；還不能正式宣稱超越論文。'],
        ['建議先看哪裡','—','先看本頁及「改了哪些參數」','其他工作表保留調參證據、每個配對和逐次數字，方便核對。']]
    table(wb,'先看結論',['比較項目','原論文 ETH','我們目前 ETH','白話說明'],overview,[25,48,50,76])
    params=[
        ['匹配距離門檻','5 cm','10 cm','已修改；讓稍有偏差的樹位仍有機會配上。這不是論文原參數。'],
        ['配準求解','只寫計算Euclidean transform，未指定細節','完整6DOF SVD＋配準RANSAC','實作選擇；RANSAC會排除不一致的匹配。不是樹幹圓柱擬合那個RANSAC。'],
        ['匹配不足的補救','最高分不足門檻時做三角形備援','最終匹配少於6對時，再做3D證據一致性補救','我們新增；避免最高分過關、但最後只有一兩對，無法求變換。'],
        ['補救的候選範圍','沒有本輪這個補救步驟','前三個描述子候選；最多60條對應證據','從已存在的描述子資訊找更多候選，不讀取真值選答案。'],
        ['補救採用條件','未提供','3D殘差<0.30m、至少6對支持，且比原匹配多','不夠一致就保留原結果；最後仍估完整三維旋轉與平移。'],
        ['距離downsampling','step：近50%、中75%、遠100%','保留step：<5m保留50%；5–10m保留75%；≥10m保留100%','保留比例沒有改；我們用到掃描器的水平距離。論文距離維度未明確交代。'],
        ['樹幹垂直度','>0.90','>0.90','維持。試過0.85，ETH配準與Liang提取都變差，所以沒有採用。'],
        ['分群距離 clustering','未列參數','0.20 m（本研究選定）','把相近候選點分成樹幹群。論文未給值，不能說作者原本用0.10m。'],
        ['地面模型 DTM','只寫產生DTM，未指定算法','格網地面＋坡度過濾＋局部平面','我們補足的實作細節；改善地面高度估計，讓樹根位置較可靠。'],
        ['圓柱擬合後檢查','未列這些檢查','半徑、傾角、內點比例、法向一致性、圓弧覆蓋','我們的提取實作有額外篩選，並非完全照論文字面的版本。'],
        ['圓柱擬合資源','未列機率與迭代上限','目標機率0.99；最多1000次','維持。試過0.999／3000次，沒有一致改善，所以未採用。'],
        ['描述子其他主要參數','N=183、w=10、Th_score=40','相同','這三個主要參數維持論文數值。'],
        ['高度、voxel、法向半徑','離地0–3m、voxel1cm、法向半徑10cm','相同','這些前處理設定維持不變。'],
        ['ICP、四自由度與六站圖路徑','本表比較兩站配準','本輪沒有使用','本輪100%來自兩站改進版；程式補救選項預設仍為關閉。']]
    table(wb,'改了哪些參數',['項目','論文說法／設定','我們目前使用','這代表什麼'],params,[28,53,63,83])
    first=read_csv(ROOT/'results/eth_trees/tuning_20261025/fresh_summary.csv')
    rates={(r['method'],r['solver']):int(r['fresh_success'])/int(r['trials']) for r in first}
    changes=[['匹配5→10cm；保留SVD','ETH seed11–20',rates['fresh_control','svd'],rates['fresh_matching','svd'],'變好','同批提取結果，只改匹配門檻。'],
             ['只加入配準RANSAC；仍為5cm','ETH seed11–20',rates['fresh_control','svd'],rates['fresh_control','svd+ransac'],'小幅變好','同批樹位、同配對。'],
             ['10cm＋配準RANSAC','ETH seed11–20',rates['fresh_control','svd'],rates['fresh_matching','svd+ransac'],'變好','上一輪推薦組；尚未加本輪補救。'],
             ['垂直度0.90→0.85；保留SVD','ETH seed11–20',rates['fresh_control','svd'],rates['fresh_vertical085','svd'],'變壞','因此保留0.90。'],
             ['圓柱機率0.999／最多3000次；保留SVD','ETH seed11–20',rates['fresh_control','svd'],rates['fresh_ransac999','svd'],'略降','因此保留0.99／1000次。'],
             ['推薦組＋匹配不足時3D補救','ETH新seed21–30',sum(int(r['success']) for r in base)/150,1.,'補回8次、退步0次','先用舊seed選方案，再用這批新seed驗證。不是拿舊seed再宣稱盲測。']]
    ws=table(wb,'調整是否有用',['修改','驗證資料','修改前成功率','修改後成功率','觀察','比較條件'],changes,[52,25,22,22,28,90])
    for i in range(2,ws.max_row+1):
        for j in [3,4]:ws.cell(i,j).number_format='0.0%'
    official={tuple(s.split()) for s in (ROOT/'data/eth_trees/raw/trees/pairs.txt').read_text().splitlines() if s.strip()}
    result=[]
    for pair in sorted({r['pair'] for r in primary}):
        new=[r for r in primary if r['pair']==pair];old=[r for r in base if r['pair']==pair]
        result.append([pair,'是' if tuple(pair.split('-')) in official else '否',10,
                       sum(int(r['success']) for r in old),sum(int(r['success']) for r in new),
                       stats.mean(float(r['e_p_m']) for r in new)*100,'論文未提供逐配對結果'])
    table(wb,'各配對結果',['配對','ETH官方列表','每組次數','原推薦組成功次數','加補救後成功次數','新方法平均點雲誤差cm','原論文逐對數字'],result,[17,22,17,29,29,34,40])
    detail=[];index={(r['seed'],r['pair']):r for r in base}
    for r in primary:
        b=index[r['seed'],r['pair']]
        detail.append([int(r['seed']),r['pair'],int(b['success']),int(r['success']),float(r['rot_err_deg']),float(r['trans_err_m']),float(r['e_p_m'])*100,int(r['recovery_accepted'])])
    table(wb,'新seed逐次',['seed','配對','原推薦組成功0或1','補救後成功0或1','補救後旋轉誤差°','補救後平移誤差m','補救後點雲誤差cm','採用補救0或1'],detail,[12,17,27,27,28,28,29,25])
    sources=[['論文','Sun et al., CircularMatch, IEEE GRSL 23 (2026), DOI 10.1109/LGRS.2025.3650098；以使用者上傳PDF為準。'],
             ['論文Table I','ETH-tree：rate=100%、e_p=0.153427m。e_p式(8)對source點雲所有點計算；表內ETH配對與彙總細節未完全列明。'],
             ['論文Table II','ETH-tree：matching strategy=3.0405ms；作者主機Intel i7-9700F 3.00GHz。'],
             ['論文主要參數','第III節：N=183、w=10、Th_score=40、matching threshold=5cm。'],
             ['論文兩站輸入','第II節第一段：input a pair of point clouds；不是我們後來加入的六站圖路徑。'],
             ['我們的來源','results/eth_trees/failure_followup_20261025/runs/validation_rows.csv；本檔另保留600列完整驗證資料。'],
             ['平均與中位數','首頁使用我們全部150案例的e_p平均值；不是先前回覆中的6.37cm中位數。兩者不能混用。'],
             ['官方列表補充',f'官方10對×10seed也全部成功，平均e_p={stats.mean(float(r["e_p_m"]) for r in primary if tuple(r["pair"].split("-")) in official)*100:.2f}cm；不能假設這就是作者使用的列表。'],
             ['共同成功案例',f'只比兩組共同成功142案例：原推薦組e_p中位數{paired["common_ep_baseline_m"]*100:.2f}cm，補救後{paired["common_ep_recovery_m"]*100:.2f}cm。這是我們內部比較，不是與論文直接對比。'],
             ['真值使用','GT只做診斷和評分；新增補救API不接收GT。更改兩個實際案例的評分GT後，匹配與估計變換保持不變。'],
             ['原作重跑狀態','尚未取得並完成作者原實作的同機、同配對重跑，因此不宣稱嚴格超越。'],
             ['泛化限制','seed21–30仍是同一片ETH森林。Liang公開處理版沒有分站配準真值，只完成提取階段驗證。'],
             ['原始實作差異','更完整的未指定細節與選擇見docs/cn_ambiguities.md及HANDOFF.md；本表不能把未指定值當成作者參數。']]
    table(wb,'來源與限制',['項目','白話說明與出處'],sources,[30,145])
    h=list(rows[0]);raw=[]
    for r in rows:
        row=[]
        for k in h:
            value=r[k]
            try:
                n=float(value);value=n if math.isfinite(n) else None
            except ValueError:pass
            row.append(None if value=='' else value)
        raw.append(row)
    table(wb,'完整驗證原始列',h,raw,[32 if k in ['method','solver','setting','label'] else 21 for k in h])
    OUT.mkdir(parents=True,exist_ok=True);path=OUT/'paper_vs_ours.xlsx';wb.save(path)
    manifest=dict(paper_eth_ep_m=.153427,paper_eth_success=1.,ours_cases=150,ours_success=150,
                  ours_mean_ep_m=mean_ep,ours_mean_match_solve_ms=mean_ms,raw_validation_rows=600,
                  sheets=wb.sheetnames,workbook_sha256=hashlib.sha256(path.read_bytes()).hexdigest())
    (OUT/'validation.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False)+'\n')
    print(path);print(json.dumps(manifest,ensure_ascii=False))


if __name__=='__main__':main()
