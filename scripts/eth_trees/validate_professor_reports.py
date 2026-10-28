#!/usr/bin/env python3
"""Reopen actual recalculated XLSX files and audit data, formulas and chart caches."""
from pathlib import Path
import json,csv,math,hashlib,re,zipfile
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'results/professor_reports_20261026'
m=json.loads((OUT/'source_manifest.json').read_text())
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def close(a,b):return isinstance(a,(int,float)) and math.isclose(a,b,rel_tol=1e-10,abs_tol=1e-9)
# Independently reaggregate every group directly from archived scalar CSV rows.
tables={}
for sid,s in m['sources'].items():
 if 'snapshot' not in s:continue
 p=OUT/s['snapshot'];assert sha(p)==s['sha256'],sid
 if p.suffix=='.csv':tables[sid]=list(csv.DictReader(p.open()))
for gid,g in m['groups'].items():
 rr=[tables[g['source']][i-2] for i in g['csv_lines']];assert len(rr)==g['stats']['n'],gid
 assert sum(r['success']=='1' for r in rr)==g['stats']['ok'],gid
 ep=[float(r['e_p_m']) for r in rr if math.isfinite(float(r['e_p_m']))]
 assert len(ep)==g['stats']['ep_n'],gid
 assert close(sum(ep)/len(ep),g['stats']['ep']),gid
 selector=g['selector']
 for r in rr:
  for field in ['method','solver']:
   if field in selector:assert r[field]==selector[field],(gid,field)
  if 'scene' in selector:
   actual=r.get('scene') or ('Plot'+r['plot'] if 'plot' in r else 'ETH');assert actual==selector['scene'],gid
results=[];total_formulas=0;total_checks=0;chart_count=0
for i in range(1,5):
 paths=list(OUT.glob(f'report_{i:02d}_*.xlsx'));assert len(paths)==1
 p=paths[0];book=load_workbook(p);values=load_workbook(p,data_only=True)
 assert len(book.worksheets)==3 and not book._external_links
 for sheet in values:
  assert not any(c.data_type=='e' for row in sheet for c in row),(p,sheet.title)
 formulas=[c for s in book for row in s for c in row if c.data_type=='f'];assert formulas
 for s in book:
  assert all(values[s.title][c.coordinate].value is not None for row in s for c in row if c.data_type=='f')
 for item in m['validation_expectations'][str(i)]:
  actual=values[item['sheet']][item['cell']].value;expected=item['expected']
  assert close(actual,expected) if isinstance(expected,(int,float)) else actual==expected,(p.name,item,actual)
  if item['kind']=='formula':assert book[item['sheet']][item['cell']].data_type=='f'
  total_checks+=1
 allcharts=[c for s in book for c in s._charts];expected=m['charts'][str(i)];assert len(allcharts)==len(expected)
 for chart,record in zip(allcharts,expected):
  assert len(chart.series)==len(record['columns'])
  assert chart.y_axis.scaling.min==0
  for series,column in zip(chart.series,record['columns']):
   ref=series.val.numRef;assert ref and ref.numCache
   expected_ref=f'{record["source_sheet"]}!${get_column_letter(column)}${record["min_row"]}:${get_column_letter(column)}${record["max_row"]}'
   assert ref.f.replace("'",'')==expected_ref,(p.name,ref.f,expected_ref)
   plotted=[values[record['source_sheet']].cell(row,column).value for row in range(record['min_row'],record['max_row']+1)]
   assert all(isinstance(x,(float,int)) and math.isfinite(x) for x in plotted),'Missing values must never plot as zero'
   cached=[x.v for x in sorted(ref.numCache.pt,key=lambda x:x.idx)]
   assert len(cached)==len(plotted) and all(close(a,b) for a,b in zip(cached,plotted)),(p.name,record)
 text=p.with_suffix('.txt').read_text();assert '【教授可能追問' in text and '【數據追溯】' in text
 for title in ['0:00','待補','來源','相同'] :assert title in text,(p.name,title)
 for name,record in m['table_ranges'][str(i)].items():
  if name in ['success','error','params','factorial','recovery','negative','paired','fair','versions','case','cold','height','time']:
   # Not every technical table needs to be spoken, but referenced ones must be exact.
   if record['ref'] in text:assert record['sheet'] in text
 body=text.split('【教授可能追問')[0];cjk=len(re.findall('[\u3400-\u9fff]',body));assert cjk>=2100
 results.append(dict(file=p.name,sha256=sha(p),bytes=p.stat().st_size,sheets=book.sheetnames,formulas=len(formulas),charts=len(allcharts),checks=len(m['validation_expectations'][str(i)]),txt=p.with_suffix('.txt').name,txt_sha256=sha(p.with_suffix('.txt')),spoken_cjk_characters=cjk))
 total_formulas+=len(formulas);chart_count+=len(allcharts)
# Key narrative claims are checked against primary raw-data groups, not only summary headers.
for i,before,after in [(1,'reference5_raw','eth10_raw'),(3,'eth10_raw','eth10_refined')]:
 text=(OUT/results[i-1]['txt']).read_text()
 for scene in ['Plot1','Plot2','Plot3','Plot4','ETH']:
  for method in [before,after]:
   g=m['groups'][f'U:{scene}:{method}']['stats'];assert f'{g["ep"]*100:.3f}' in text
text=(OUT/results[3]['txt']).read_text()
for scene in ['Plot1','Plot2','Plot3','Plot4','ETH']:
 g=m['groups'][f'U:{scene}:multi_v3_refined']['stats'];assert f'{g["ep"]*100:.3f}' in text and g['ok']==g['n']
assert '0.127' in text and '0.733' in text
# Check every distributed file, including evidence snapshots, for private Drive IDs.
private=[]
f=ROOT/'data/tongji_trees/folder_listing.json'
if f.exists():private.extend(str(x[0]).encode() for x in json.loads(f.read_text()))
f=ROOT/'data/unified/private_sources.json'
if f.exists():private.append(json.loads(f.read_text())['google_drive_read_only_folder'].rstrip('/').split('/')[-1].split('?')[0].encode())
private=[x for x in private if x]
for path in OUT.rglob('*'):
 if not path.is_file() or 'runs' in path.relative_to(OUT).parts or path.suffix=='.zip':continue
 if path.suffix=='.xlsx':
  with zipfile.ZipFile(path) as z:content=b''.join(z.read(n) for n in z.namelist())
 else:content=path.read_bytes()
 assert not any(x in content for x in private),str(path)
 assert b'C:\\Users\\' not in content,str(path)
validation=dict(validated_at=datetime.now(ZoneInfo('Asia/Taipei')).isoformat(timespec='seconds'),reports=results,independently_reaggregated_groups=len(m['groups']),literal_and_formula_checks=total_checks,formula_count=total_formulas,charts=chart_count,formula_caches_present=True,no_excel_error_cells=True,missing_not_plotted_as_zero=True,private_ids_checked=len(private),private_ids_absent=True,raw_clouds_distributed=False,source_manifest_sha256=sha(OUT/'source_manifest.json'),note='No new registration experiment; all report data derive from completed runs. Original author baseline and untested branches remain marked 待補.')
(OUT/'validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(validation,ensure_ascii=False,indent=2))
