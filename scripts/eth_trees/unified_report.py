#!/usr/bin/env python3
"""Two-source benchmark report: failures, paper references and costs stay visible."""
import json
import math
from pathlib import Path
import statistics as st
import re
from datetime import datetime
from zoneinfo import ZoneInfo
import numpy as np
from openpyxl import Workbook
from make_study_workbook import table
import unified_study as study
from unified_study import ROOT,OUT,RUNS,SEEDS,read_csv,write_csv,sha

METHODS=['reference5_raw','reference5_refined','eth10_raw','eth10_refined','multi_refined','multi_v2_refined','reference5_refined_v3','eth10_refined_v3','multi_v3_refined']
NAMES={'reference5_raw':'5cm參考實作','reference5_refined':'5cm參考＋相同refinement','eth10_raw':'ETH固定改進版','eth10_refined':'ETH改進＋refinement','multi_refined':'多初始值v1（保留）','multi_v2_refined':'擴大候選v2（保留）','reference5_refined_v3':'5cm參考＋相同結構refinement','eth10_refined_v3':'ETH固定版＋結構refinement','multi_v3_refined':'結構驗證＋6DOF refinement（統一版）'}
PRIMARY='multi_v3_refined'
PAPER={'Plot1':.052558,'Plot2':.114808,'Plot3':.048756,'Plot4':.043560,'ETH':.153427}
def finite(values):
    out=[]
    for x in values:
        try:
            f=float(x)
            if math.isfinite(f):out.append(f)
        except (ValueError,TypeError):pass
    return out
def mean(values):return st.mean(values) if values else None
def median(values):return st.median(values) if values else None
def p90(values):return float(np.quantile(values,.9)) if values else None
def value(x):
    if x in ('',None):return None
    if isinstance(x,str) and len(x) in [40,64] and all(c in '0123456789abcdef' for c in x.lower()):return x
    try:f=float(x);return f if math.isfinite(f) else None
    except (ValueError,TypeError):return x

def collect():
    rows=[]
    for pair in study.manifest()['pairs']:
        filename=RUNS/'final_evaluation'/pair['scene']/f'{pair["source"]}-{pair["target"]}.csv'
        assert filename.exists() and Path(str(filename)+'.complete').exists(),filename
        raw=read_csv(filename);assert len(raw)==50
        newer=RUNS/'v2_evaluation'/pair['scene']/filename.name
        assert newer.exists() and Path(str(newer)+'.complete').exists(),newer
        extra=read_csv(newer);assert len(extra)==10;raw+=extra
        third=RUNS/'v3_evaluation'/pair['scene']/filename.name
        assert third.exists() and Path(str(third)+'.complete').exists(),third
        third_rows=read_csv(third);assert len(third_rows)==30;raw+=third_rows
        keyed={(int(r['label']),r['method']):r for r in raw};assert len(keyed)==90
        for seed in SEEDS:
            for method in METHODS:
                r=dict(dataset=pair['dataset'],scene=pair['scene'],pair=f'{pair["source"]}-{pair["target"]}',seed=seed,**keyed[(seed,method)])
                r['success']=int(r['success']);assert r['success']==int(r['status']=='computed' and float(r['rot_err_deg'])<2 and float(r['trans_err_m'])<.5)
                rows.append(r)
    assert len(rows)==4950;write_csv(OUT/'registration_rows.csv',rows);return rows

def summaries(rows):
    summary=[];datasets=[];seedrows=[];paired=[]
    for scene in PAPER:
        for method in METHODS:
            rr=[r for r in rows if r['scene']==scene and r['method']==method]
            ep=finite(r['e_p_m'] for r in rr);good=finite(r['e_p_m'] for r in rr if r['success'])
            pairs={r['pair'] for r in rr};pair_pass=sum(all(r['success'] for r in rr if r['pair']==pair) for pair in pairs)
            summary.append(dict(dataset=rr[0]['dataset'],scene=scene,method=method,trials=len(rr),pairs=len(pairs),success_trials=sum(r['success'] for r in rr),success_rate=sum(r['success'] for r in rr)/len(rr),pairs_all_seeds_success=pair_pass,finite_error_cases=len(ep),missing_transform=sum(r['pose_ok']!='1' for r in rr),mean_finite_ep_m=mean(ep),median_finite_ep_m=median(ep),p90_finite_ep_m=p90(ep),mean_success_ep_m=mean(good),median_success_ep_m=median(good),p90_success_ep_m=p90(good),mean_rotation_deg=mean(finite(r['rot_err_deg'] for r in rr)),mean_translation_m=mean(finite(r['trans_err_m'] for r in rr)),mean_online_ms=mean(finite(r['online_ms'] for r in rr)),median_online_ms=median(finite(r['online_ms'] for r in rr)),p90_online_ms=p90(finite(r['online_ms'] for r in rr)),mean_descriptor_ms=mean(finite(r['descriptor_ms'] for r in rr)),mean_matching_ms=mean(finite(r['matching_ms'] for r in rr)),mean_solve_ms=mean(finite(r['solve_ms'] for r in rr)),mean_refine_ms=mean(finite(r['refine_ms'] for r in rr)),mean_other_ms=mean([float(r['online_ms'])-sum(float(r[k]) for k in ['descriptor_ms','matching_ms','solve_ms','refine_ms']) for r in rr])))
            for seed in SEEDS:
                selected=[r for r in rr if r['seed']==seed];seedrows.append(dict(scene=scene,method=method,seed=seed,trials=len(selected),success=sum(r['success'] for r in selected),mean_finite_ep_m=mean(finite(r['e_p_m'] for r in selected))))
        for before,after in [('reference5_raw','reference5_refined'),('eth10_raw','eth10_refined'),('eth10_refined','multi_refined'),('reference5_refined','multi_refined'),('reference5_raw','multi_refined'),('multi_refined','multi_v2_refined'),('reference5_refined','multi_v2_refined'),('reference5_raw','multi_v2_refined'),('multi_v2_refined','multi_v3_refined'),('multi_refined','multi_v3_refined'),('eth10_refined_v3','multi_v3_refined'),('reference5_raw','reference5_refined_v3'),('reference5_refined_v3','multi_v3_refined'),('reference5_raw','multi_v3_refined')]:
            a={(r['seed'],r['pair']):r for r in rows if r['scene']==scene and r['method']==before};b={(r['seed'],r['pair']):r for r in rows if r['scene']==scene and r['method']==after};assert a.keys()==b.keys()
            shared=[(r,b[k]) for k,r in a.items() if r['success'] and b[k]['success']]
            paired.append(dict(scene=scene,before=before,after=after,trials=len(a),gained=sum(not r['success'] and b[k]['success'] for k,r in a.items()),lost=sum(r['success'] and not b[k]['success'] for k,r in a.items()),common_success=len(shared),mean_common_ep_change_m=mean([float(y['e_p_m'])-float(x['e_p_m']) for x,y in shared])))
    for dataset in ['Tongji','ETH']:
        for method in METHODS:
            rr=[r for r in rows if r['dataset']==dataset and r['method']==method];ss=[r for r in summary if r['dataset']==dataset and r['method']==method]
            for weighting in ['scene_equal','pair_weighted']:
                datasets.append(dict(dataset=dataset,method=method,weighting=weighting,scenes=len(ss),trials=len(rr),success=sum(r['success'] for r in rr),success_rate=st.mean(r['success_rate'] for r in ss) if weighting=='scene_equal' else sum(r['success'] for r in rr)/len(rr),finite_error_cases=len(finite(r['e_p_m'] for r in rr)),mean_finite_ep_m=mean(finite(r['mean_finite_ep_m'] for r in ss)) if weighting=='scene_equal' else mean(finite(r['e_p_m'] for r in rr)),interpretation='finite errors include failed cases; missing transforms remain failures and have no computable error'))
    write_csv(OUT/'scene_summary.csv',summary);write_csv(OUT/'dataset_summary.csv',datasets);write_csv(OUT/'seed_summary.csv',seedrows);write_csv(OUT/'ablation.csv',paired)
    official={tuple(line.split()) for line in (ROOT/'data/eth_trees/raw/trees/pairs.txt').read_text().splitlines() if line.strip()};assert len(official)==10
    supplemental=[]
    for method in METHODS:
        rr=[r for r in rows if r['dataset']=='ETH' and r['method']==method and tuple(r['pair'].split('-')) in official];assert len(rr)==100
        ep=finite(r['e_p_m'] for r in rr)
        supplemental.append(dict(method=method,trials=100,pairs=10,success=sum(r['success'] for r in rr),finite_error_cases=len(ep),mean_finite_ep_m=mean(ep),median_finite_ep_m=median(ep),p90_finite_ep_m=p90(ep),scope='supplement only; primary ETH benchmark remains all15 pairs'))
    write_csv(OUT/'eth_official10_summary.csv',supplemental)
    goals=[]
    for r in summary:
        if r['method']!=PRIMARY:continue
        paper=PAPER[r['scene']];allfinite=r['finite_error_cases']==r['trials'];error=r['mean_finite_ep_m']
        goals.append(dict(scene=r['scene'],trials=r['trials'],success=r['success_trials'],strict_pairs=r['pairs_all_seeds_success'],pairs=r['pairs'],finite_cases=r['finite_error_cases'],mean_ep_m=error,paper_ep_m=paper,relative_decrease=(paper-error)/paper if error is not None else None,success_goal=r['success_trials']==r['trials'],error_goal=bool(allfinite and error<paper-1e-6),ten_percent_goal=bool(allfinite and error<=.9*paper),comparison='paper values only; exact pair list/success definition not independently established'))
    write_csv(OUT/'paper_reference_comparison.csv',goals)
    return summary,datasets,seedrows,paired,goals

def costs():
    resources=[json.loads(line) for line in (RUNS/'resources.jsonl').read_text().splitlines()];last={r['label']:r for r in resources};station=[]
    for scan in study.manifest()['scans']:
        label='geometry_'+scan['scene']+'_'+scan['scan'];job=last[label]
        timing=read_csv(RUNS/'geometry'/scan['scene']/(scan['scan']+'.bin.timing.csv'))
        station.append(dict(dataset=scan['dataset'],scene=scan['scene'],scan=scan['scan'],raw_read_s=float(timing[0]['read_ms'])/1000,voxel_s=sum(float(r['voxel_ms']) for r in timing)/1000,normal_s=sum(float(r['normals_ms']) for r in timing)/1000,wall_s=job['wall_s'],cpu_s=job['cpu_s'],peak_process_mib=job['peak_rss_kib']/1024,concurrent_jobs=4,scope='one-time per station geometry cache; not per pair'))
    write_csv(OUT/'station_geometry_cost.csv',station)
    if (OUT/'feature_benchmark.json').exists():
        features=[]
        for r in json.loads((OUT/'feature_benchmark.json').read_text()):
            if 'stage_times_s' in r:stages={int(k):float(v) for k,v in r['stage_times_s'].items()}
            else:
                log=(RUNS/'logs'/(r['label']+'.log')).read_text()
                stages={int(k):float(ms)/1000 for k,ms in re.findall(r'^\s+(\d) [\w+]+\s+\d+\s+([\d.]+)',log,re.M)}
            assert set(stages)==set(range(9))
            features.append(dict(dataset=r['dataset'],scene=r['scene'],scan=r['scan'],seed=31,wall_s=r['wall_s'],cpu_s=r['cpu_s'],peak_process_mib=r['peak_rss_kib']/1024,dtm_height_s=stages[1],voxel_s=stages[2],normal_s=stages[3],verticality_s=stages[4],retention_s=stages[5],clustering_s=stages[6],cylinder_s=stages[7],position_s=stages[8],io_and_other_residual_s=r['wall_s']-sum(stages.values()),concurrent_jobs=4))
        write_csv(OUT/'station_feature_cost.csv',features)
    cold=json.loads((OUT/'cold_benchmark.json').read_text()) if (OUT/'cold_benchmark.json').exists() else []
    normalized=[dict(label=r['label'],exit_code=r['exit_code'],wall_s=r['wall_s'],cpu_s=r['cpu_s'],peak_process_mib=r['peak_rss_kib']/1024,command=json.dumps(r['command']),runtime_epoch=r.get('runtime_epoch','legacy_before_environment_restart')) for r in resources]
    write_csv(OUT/'resource_records.csv',normalized)
    return station,cold,normalized

def make_books(rows,summary,datasets,seeds,paired,goals,station,cold,resources):
    primary=[r for r in summary if r['method']==PRIMARY];numeric_success=all(g['success_goal'] for g in goals);numeric_error=all(g['error_goal'] for g in goals);ten=all(g['success_goal'] and g['ten_percent_goal'] for g in goals)
    regressions=[]
    for r in primary:
        old=next(x for x in summary if x['scene']==r['scene'] and x['method']=='multi_v2_refined')
        if r['mean_finite_ep_m'] is not None and old['mean_finite_ep_m'] is not None and r['mean_finite_ep_m']>old['mean_finite_ep_m']+1e-6:
            regressions.append(f'{r["scene"]}平均e_p比v2增加{1000*(r["mean_finite_ep_m"]-old["mean_finite_ep_m"]):.3f}毫米，成功率{old["success_trials"]}/{old["trials"]}→{r["success_trials"]}/{r["trials"]}')
    regression_note='；'.join(regressions) if regressions else '各場景平均有限e_p相對v2未增加超過0.001毫米。'
    versions_path=OUT/'versions.json'
    versions=json.loads(versions_path.read_text())
    for version in versions:
        if version['name']=='unified_v3':version['status']='full five-scene evaluation completed; preselected primary retained, both matched refinement controls included'
    versions_path.write_text(json.dumps(versions,ensure_ascii=False,indent=2)+'\n')
    conclusion=[['本輪統一方法',NAMES[PRIMARY]],['資料範圍','2個獨立來源：Tongji四個Plot、ETH一個場景；共26站、55配對，各seed31–40。'],['五場景100%成功','達成（限本輪固定評估範圍）' if numeric_success else '未達成：'+', '.join(g['scene'] for g in goals if not g['success_goal'])],['五場景平均e_p低於論文','參考數值比較達成' if numeric_error else '未達成：'+', '.join(g['scene'] for g in goals if not g['error_goal'])],['五場景各降低至少10%','參考數值比較達成' if ten else '未達成：'+', '.join(g['scene'] for g in goals if not (g['success_goal'] and g['ten_percent_goal']))],['同環境重跑原作者程式','尚未完成：作者GitHub只有README，無程式／tags／releases。5cm參考為我們實作，不冒稱作者原作。'],['相同規則跨兩資料使用','本輪所有場景使用同一份已凍結方法；推論程式不接收dataset名稱或GT。'],['未見森林的一般化','未驗證：ETH及Tongji均已參與開發；只換seed不構成新場景。'],['與論文比較限制','論文配對清單與成功率門檻未完全核對；低於報告值不能當成嚴格超越原作。'],['成功門檻','本專案自訂：rotation<2°且translation<0.5m，未稱為論文門檻。'],['誤差統計','主要表含所有可計算失敗；沒有變換仍算失敗，並列有限誤差數。成功案例另算。'],['成本解讀','每站成本另列；online不含提取／GT評分，部分採集有2程序CPU共用。嚴格速度看串行冷啟動；v3參考控制與統一版於同一新epoch比較。'],['研究完成狀態','保留本輪證據與未達成項目；原作baseline與獨立森林驗證仍缺，不能宣布整個研究完成。']]
    parameters=[['共用樹幹提取','沿用固定ETH版：DTM原設定、1cm體素、10cm法向半徑、垂直度0.90、分群20cm、階梯抽樣。未針對Plot切換DTM。'],['5cm參考','Thscore40、w10、N183、匹配容差0.05m、SVD；DTM與額外擬合檢查來自我們的實作，作者未釋出細節。'],['ETH固定版','匹配容差10cm、配準RANSAC、初始CN配對<6時證據補救，保留先前ETH設計。'],['新增refinement','完整6DOF點對平面迭代，不鎖roll/pitch或使用GT。'],['三個尺度','0.5m→0.2m→0.1m體素；最多30→25→20次迭代。'],['法向量規則','refinement半徑=2.5×體素，至少6鄰點，曲率≤0.1；兩站法向量夾角餘弦絕對值至少0.7。'],['錯配排除','最近鄰距離<3×體素、一對一目標、保留較近的70%；Huber尺度=max(0.05×體素,2×絕對殘差中位數)。'],['點數預算','每尺度來源最多20000點；雙向幾何評分各最多8000點；均用確定性取樣。'],['多初始值觸發','原CN配對<6時，取前20描述子候選、最多120證據邊、7140組水平雙點種子及10000相容三元組。'],['初始值選擇','保留最多256個相異姿態，加既有粗配準；既有姿態保留於最多8個refinement候選，再依雙向幾何及全樹位XY支持選擇。'],['選擇分數','雙向overlap/(0.03+最近鄰RMSE)乘log(1+全樹位XY支持數)；這是推論分數，不是e_p。'],['水平種子限制','僅初始值使用直立掃描假設；最終refinement可估roll/pitch及高度位移，沒有鎖成4DOF。'],['結構驗證修正','v3只用|normal.z|<0.3的垂直表面驗證，至少各30點，否則退回所有方向；驗證距離min(體素,0.2m)，降低水平地面偏差。'],['公平refinement','v3的5cm參考與ETH固定版都加上相同結構refinement；優化仍使用所有表面方向。'],['精度代價','refinement增加幾何快取、法向量與逐配對計算；不宣稱同時比原粗配準更快。']]
    evidence=[['第一次ETH→Tongji','保留fixed_eth_20261025：ETH選定補救在Tongji329/400，未看Tongji前凍結；未覆寫。'],['Tongji內部DTM','Plot2開發後在其餘Plot定位測試；已參與本輪方法假設，不再稱未見。'],['本輪統一benchmark','同一方法用於Tongji/ETH；兩資料均已開發接觸，可證benchmark改善，不能抹去接觸歷史。'],['完全獨立資料','目前沒有；WHU-Park只作額外診斷，不當作原論文WHU-FGI，本輪不混入它的分數。'],['原作者baseline','https://github.com/MercerSU/CircularMatch；commit1e9623f目前僅README，缺少可建置程式。'],['兩站限制','每次native inference僅讀來源與目標快取及其樹位；沒有GT參數、其他站圖或跨站地圖。'],['seed解讀','55配對重複10次，不是550個獨立場景；同時列所有seed均成功的唯一配對數。']]
    failures=[r for r in rows if r['method']==PRIMARY and not r['success']]
    def common(book,simple=False):
        simple_conclusion=[
            ['測了什麼','Tongji四個Plot＋ETH；共55組兩站配對，每組測10個固定seed。'],
            ['是否全部成功','五個場景全部成功（限本輪測試）' if numeric_success else '仍有失敗：'+', '.join(g['scene'] for g in goals if not g['success_goal'])],
            ['誤差有沒有更小','五個場景平均誤差都低於論文報告值' if numeric_error else '還沒全部低於論文：'+', '.join(g['scene'] for g in goals if not g['error_goal'])],
            ['有沒有各降低10%','五個場景均成功且各降低至少10%' if ten else '尚未全部達成，請看「論文報告值並列」。'],
            ['怎麼看數字','e_p越小越好，這份表用公分；平均保留能計算的失敗，無輸出另算失敗。'],
            ['改善的代價','增加點雲幾何處理、候選搜尋及最後微調；完整時間和記憶體在完整版。'],
            ['能說超越論文嗎','目前只能比較報告數字。作者未釋出程式，還不能完成同環境原作對照；表中的5cm參考是我們的實作。'],
            ['跨資料集做到哪裡','兩組資料共用同一設定；兩組都已參與開發，還需要另一個沒看過的森林驗證。'],
            ['成功怎麼算','本專案：旋轉誤差小於2度且平移誤差小於50公分；論文門檻未查明，配對清單也未完全核對。']]
        table(book,'先看結論',['項目','說明'],simple_conclusion if simple else conclusion,[30,100] if simple else [32,145])
        chosen=[r for r in summary if not simple or r['method'] in ['reference5_raw','reference5_refined_v3',PRIMARY]]
        ws=table(book,'逐場景比較',['資料集','場景','方法','成功次數','試驗次數','成功率','所有seed皆成功配對','配對數','有限誤差案例數','平均e_p cm（含可算失敗）','中位e_p cm','P90 e_p cm','成功案例平均e_p cm','平均online ms'],[[r['dataset'],r['scene'],NAMES[r['method']],r['success_trials'],r['trials'],r['success_rate'],r['pairs_all_seeds_success'],r['pairs'],r['finite_error_cases'],100*r['mean_finite_ep_m'] if r['mean_finite_ep_m'] is not None else None,100*r['median_finite_ep_m'] if r['median_finite_ep_m'] is not None else None,100*r['p90_finite_ep_m'] if r['p90_finite_ep_m'] is not None else None,100*r['mean_success_ep_m'] if r['mean_success_ep_m'] is not None else None,r['mean_online_ms']] for r in chosen],[20,18,48,18,18,20,30,18,26,36,24,24,32,28])
        if simple:
            book.remove(ws);compact=[]
            for scene in PAPER:
                rr={r['method']:r for r in summary if r['scene']==scene}
                original,control,new=rr['reference5_raw'],rr['reference5_refined_v3'],rr[PRIMARY]
                cm=lambda r:100*r['mean_finite_ep_m'] if r['mean_finite_ep_m'] is not None else None
                compact.append(['Tongji / '+scene if scene!='ETH' else 'ETH-tree',f'{original["success_trials"]}/{original["trials"]}',cm(original),f'{control["success_trials"]}/{control["trials"]}',cm(control),f'{new["success_trials"]}/{new["trials"]}',cm(new),100*PAPER[scene]])
            table(book,'逐場景比較',['來源／場景','5cm參考成功','參考平均誤差cm','參考＋微調成功','微調後平均誤差cm','統一版成功','統一版平均誤差cm','論文報告誤差cm'],compact,[19,16,22,22,24,18,24,22])
        else:
            for row in ws.iter_rows(min_row=2):row[5].number_format='0.0%'
        ws=table(book,'論文報告值並列',['場景','成功次數','試驗次數','本輪平均e_p m','論文報告e_p m','降低比例','100%門檻','誤差門檻','10%目標','限制'],[[g['scene'],g['success'],g['trials'],g['mean_ep_m'],g['paper_ep_m'],g['relative_decrease'],'達成' if g['success_goal'] else '未達成','達成' if g['error_goal'] else '未達成','達成' if g['success_goal'] and g['ten_percent_goal'] else '未達成','論文配對／成功門檻未完全確認；數值參考比較'] for g in goals],[18,18,18,30,30,24,22,22,22,100])
        if simple:
            book.remove(ws)
            ws=table(book,'論文報告值並列',['場景','成功／試驗','統一版平均誤差cm','論文誤差cm','降低比例','100%且降低10%'],[[g['scene'],f'{g["success"]}/{g["trials"]}',100*g['mean_ep_m'] if g['mean_ep_m'] is not None else None,100*g['paper_ep_m'],g['relative_decrease'],'達成' if g['success_goal'] and g['ten_percent_goal'] else '未達成'] for g in goals],[18,18,25,20,18,25])
            for row in ws.iter_rows(min_row=2):row[4].number_format='0.0%'
        else:
            for row in ws.iter_rows(min_row=2):row[5].number_format='0.0%'
        if simple:
            by_scene={scene:{r['method']:r for r in summary if r['scene']==scene} for scene in PAPER}
            def successes(method):
                return '；'.join(f'{scene} {by_scene[scene][method]["success_trials"]}/{by_scene[scene][method]["trials"]}' for scene in PAPER)
            simple_parameters=[
                ['樹幹提取','本輪所有版本用相同樹位輸入；沒有替不同Plot換提取參數。'],
                ['粗配準組合','從5cm匹配容差＋SVD，改為10cm＋RANSAC及弱配對補救；這些是一起改，不能把效果全歸給10cm。'],
                ['粗配準效果','5cm：'+successes('reference5_raw')+'。10cm組合：'+successes('eth10_raw')+'。'],
                ['最後微調','新增由粗到細的點對平面微調（50→20→10cm尺度），能修正旋轉與平移；5cm參考也加相同微調作公平比較。'],
                ['擴大候選曾變壞','弱配對時多試幾個起始位置；Plot4由v1的96/100退到v2的78/100，原因是錯誤位置的地面也可能重疊。'],
                ['本輪修正','用接近垂直的表面驗證位置，減少地面誤導；最後微調仍使用各方向的表面。Plot4本輪為'+str(by_scene['Plot4'][PRIMARY]['success_trials'])+'/100。'],
                ['本輪仍有的誤差退步',regression_note],
                ['哪些變好、哪些變壞','「逐場景比較」並列參考、參考加相同微調、統一版；完整消融保留每次新增成功與退步，沒有挑最佳seed。']]
            table(book,'改了哪些參數',['改動','白話說明與結果'],simple_parameters,[30,100])
        else:table(book,'改了哪些參數',['項目','目的與設定'],parameters,[32,145])
    simple=Workbook();simple.remove(simple.active);common(simple,True);simple.save(OUT/'comparison_simple.xlsx')
    full=Workbook();full.remove(full.active);common(full)
    ws=table(full,'Tongji與ETH彙整',['資料集','方法','彙整方式','場景數','成功次數','試驗次數','成功率','有限誤差案例數','平均有限e_p cm'],[[r['dataset'],NAMES[r['method']],'場景等權' if r['weighting']=='scene_equal' else '配對加權',r['scenes'],r['success'],r['trials'],r['success_rate'],r['finite_error_cases'],100*r['mean_finite_ep_m'] if r['mean_finite_ep_m'] is not None else None] for r in datasets],[22,48,26,18,18,18,20,26,32])
    for row in ws.iter_rows(min_row=2):row[6].number_format='0.0%'
    table(full,'消融實驗',['場景','原方法','新方法','新增成功','退步失敗','兩法皆成功','平均e_p變化cm（負數較好）'],[[r['scene'],NAMES[r['before']],NAMES[r['after']],r['gained'],r['lost'],r['common_success'],100*r['mean_common_ep_change_m'] if r['mean_common_ep_change_m'] is not None else None] for r in paired],[18,48,48,20,20,24,38])
    costrows=[]
    for r in station:
        for name,key in [('原始讀檔','raw_read_s'),('體素前處理','voxel_s'),('refinement法向量','normal_s'),('幾何快取總計','wall_s')]:costrows.append([r['dataset'],r['scene'],r['scan'],name,r[key],r['peak_process_mib'] if name=='幾何快取總計' else None,'每站一次，4個並行工作；子程序RSS含載入原始點雲'])
    for r in summary:
        for name,key in [('descriptor','mean_descriptor_ms'),('matching','mean_matching_ms'),('求粗變換','mean_solve_ms'),('refinement','mean_refine_ms'),('候選及未細分開銷','mean_other_ms'),('配對online總計','mean_online_ms')]:costrows.append([r['dataset'],r['scene'],NAMES[r['method']],name,r[key]/1000,None,'逐配對／seed平均；不含一次性站點成本及GT評分'])
    for r in cold:costrows.append([r['dataset'],r['scene'],NAMES[r['method']],'從原始點雲冷啟動',r['wall_s'],r['peak_child_rss_kib']/1024,'seed31固定配對；重提取兩站；OS快取未重置；epoch='+r.get('runtime_epoch','legacy_before_environment_restart')])
    for r in resources:
        if r['label'].startswith('evaluate_'):
            parts=r['label'].split('_');scene=parts[2]
            costrows.append(['ETH' if scene=='ETH' else 'Tongji',scene,parts[-1],'GT全來源點評分（非推論）',r['wall_s'],r['peak_process_mib'],'包含評分讀原始來源點；另列成本，不與論文matching時間比較'])
    if (OUT/'station_feature_cost.csv').exists():
        for r in read_csv(OUT/'station_feature_cost.csv'):
            for name,key in [('提取DTM及高度切片','dtm_height_s'),('提取體素','voxel_s'),('提取法向量','normal_s'),('提取垂直度','verticality_s'),('距離抽樣','retention_s'),('分群','clustering_s'),('圓柱擬合','cylinder_s'),('地面交點','position_s'),('讀檔及未細分開銷','io_and_other_residual_s'),('單站完整提取seed31','wall_s')]:costrows.append([r['dataset'],r['scene'],r['scan'],name,float(r[key]),float(r['peak_process_mib']) if key=='wall_s' else None,'全26站重測一次seed31；4並行／每程序4threads；其他seed需另加抽樣至交點成本'])
    table(full,'時間與記憶體',['資料集','場景','站或方法','階段','秒','峰值子程序MiB','成本範圍'],costrows,[20,18,48,32,24,30,125])
    budget=[]
    for title,prefix in [('幾何快取','geometry_'),('v1準確度採集','final_'),('v2準確度採集','v2_'),('v3準確度採集','v3_'),('GT評分','evaluate_'),('冷啟動對照','cold_'),('26站提取成本','feature_benchmark_')]:
        rr=[r for r in resources if r['label'].startswith(prefix)]
        if rr:budget.append([title,len(rr),sum(r['wall_s'] for r in rr),sum(r['cpu_s'] for r in rr),max(r['peak_process_mib'] for r in rr),'wall為各工作加總，並行時不等於整輪經過時間；包含實際重複執行，不增加準確度樣本數'])
    table(full,'本輪計算預算',['階段','程序次數','程序wall加總秒','CPU加總秒','最大子程序MiB','解讀'],budget,[30,20,30,26,28,135])
    table(full,'一般化證據與限制',['證據','解讀'],evidence,[32,145])
    official=read_csv(OUT/'eth_official10_summary.csv')
    table(full,'ETH官方10補充',['方法','成功試驗','試驗數','配對數','有限誤差案例數','平均有限e_p cm','中位e_p cm','P90 e_p cm','用途'],[[NAMES[r['method']],int(r['success']),100,10,int(r['finite_error_cases']),100*float(r['mean_finite_ep_m']) if r['mean_finite_ep_m'] else None,100*float(r['median_finite_ep_m']) if r['median_finite_ep_m'] else None,100*float(r['p90_finite_ep_m']) if r['p90_finite_ep_m'] else None,'補充；主要ETH仍全部15對，未以此刪除困難配對'] for r in official],[48,22,18,18,26,30,24,24,110])
    if failures:table(full,'失敗分析',['場景','配對','seed','狀態','旋轉度','平移m','e_p m','候選數','幾何overlap'],[[r['scene'],r['pair'],r['seed'],r['status'],value(r['rot_err_deg']),value(r['trans_err_m']),value(r['e_p_m']),value(r['proposal_count']),value(r['geo_overlap'])] for r in failures],[18,25,18,28,24,24,24,24,26])
    else:table(full,'失敗分析',['項目','說明'],[['主要方法','本輪550次均通過專案門檻；有限測試不保證所有森林。'],['原始失敗','其他版本失敗與有限誤差均保留於全部逐次結果；消融列出新增／退步。'],['相對v2的平均誤差退步',regression_note]],[32,145])
    for title,records in [('全部逐次結果',rows),('每seed變異',seeds),('資源原始紀錄',resources),('固定配對清單',study.manifest()['pairs'])]:
        h=list(records[0]);table(full,title,h,[[value(r.get(k,'')) for k in h] for r in records],[38 if k in ['method','command','gt'] else 26 for k in h])
    table(full,'欄位說明',['欄位','說明'],[['e_p_m','Eq8：全部來源點在估計與GT變換下位置差的平均；不是ICP fitness或最近鄰。'],['geo_overlap / geo_rmse','推論幾何評分；v3使用垂直結構、v1/v2用全體表面，不宜直接比數值；與GT無關，不能取代e_p。'],['status','computed可計算變換誤差；no_transform沒有有效輸出仍算失敗。'],['mean_finite_ep','包含可計算失敗；沒有變換者無法計算，須同時核對finite_error_cases。'],['online_ms','原生配對推論時間，不含快取讀檔及站點預處理；其中未細分的提案生成包含在總計。'],['peak_process_mib','Linux wait4最大子程序RSS；並行快取時不等於整台機器峰值；不含OS page cache。'],['matching / refinement公平性','baseline亦有同一refinement版本；先比較成功率，再比較共同成功案例誤差。'],['原作baseline','尚缺作者程式；reference5是參考實作，不可標為嚴格原作重現。']],[42,145])
    full.save(OUT/'comparison_full.xlsx')
    validation=dict(trials=len(rows),unique_pairs=55,datasets=2,scenes=5,seeds=SEEDS,numeric_success_all_scenes=numeric_success,numeric_error_below_paper_all_scenes=numeric_error,ten_percent_all_scenes=ten,strict_author_baseline_verified=False,unseen_forest_generalization_verified=False,code_commit=json.loads((OUT/'method_v3_frozen.json').read_text())['code_commit'],cold_benchmark_cases=len(cold),files={name:sha(OUT/name) for name in ['registration_rows.csv','scene_summary.csv','dataset_summary.csv','ablation.csv','paper_reference_comparison.csv','comparison_simple.xlsx','comparison_full.xlsx']})
    (OUT/'validation.json').write_text(json.dumps(validation,indent=2)+'\n')
    scene_lines=[]
    for g in goals:
        formatted=f'{g["mean_ep_m"]:.6f}' if g['mean_ep_m'] is not None else '無有限誤差'
        decrease=f'{100*g["relative_decrease"]:.1f}%' if g['relative_decrease'] is not None else '—'
        scene_lines.append(f'| {g["scene"]} | {g["success"]}/{g["trials"]} | {g["strict_pairs"]}/{g["pairs"]} | {formatted} | {g["paper_ep_m"]:.6f} | {decrease} |')
    text='''# 統一森林配準benchmark

本輪只計Tongji-trees（四Plot）及ETH-tree兩個獨立來源。全部55個固定配對、seed31–40；每次推論只使用當前兩站，不讀GT、不用其他站圖。兩資料均已參與開發，不能稱未見森林驗證。

主要統一版本為v3：保留ETH粗配準，弱CN時擴大有限候選，使用垂直結構驗證降低水平地面錯配偏差，再做完整6DOF點對平面refinement。5cm參考與ETH固定版也套相同v3 refinement作公平對照。沒有依Plot或dataset切換方法。

## 逐場景主要結果

| 場景 | 成功試驗 | 所有seed皆成功配對 | 平均有限e_p m | 論文報告e_p m | 降低比例 |
|---|---:|---:|---:|---:|---:|
'''+ '\n'.join(scene_lines)+f'''

成功門檻為本專案rotation<2°且translation<0.5m，未聲稱是論文門檻。e_p依Eq8計算全部來源點；有限失敗也納入平均，無變換仍計失敗且另列無法計算誤差數。550次是55個配對重複10次，不能當作550個獨立場景。

本輪五場景全成功：{numeric_success}；五場景誤差數值低於論文：{numeric_error}；五場景各降低至少10%且全成功：{ten}。這些是固定benchmark的數值判定，論文配對及成功率門檻未完全核對。

## 尚未完成的研究驗收

- 作者原作baseline未重跑：論文指定 https://github.com/MercerSU/CircularMatch 的完整history僅README，無程式、tags、releases或forks。5cm參考是我們實作，不冒稱未修改作者原作。
- ETH及全部Tongji已參與開發／失敗分析；沒有完全獨立未見森林。相同固定方法的兩來源benchmark改善，不等同更廣泛一般化。
- 不以整體平均掩蓋場景退步；所有九版本逐場景及逐案例保留。

## 版本與消融

v1在Plot4為96/100；v2擴大候選後退到78/100，說明候選多而幾何評分偏好重疊地面會造成退步。v3改結構驗證，不改全6DOF求解。全55配對均重新評估，沒有只補失敗。

相對v2仍有的平均誤差退步：{regression_note}。統一版保留此結果，沒有為該場景換回其他版本。

第一個ETH凍結後的Tongji遷移仍在 `../tongji_trees/fixed_eth_20261025`；DTM定位／完整Plot2追測仍在 `../tongji_trees/dtm_height_20261025`。本輪未覆寫這些紀錄。配置、程式hash與各候選預算見protocol、method_frozen、method_v2_frozen、method_v3_frozen、unified_config及versions。

## 報告及成本

- comparison_simple.xlsx：白話結論、五場景、論文數值與參數。
- comparison_full.xlsx：全部4950列、消融、Tongji等權／配對加權、ETH單列、時間／記憶體、失敗、seed變異、一般化限制。
- registration_rows.csv與scene_summary.csv保留可計算失敗；status及finite_error_cases避免把缺輸出當零誤差。
- v3的相同refinement控制與統一版冷啟動在同一新runtime epoch比較。舊epoch時間只作歷史紀錄；並行準確度採集online時間不是獨占速度基準。
- 每站幾何快取與seed31完整提取另量測，不當作免費；12筆冷啟動重建兩站應用快取，OS檔案快取未重置並反向排序。記憶體為最大子程序RSS，不含Python父程序及OS page cache。
- GT全點評分成本獨立列出，不混入推論；不同硬體的論文matching-only時間不作直接比較。

## 續跑與重現

先看RESUME.md。原始點雲及私人Drive識別碼不隨Git分發；manifest記錄26站、55配對、checksum、方向、單位與開發接觸歷史。checkpoint可恢復完整及部分seed，SHA先核驗。環境重啟會殺背景程序，需先核對活程序再啟動resume_unified.py。

核驗文件為evaluator_audit、evaluator_full_cloud_audit、matrix_direction_audit、validation及workbook_validation。DELIVERY.json提供核對過下載SHA的固定commit連結。FINAL_COMPLETE只標示本輪計算及交付結束，不代表整個研究已驗收。

報表生成時間：{datetime.now(ZoneInfo('Asia/Taipei')).isoformat(timespec='seconds')}（Asia/Taipei）。
'''
    (OUT/'README.md').write_text(text)
    log=OUT/'RESEARCH_LOG.md'
    if log.exists():
        content=log.read_text().split('\n## 本輪完成節點')[0]
        content=content.replace('結果仍待完成。','本輪結果見下方完成節點，規則未再更換。')
        content=content.replace('五場景全成功與誤差目標尚待v3驗證；','五場景數值目標結果見下方完成節點；')
        content+='\n## 本輪完成節點\n\n'+ '\n'.join(f'- {g["scene"]}：{g["success"]}/{g["trials"]}成功，平均有限e_p={g["mean_ep_m"]}m。' for g in goals)+f'\n\n數值全成功={numeric_success}，五場景低於論文={numeric_error}，各降低10%={ten}；嚴格原作者baseline與未見森林驗證仍未完成。\n'
        log.write_text(content)
    print(json.dumps(goals,indent=2),flush=True)

def main():
    rows=collect();summary,datasets,seeds,paired,goals=summaries(rows);station,cold,resources=costs()
    assert len(cold)==12,'Cold benchmarks incomplete; keep CSV summaries but do not publish final workbooks yet'
    assert (OUT/'station_feature_cost.csv').exists() and len(read_csv(OUT/'station_feature_cost.csv'))==26,'Station cost measurements incomplete'
    make_books(rows,summary,datasets,seeds,paired,goals,station,cold,resources)

if __name__=='__main__':main()
