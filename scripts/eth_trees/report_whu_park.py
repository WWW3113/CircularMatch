#!/usr/bin/env python3
"""Report unchanged transfer and a separately labeled density-adaptation attempt."""
import hashlib
import json
import math
import re
import statistics as stats

from openpyxl import Workbook
from make_study_workbook import table
from tune_registration import ROOT, read_csv, write_csv
from transfer_whu_park import OUT, pairs


def number(value):
    if value in ['',None]:return None
    try:
        n=float(value);return n if math.isfinite(n) else None
    except ValueError:return value


def main():
    plan=json.loads((OUT/'plan.json').read_text());adapt=json.loads((OUT/'density_adaptation.json').read_text())
    cases=[];extract=[];summary=[]
    for preprocessing,base in [('固定ETH參數',OUT/'runs'),('密度適配探索',OUT/'runs/adapted')]:
        rr=read_csv(base/'registration_rows.csv');assert len(rr)==1240
        assert len({(r['method'],r['solver'],r['seed'],r['pair']) for r in rr})==1240
        cases.extend(dict(preprocessing=preprocessing,**r) for r in rr)
        for method in ['baseline','recovery']:
            for solver in ['svd','svd+ransac']:
                rows=[r for r in rr if r['method']==method and r['solver']==solver];assert len(rows)==310
                successful=[float(r['e_p_m']) for r in rows if r['success']=='1']
                summary.append(dict(preprocessing=preprocessing,method=method,solver=solver,cases=len(rows),
                                    success=sum(int(r['success']) for r in rows),success_rate=sum(int(r['success']) for r in rows)/len(rows),
                                    transforms_found=sum(int(r['transform_ok']) for r in rows),
                                    mean_ep_success_m=stats.mean(successful) if successful else '',
                                    mean_match_solve_ms=stats.mean(float(r['time_ms']) for r in rows)))
        for scan in range(32):
            text=(base/'logs'/f'extract_{scan}.log').read_text()
            def stage(i,label):return int(re.search(rf'{i} {label}\s+(\d+)',text).group(1))
            counts=[]
            for seed in range(21,31):
                trees=read_csv(base/'extraction'/str(scan)/f'step_seed{seed}_trees.csv');counts.append(len(trees))
                assert all(math.isfinite(float(r[k])) for r in trees for k in ['x','y','z'])
                extract.append(dict(preprocessing=preprocessing,scan=scan,seed=seed,
                                    input_points=stage(0,'input'),height_points=stage(1,r'dtm\+height'),
                                    normal_valid_points=stage(3,'normals'),vertical_points=stage(4,'verticality'),trees=len(trees)))
            assert counts==list(map(int,re.findall(r'trees: (\d+) ;',text)))
    assert len(cases)==2480 and len(extract)==640
    primary=[r for r in summary if r['solver']=='svd+ransac']
    assert all(r['success']==(0 if r['preprocessing']=='固定ETH參數' else 2) for r in primary)
    write_csv(OUT/'summary.csv',summary);write_csv(OUT/'extraction_rows.csv',extract)
    source_checks=[
        ['WHU-Park','已取得並執行','https://github.com/HpWang-whu/YOHO','32站、31個原始GT配對；不是FGI-Forest。'],
        ['Tongji-Trees','原始資料仍需官方申請','https://github.com/zexinyang/GlobalMatch','公開README提供申請表，尚無可匿名直接下載的原始資料。'],
        ['WHU-FGI','未取得該森林子集','https://github.com/WHU-USI3DV/WHU-TLS','已下載的YOHO版只有其他六個場景，沒有FGI。'],
        ['Liang / TreeTool','處理後六樣區已提取驗證','https://github.com/porteratzo/TreeTool','只有每樣區一份PCD與參考樹，缺少分站配準真值。'],
        ['FinnWoodlands','找到公開資料連結，未下載','https://github.com/juanb09111/FinnForest','README主要提供分割標註與點雲，尚未確認可用分站姿態／配準真值。'],
        ['Lin3D','找到公開資料連結，未下載','https://github.com/bjfu-lidar/Lin3D-Large-scale-forest-scene-INterpretation-3D-point-cloud-dataset','README為森林語義分割的xyz+label，未提供配準GT說明。'],
        ['Treeseg / 3D Forest','找到另一個TLS來源，未下載','https://github.com/chaofeng2007/Treeseg-TLS-data','連到SourceForge的3D Forest資料，尚未確認分站與姿態，不能當成已可配準的森林資料。']]
    wb=Workbook();wb.remove(wb.active)
    front=[['這是什麼資料','WHU-TLS的Park公園場景，公開YOHO整理版；不是原論文FGI森林。'],
           ['主要比較','原推薦組與補救組都使用配準RANSAC；SVD另保留於完整逐次資料。'],
           ['有多少資料','32站，31個提供的真值配對，每對跑seed21–30，共310次／方法。'],
           ['固定ETH參數','原推薦組0/310；加三維補救也是0/310。'],
           ['直接遷移卡在哪','32站的10cm鄰域都沒有足夠鄰居算有效法向量，後面無法提取樹木。'],
           ['點間距','跨站等權的最近鄰距離中位數約11.15cm，第五近鄰約16.85cm。'],
           ['另外試了什麼',f'只依輸入點間距規則，把法向量半徑改為{adapt["normal_radius_m"]*100:.0f}cm；分群距離仍為{adapt["cluster_tolerance_m"]*100:.0f}cm。未用GT選這些值。'],
           ['適配後結果','原推薦組2/310；加三維補救也是2/310（0.6%）。有效法向量恢復，但樹木提取數仍太少。'],
           ['能不能說跨資料集成功','不能。目前ETH的好結果沒有可靠遷移到這份公開公園資料。'],
           ['這輪是否重新調匹配','沒有。匹配閾值10cm、配準RANSAC及少於6對啟動補救均維持ETH固定設定。'],
           ['原點限制','使用檔案座標原點0；未獨立驗證實際掃描中心。這與ETH的原點檢查條件不同。'],
           ['探索性適配限制','適配是在固定參數失敗後提出，依輸入幾何選值；不是原封不動遷移，也不是另留獨立資料驗證。'],
           ['目前最有用的下一步','找原始密度、具有分站及真值的森林資料；或另立計畫研究密度適配與稀疏樹木提取。']]
    table(wb,'先看結果',['項目','白話說明'],front,[28,140])
    eth=read_csv(ROOT/'results/eth_trees/failure_followup_20261025/runs/validation_rows.csv')
    comparison=[]
    for method in ['baseline','recovery']:
        rr=[r for r in eth if r['method']==method and r['solver']=='svd+ransac']
        comparison.append(['ETH原始Trees','ETH固定設定',method,len(rr),sum(int(r['success']) for r in rr),sum(int(r['success']) for r in rr)/len(rr),'同一ETH森林的新seed驗證'])
    for r in summary:
        if r['solver']=='svd+ransac':comparison.append(['WHU-Park公開版',r['preprocessing'],r['method'],r['cases'],r['success'],r['success_rate'],'混合公園場景；不是論文FGI'])
    ws=table(wb,'ETH與Park比較',['資料集','前處理','匹配方法','案例數','成功數','成功率','範圍'],comparison,[25,27,24,15,15,18,60])
    for i in range(2,ws.max_row+1):ws.cell(i,6).number_format='0.0%'
    pairrows=[]
    for a,b,_ in pairs():
        p=f'{a}-{b}';values=[]
        for pre,method in [('固定ETH參數','baseline'),('固定ETH參數','recovery'),('密度適配探索','baseline'),('密度適配探索','recovery')]:
            rr=[r for r in cases if r['preprocessing']==pre and r['method']==method and r['solver']=='svd+ransac' and r['pair']==p]
            assert len(rr)==10;values.append(sum(int(r['success']) for r in rr))
        pairrows.append([p,*values])
    table(wb,'Park各配對',['source→target','固定參數原組／10','固定參數補救／10','密度適配原組／10','密度適配補救／10'],pairrows,[25,30,30,30,30])
    stages=[]
    for pre in ['固定ETH參數','密度適配探索']:
        for scan in range(32):
            rr=[r for r in extract if r['preprocessing']==pre and r['scan']==scan]
            stages.append([pre,scan,rr[0]['input_points'],rr[0]['height_points'],rr[0]['normal_valid_points'],rr[0]['vertical_points'],stats.mean(r['trees'] for r in rr),min(r['trees'] for r in rr),max(r['trees'] for r in rr)])
    table(wb,'提取卡在哪裡',['前處理','站點','原始點數','離地0–3m點數','有效法向量點數','垂直候選點數','10seed平均樹數','最少樹數','最多樹數'],stages,[25,12,22,24,25,25,25,20,20])
    sources=[['原始公開來源',plan['source']],['公開下載',plan['download']],['壓縮檔SHA256',plan['archive_sha256']],
             ['GT方向','utils/dataset.py明寫target=id0、source=id1；GT矩陣將source轉到target。本輪也用前三對的點雲重疊檢查方向，未刪選配對。'],
             ['密度適配規則',adapt['rule']],['採用法向半徑m',adapt['normal_radius_m']],['採用分群距離m',adapt['cluster_tolerance_m']],
             ['成功標準','旋轉誤差<2°且平移誤差<0.5m；沿用ETH，未因Park改動。'],
             ['誤差空白','無法求出變換時沒有e_p；空白不是0。僅2個成功案例的平均誤差不能代表整體配準品質。'],
             ['探索標示',adapt['phase']],['物理原點驗證',False]]
    table(wb,'參數與來源',['項目','內容'],sources,[35,140])
    table(wb,'其他森林來源',['資料','目前狀態','來源','是否能做配準'],source_checks,[26,39,100,100])
    h=list(cases[0]);table(wb,'全部逐次結果',h,[[number(r.get(k)) for k in h] for r in cases],[32 if k in ['preprocessing','method','solver','setting','label'] else 21 for k in h])
    path=OUT/'cross_dataset_park.xlsx';wb.save(path)
    validation=dict(scans=32,provided_gt_pairs=31,seeds=list(range(21,31)),extractions_checked=640,
                    registration_rows=2480,zero_shot_primary_success=0,adapted_primary_success=2,primary_cases_per_method=310,
                    no_gt_parameter_tuning=True,physical_origin_independently_verified=False,
                    workbook_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),sheets=wb.sheetnames)
    (OUT/'validation.json').write_text(json.dumps(validation,indent=2)+'\n')
    lines=['# WHU-Park 外部場景測試','',
           '資料為YOHO專案提供的WHU-TLS公開整理版Park，不是原論文WHU-FGI。32站、31個提供的GT配對，每對使用seed21–30。兩種前處理×原推薦組／三維補救×SVD／配準RANSAC，共2480列；640次提取均核對樹位CSV與日誌。','',
           '## 結果','',
           'ETH參數直接遷移：兩種匹配方法皆0/310成功。32站有效法向量都是0，後面無法提取樹木，三維匹配補救沒有輸入。','',
           '診斷：點的最近鄰距離中位數約0.1115m，第五近鄰約0.1685m；10cm鄰域不足以取得PCA所需的5個鄰居。','',
           '另作探索性密度適配：只用輸入點間距的固定規則，法向半徑升至0.20m；分群距離仍0.20m，其他參數不改，且未用GT選值。法向量恢復，但樹木提取依然很少；原推薦組與三維補救皆2/310成功（0.6%）。沒有在看到適配結果後繼續搜尋參數。','',
           '因此目前沒有可靠的跨資料集配準效果。這份資料為處理後混合公園場景，實際掃描中心也未獨立驗證；不能用本次結果聲稱原始森林資料的泛化表現。','',
           '## 來源與真值','',
           f'公開來源：{plan["source"]}；壓縮檔SHA256：{plan["archive_sha256"]}。只提取既有Park點雲與gt.log/gt.info。原始GT方向依YOHO utils/dataset.py：header target,source,N，矩陣把source映至target。前三對的原始點雲重疊亦確認此方向，未依GT刪選配對。','',
           '另找到FinnWoodlands、Lin3D及Treeseg/3D Forest公開來源，但目前README沒有確認可直接使用的分站配準真值；Tongji仍為申請表、公開WHU整理版仍不含FGI；Liang處理版僅能作提取驗證。','',
           '## 重現','',
           '依序執行transfer_whu_park.py prepare、extract、register，再執行density、extract --adapted、register --adapted，最後執行report_whu_park.py。prepare要求data/whu_tls_yoho/extracted/WHU-TLS/Park/PointCloud已有核驗壓縮檔內的32份點雲與GT。密度適配計畫在試跑前保存於density_adaptation.json，固定ETH參數計畫保存於plan.json。','',
           '[跨資料集比較Excel](cross_dataset_park.xlsx)含主要結果、ETH與Park比較、31個配對、提取瓶頸、參數來源與2480筆原始逐次結果。']
    (OUT/'README.md').write_text('\n'.join(lines)+'\n')
    print(path);print(json.dumps(summary,indent=2))


if __name__=='__main__':main()
