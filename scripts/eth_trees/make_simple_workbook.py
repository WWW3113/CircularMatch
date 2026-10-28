#!/usr/bin/env python3
"""Three small, plain-language sheets derived from the published study results."""
import csv
from pathlib import Path
import statistics as stats

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'results/cross_dataset_20261025'


def csv_rows(path):
    with path.open() as f:
        return list(csv.DictReader(f))


def sheet(wb, title, explanation, headers, rows, widths):
    ws = wb.create_sheet(title)
    end = get_column_letter(len(headers))
    ws.merge_cells(f'A1:{end}1'); ws['A1'] = title
    ws['A1'].font = Font(name='Arial', size=15, bold=True)
    ws.row_dimensions[1].height = 28
    ws.merge_cells(f'A2:{end}3'); ws['A2'] = explanation
    ws['A2'].font = Font(name='Arial', size=11)
    ws['A2'].alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[2].height = 22; ws.row_dimensions[3].height = 22
    for index, row in enumerate([headers, *rows], 4):
        for column, value in enumerate(row, 1): ws.cell(index, column, value)
    for c in ws[4]:
        c.font = Font(name='Arial', size=11, bold=True)
        c.fill = PatternFill('solid', fgColor='E7E6E6')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[4].height = 32
    for row in ws.iter_rows(min_row=5):
        for c in row:
            c.font = Font(name='Arial', size=11)
            c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[row[0].row].height = 42
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = 'A5'; ws.auto_filter.ref = f'A4:{end}{ws.max_row}'
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
    return ws


def note(ws, text):
    row = ws.max_row + 2; end = get_column_letter(ws.max_column)
    ws.merge_cells(f'A{row}:{end}{row}'); ws.cell(row, 1, text)
    ws.cell(row, 1).font = Font(name='Arial', size=11)
    ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[row].height = 36


def main():
    eth = csv_rows(ROOT/'results/eth_trees/tuning_20261025/fresh_summary.csv')
    ext = csv_rows(ROOT/'results/liang_trees/transfer_20261025/summary.csv')
    rates = {(r['method'], r['solver']): int(r['fresh_success'])/int(r['trials']) for r in eth}
    base = rates['fresh_control', 'svd']
    coverage = {v: stats.mean(float(r['mean_reference_match_rate']) for r in ext if r['variant'] == v)
                for v in ['control', 'vertical085', 'ransac999']}

    def change(before, after):
        return f'{before:.1%} → {after:.1%}\n{100*(after-before):+.1f} 個百分點'

    wb = Workbook(); wb.remove(wb.active)
    rows = [
        ['只放寬匹配距離', '匹配閾值：5 → 10 cm\n仍用 SVD 求解',
         change(base, rates['fresh_matching', 'svd']), '未測：缺少分站與配準真值', '變好；容許較大的樹位偏差'],
        ['只加配準 RANSAC', '匹配仍為 5 cm\n加上排除錯誤配對的步驟',
         change(base, rates['fresh_control', 'svd+ransac']), '未測：缺少分站與配準真值', '小幅變好'],
        ['兩項一起用（推薦）', '匹配閾值 10 cm\n＋配準 RANSAC',
         change(base, rates['fresh_matching', 'svd+ransac']), '未測：缺少分站與配準真值', '本輪兩站方法最好；先採用此組合'],
        ['放寬樹幹垂直度', '垂直度門檻：0.90 → 0.85\n讓更多不那麼垂直的點進入提取',
         change(base, rates['fresh_vertical085', 'svd']), change(coverage['control'], coverage['vertical085']),
         '兩份資料的指標都變差；保留 0.90'],
        ['加強圓柱擬合', 'RANSAC目標機率：0.99 → 0.999\n迭代上限：1000 → 3000',
         change(base, rates['fresh_ransac999', 'svd']), change(coverage['control'], coverage['ransac999']),
         'ETH略降、Liang微升；先保留原設定'],
        ['改走六站圖路徑', '由獨立兩站，改成使用六站\n透過品質較好的中間站連接',
         f'六站方法：{rates["graph_mst", "graph_mst"]:.1%}\n使用更多資料，條件不同', '未測：缺少原始分站資料',
         '適合有全部站點時使用；另作多站任務']]
    ws = sheet(wb, '調參重點',
               'ETH看「配準成功率」；Liang看「找到參考樹的比例」。數字越高越好，兩種指標不能互相比高低。ETH前五列均與5 cm＋SVD基準比較。',
               ['調整項目', '改了什麼', 'ETH：配準成功率', 'Liang：參考樹匹配比例', '目前建議'],
               rows, [24, 39, 34, 34, 43])
    for i in range(5, 11): ws.row_dimensions[i].height = 64
    best = rates['fresh_matching', 'svd+ransac']
    note(ws, '加在推薦組也沒有更好：再放寬垂直度，' +
         f'{best:.1%} → {rates["fresh_vertical085_matching", "svd+ransac"]:.1%}；再加強圓柱擬合，' +
         f'{best:.1%} → {rates["fresh_ransac999_matching", "svd+ransac"]:.1%}。')
    note(ws, '範圍：ETH為同一森林、15個配對×10個seed；Liang為處理後六樣區等權平均。尚未證明配準可跨資料集泛化。這些是實驗設定，程式預設值未改。')

    full = load_workbook(OUT/'comparison.xlsx', data_only=True, read_only=True)
    raw = list(full['逐對結果'].values); headers = raw[0]
    records = [dict(zip(headers, r)) for r in raw[1:]]
    pairs = sorted({r['pair'] for r in records})
    pair_rows = []
    for pair in pairs:
        old = [r for r in records if r['pair']==pair and r['方法代碼']=='fresh_control' and r['solver']=='svd']
        new = [r for r in records if r['pair']==pair and r['方法代碼']=='fresh_matching' and r['solver']=='svd+ransac']
        assert len(old)==len(new)==10
        a=sum(int(r['成功']) for r in old); b=sum(int(r['成功']) for r in new)
        pair_rows.append([pair, a, b, b-a, '變好' if b>a else ('變壞' if b<a else '相同')])
    full.close()
    ws=sheet(wb, 'ETH哪些配對改善',
             '只比較「原本5 cm＋SVD」與「推薦10 cm＋配準RANSAC」。每個配對都試10次；下表是成功次數，最高為10。',
             ['配對', '原本：成功幾次／10', '推薦組：成功幾次／10', '多成功幾次', '變好或變壞'], pair_rows, [20, 28, 30, 24, 24])
    for i in range(5, ws.max_row+1): ws.row_dimensions[i].height=24
    note(ws, '成功＝旋轉誤差小於2度，且平移誤差小於0.5公尺。兩組皆使用同一組seed11–20；不是十片獨立森林。')

    plots=[]
    for plot in range(1,7):
        r={x['variant']:float(x['mean_reference_match_rate']) for x in ext if int(x['plot'])==plot}
        plots.append([f'Plot {plot}', r['control'], r['vertical085'], 100*(r['vertical085']-r['control']),
                      r['ransac999'], 100*(r['ransac999']-r['control'])])
    ws=sheet(wb, 'Liang哪些樣區變化',
             '這張只看樹木提取：找到的參考樹÷提供的參考樹，越高越好。每個樣區平均10次。差值以百分點表示；Plot1、3、5是論文使用的樣區。',
             ['樣區', '原本設定', '垂直度0.85', '垂直度調整差值', '加強圓柱擬合', '圓柱擬合調整差值'], plots, [18, 23, 24, 27, 25, 29])
    for i in range(5, 11):
        for j in [2,3,5]: ws.cell(i,j).number_format='0.0%'
        for j in [4,6]: ws.cell(i,j).number_format='+0.0" 百分點";-0.0" 百分點";0.0" 百分點"'
    note(ws, '放寬垂直度在六個樣區都變差；加強圓柱擬合的差異很小。匹配門檻固定為0.5公尺，未針對這份新資料重新調參。')
    note(ws, '限制：這是公開處理後點雲，標註範圍與真實掃描中心未完全核實；沒有各站配準真值，因此不能用本表判斷10 cm配準設定能否泛化。')
    target=OUT/'comparison_simple.xlsx'; wb.save(target)
    print(target)


if __name__=='__main__': main()
