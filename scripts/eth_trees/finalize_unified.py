#!/usr/bin/env python3
"""Validate and publish this benchmark round after all computations finish."""
import fcntl
import hashlib
import json
import math
from pathlib import Path
import subprocess
import time
import urllib.request
import zipfile
from openpyxl import load_workbook
import unified_study as s

def ready():
    if len(list((s.RUNS/'v3_evaluation').glob('*/*.csv.complete')))!=55:return False
    for filename,count in [('cold_benchmark.json',12),('feature_benchmark.json',26)]:
        p=s.OUT/filename
        if not p.exists():return False
        try:
            if len(json.loads(p.read_text()))!=count:return False
        except json.JSONDecodeError:return False
    return True

def compare(raw,actual):
    if raw=='':return actual is None
    try:
        expected=float(raw)
        return (isinstance(actual,(float,int)) and math.isclose(expected,actual,rel_tol=1e-12,abs_tol=1e-12)) if math.isfinite(expected) else actual is None
    except (ValueError,TypeError):return raw==actual

def validate():
    meta=json.loads((s.OUT/'validation.json').read_text());rows=s.read_csv(s.OUT/'registration_rows.csv')
    assert len(rows)==4950
    assert len({(r['scene'],r['pair'],r['seed'],r['method']) for r in rows})==4950
    methods={r['method'] for r in rows};assert len(methods)==9
    for method in methods:assert sum(r['method']==method for r in rows)==550
    for name,digest in meta['files'].items():assert s.sha(s.OUT/name)==digest
    for record in json.loads((s.OUT/'feature_benchmark.json').read_text()):
        assert record['exit_code']==0
    frozen=json.loads((s.OUT/'method_v3_frozen.json').read_text())
    for name,digest in frozen['sha256'].items():assert s.sha(s.ROOT/name)==digest
    books=[]
    for name in ['comparison_full.xlsx','comparison_simple.xlsx']:
        workbook=load_workbook(s.OUT/name)
        assert not any(c.data_type in ('e','f') for sheet in workbook for row in sheet for c in row)
        assert all(sheet.freeze_panes=='A2' and sheet.auto_filter.ref for sheet in workbook)
        if name=='comparison_full.xlsx':
            sheet=workbook['全部逐次結果'];assert sheet.max_row==4951
            keys=[c.value for c in sheet[1]];assert keys==list(rows[0])
            assert all(compare(r[k],v) for got,r in zip(sheet.iter_rows(min_row=2,values_only=True),rows) for k,v in zip(keys,got))
            sheet=workbook['逐場景比較']
            assert sheet.max_row==46
            for row in sheet.iter_rows(min_row=2,values_only=True):assert math.isclose(row[5],row[3]/row[4],abs_tol=1e-12)
        books.append(dict(file=name,bytes=(s.OUT/name).stat().st_size,sha256=s.sha(s.OUT/name),sheets=[dict(name=sheet.title,rows=sheet.max_row,columns=sheet.max_column) for sheet in workbook]))
    private_ids=[]
    listing=s.ROOT/'data/tongji_trees/folder_listing.json'
    if listing.exists():private_ids.extend(str(row[0]).encode() for row in json.loads(listing.read_text()))
    private=s.ROOT/'data/unified/private_sources.json'
    if private.exists():private_ids.append(json.loads(private.read_text())['google_drive_read_only_folder'].rstrip('/').split('/')[-1].encode())
    for path in s.OUT.iterdir():
        if not path.is_file():continue
        if path.suffix=='.xlsx':
            with zipfile.ZipFile(path) as z:content=b''.join(z.read(n) for n in z.namelist())
        else:content=path.read_bytes()
        assert not any(x in content for x in private_ids),path
        assert b'C:\\Users\\' not in content,path
    result=dict(final_trials=4950,unique_pair_seed_methods=True,full_workbook_rows_verified=4950,scenes=5,datasets=2,private_identifiers_checked=len(private_ids),private_ids_absent=True,workbooks=books,cold_benchmarks=12,station_cost_cases=26)
    (s.OUT/'workbook_validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    return result

def main():
    while not ready():time.sleep(20)
    subprocess.run(['python','scripts/eth_trees/unified_report.py'],cwd=s.ROOT,check=True)
    validation=validate()
    flag=s.OUT/'FINAL_COMPLETE.json'
    # This stops periodic checkpointing, not the research acceptance checks.
    flag.write_text(json.dumps(dict(round_complete=True,research_project_complete=False,note='Strict author baseline and unseen-forest validation remain unverified; see report for any additional unmet numerical goals',validation_sha256=s.sha(s.OUT/'workbook_validation.json')),indent=2)+'\n')
    with (s.RUNS/'checkpoint-git.lock').open('a') as lock:
        fcntl.flock(lock,fcntl.LOCK_EX)
        subprocess.run(['git','add','--','scripts/eth_trees','CMakeLists.txt',str(s.OUT.relative_to(s.ROOT))],cwd=s.ROOT,check=True)
        subprocess.run(['git','diff','--cached','--check'],cwd=s.ROOT,check=True,stdout=subprocess.DEVNULL)
        subprocess.run(['git','commit','-m','Publish audited unified forest benchmark, costs and both workbooks'],cwd=s.ROOT,check=True)
        subprocess.run(['git','push','origin','HEAD'],cwd=s.ROOT,check=True)
        commit=subprocess.check_output(['git','rev-parse','HEAD'],cwd=s.ROOT,text=True).strip();links={}
        for name in ['comparison_simple.xlsx','comparison_full.xlsx']:
            relative=str((s.OUT/name).relative_to(s.ROOT));url=f'https://raw.githubusercontent.com/WWW3113/CircularMatch/{commit}/{relative}'
            body=None
            for attempt in range(3):
                try:
                    with urllib.request.urlopen(url,timeout=30) as response:body=response.read()
                    assert hashlib.sha256(body).hexdigest()==s.sha(s.OUT/name)
                    break
                except Exception:
                    if attempt==2:raise
                    time.sleep(5)
            links[name]=url
        (s.OUT/'DELIVERY.json').write_text(json.dumps(dict(commit=commit,verified_downloads=links,workbook_validation_sha256=s.sha(s.OUT/'workbook_validation.json')),indent=2)+'\n')
        subprocess.run(['git','add','--',str((s.OUT/'DELIVERY.json').relative_to(s.ROOT))],cwd=s.ROOT,check=True)
        subprocess.run(['git','commit','-m','Record verified immutable workbook download links'],cwd=s.ROOT,check=True)
        subprocess.run(['git','push','origin','HEAD'],cwd=s.ROOT,check=True)
    print(json.dumps(links),flush=True)

if __name__=='__main__':main()
