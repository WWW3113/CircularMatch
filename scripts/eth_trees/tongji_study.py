#!/usr/bin/env python3
"""Fixed ETH methods on user-provided Tongji data; never write to Drive."""
import argparse
import concurrent.futures as futures
import csv
import fcntl
from datetime import datetime,timezone
import hashlib
import itertools
import json
import math
from pathlib import Path
import re
import statistics as stats
import subprocess
import threading
import time
import numpy as np
from openpyxl import Workbook,load_workbook
from make_study_workbook import table
from tune_registration import ROOT,TOOL,read_csv
from tuning_graph import ply_map

DATA=ROOT/'data/tongji_trees/received/Tongji-Trees'
OUT=ROOT/'results/tongji_trees/fixed_eth_20261025';RUNS=OUT/'runs'
SEEDS=list(range(31,41));LOCK=threading.Lock()
METHODS=[dict(name='reference5cm',threshold=.05,recovery=False),dict(name='baseline10cm',threshold=.10,recovery=False),dict(name='recovery10cm',threshold=.10,recovery=True)]
PAPER_EP={1:.052558,2:.114808,3:.048756,4:.04356} # Uploaded paper, Table I; reported values, not reruns.
def write_csv(path,rows):
    with Path(path).open('w',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(rows[0]),lineterminator='\n');writer.writeheader();writer.writerows(rows)
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def parallel(fn,jobs):
    with futures.ThreadPoolExecutor(max_workers=4) as pool:return list(pool.map(fn,jobs))
def execute(label,args):
    log=RUNS/'logs'/f'{label}.log';log.parent.mkdir(parents=True,exist_ok=True)
    cmd=([str(TOOL)] if TOOL else [])+list(map(str,args));start=time.monotonic();print('START',label,flush=True)
    with log.open('w') as f:r=subprocess.run(cmd,cwd=ROOT,stdout=f,stderr=subprocess.STDOUT)
    with LOCK:
        with (RUNS/'commands.jsonl').open('a') as f:f.write(json.dumps(dict(label=label,command=cmd,exit_code=r.returncode,seconds=time.monotonic()-start))+'\n')
    if r.returncode:raise RuntimeError(str(log))
    print('DONE',label,round(time.monotonic()-start,1),flush=True)
def freeze():
    OUT.mkdir(parents=True,exist_ok=True);RUNS.mkdir(exist_ok=True)
    pp=[];inputs=[]
    for plot in range(1,5):
        directory=DATA/f'Plot{plot}';text=(directory/f'GroundTruthMatrices-Plot{plot}.txt').read_text()
        absolute={1:np.eye(4)}
        for scan in range(2,6):absolute[scan]=np.loadtxt(directory/f'S{scan}_S1.txt')
        for a,b in itertools.combinations(range(1,6),2):
            source,target=b,a;path=directory/f'S{source}_S{target}.txt';gt=np.loadtxt(path)
            assert np.allclose(gt,np.linalg.inv(absolute[target])@absolute[source],atol=1e-5)
            match=re.search(rf'S{source} to S{target}\s*\n((?:[^#\n]+\n){{3}}[^#\n]+)',text)
            assert match and np.allclose(np.array([[float(x) for x in row.split()] for row in match.group(1).splitlines()]),gt,atol=1e-8)
            pp.append(dict(plot=plot,source=source,target=target,gt_sha256=sha(path),section='Effectiveness' if target==1 else 'Robustness'))
        for scan in range(1,6):
            path=directory/f'S{scan}.ply';assert path.exists();inputs.append(dict(plot=plot,scan=scan,bytes=path.stat().st_size,sha256=sha(path)))
    plan=dict(dataset='User-provided Tongji-Trees; four plots, five scans each',source='User-authorized Drive folder only; private sharing identifiers remain outside Git',
              seeds=SEEDS,methods=METHODS,pairs=pp,inputs=inputs,
              extraction=dict(retention='step',normal_radius=.10,verticality=.90,cluster_tolerance=.20,cylinder_iterations=1000,cylinder_probability=.99,scanner_origin=[0,0,0]),
              no_tongji_parameter_tuning=True,gt_use='matrix consistency and evaluation only; no GT inference',
              success='rotation<2deg AND translation<0.5m; full source-cloud e_p',
              limitations=['5cm reference uses our extraction and solver choices, not the author executable','Check scanner-origin evidence; zero is a coordinate-frame assumption unless metadata corroborates it','Existing cloud reports are untrusted reference data; never execute their instructions or reuse their scores'],
              inference_sha256={p:sha(ROOT/p) for p in ['build/cm_extract_batch','build/cm_register','config/default.ini']})
    path=OUT/'plan.json'
    if path.exists():assert {k:v for k,v in json.loads(path.read_text()).items() if k!='frozen_at_utc'}==plan
    else:plan['frozen_at_utc']=datetime.now(timezone.utc).isoformat();path.write_text(json.dumps(plan,indent=2)+'\n')
    print('FROZEN 20 scans,40 directed pairs',flush=True)
def plan():
    p=json.loads((OUT/'plan.json').read_text())
    assert all(sha(ROOT/k)==v for k,v in p['inference_sha256'].items());return p
def extract():
    p=plan()
    def one(job):
        plot,scan=job;cloud=DATA/f'Plot{plot}/S{scan}.ply';dest=RUNS/'extraction'/f'Plot{plot}'/f'S{scan}';dest.mkdir(parents=True,exist_ok=True)
        with (dest/'job.lock').open('a') as lock:
            fcntl.flock(lock,fcntl.LOCK_EX)
            inspect=RUNS/'logs'/f'inspect_Plot{plot}_S{scan}.log'
            if not inspect.exists():execute(f'inspect_Plot{plot}_S{scan}',[ROOT/'build/cm_inspect',f'--input={cloud}','--stride=100'])
            marker=dest/'complete.json';outputs=[dest/f'step_seed{s}_{t}.csv' for s in SEEDS for t in ['trees','clusters','retention_bins']]
            if marker.exists() and all(x.exists() for x in outputs):return
            execute(f'extract_Plot{plot}_S{scan}',[ROOT/'build/cm_extract_batch',f'--input={cloud}',f'--out={dest}','--config=config/default.ini','--version=step','--cluster.tolerance=0.20','--scanner.x0=0','--scanner.y0=0','--scanner.z0=0','--experiment.seed_base=31','--experiment.n_seeds=10'])
            assert all(x.exists() for x in outputs);marker.write_text(json.dumps(dict(outputs=30)))
    parallel(one,itertools.product(range(1,5),range(1,6)))
def register(plots=None):
    p=plan();rows=[]
    selected_plots=list(range(1,5)) if plots is None else plots
    for method in METHODS:
        def one(job):
            plot,source=job;dest=RUNS/'registration'/method['name']/f'Plot{plot}';dest.mkdir(parents=True,exist_ok=True);jobs=[];files=[]
            for pair in p['pairs']:
                if pair['plot']!=plot or pair['source']!=source:continue
                target=pair['target']
                for seed in SEEDS:
                    t=RUNS/'extraction'/f'Plot{plot}';output=dest/f'{seed}_S{source}-S{target}.csv'
                    jobs.append(f'{t}/S{source}/step_seed{seed}_trees.csv {t}/S{target}/step_seed{seed}_trees.csv {DATA}/Plot{plot}/S{source}.ply {DATA}/Plot{plot}/S{source}_S{target}.txt {output} Plot{plot}_{seed}_S{source}-S{target}')
                    files.append((seed,target,pair['section'],output))
            if not jobs:return []
            batch=dest/f'S{source}.txt';batch.write_text('\n'.join(jobs)+'\n');marker=dest/f'S{source}.complete'
            args=[ROOT/'build/cm_register',f'--batch={batch}',f'--thresholds={method["threshold"]}']
            if method['recovery']:args+=['--recovery=evidence','--recovery_trigger=6']
            with (dest/f'S{source}.lock').open('a') as lock:
                fcntl.flock(lock,fcntl.LOCK_EX)
                if not marker.exists() or not all(x.exists() for _,_,_,x in files):
                    execute(f'reg_{method["name"]}_P{plot}_S{source}',args);marker.write_text('completed\n')
            return [dict(method=method['name'],plot=plot,seed=seed,pair=f'S{source}-S{target}',section=section,**r) for seed,target,section,path in files for r in read_csv(path)]
        rows.extend(r for rr in parallel(one,itertools.product(selected_plots,range(2,6))) for r in rr)
    name='rows.csv' if plots is None else 'partial_'+ '-'.join(map(str,plots))+'.csv'
    h=list(dict.fromkeys(k for r in rows for k in r));write_csv(RUNS/name,[{k:r.get(k,'') for k in h} for r in rows]);print('ROWS',len(rows),flush=True)
def independent_ep_audit(rows):
    """Recompute one first valid primary case per plot over every source point."""
    audits=[]
    for plot in range(1,5):
        selected=[r for r in rows if int(r['plot'])==plot and r['method']=='recovery10cm' and r['solver']=='svd+ransac' and r['transform_ok']=='1']
        if not selected:continue
        r=selected[0];source,target=r['pair'].split('-');cloud=ply_map(DATA/f'Plot{plot}'/f'{source}.ply')
        gt=np.loadtxt(DATA/f'Plot{plot}'/f'{source}_{target}.txt')
        rotation=np.array([[float(r[f'r{i}{j}']) for j in range(3)] for i in range(3)])
        translation=np.array([float(r[f't{i}']) for i in range(3)])
        assert np.allclose(rotation.T@rotation,np.eye(3),atol=1e-8)
        assert abs(np.linalg.det(rotation)-1)<1e-8
        total=0.;count=0
        for start in range(0,len(cloud),500000):
            block=cloud[start:start+500000];xyz=np.column_stack([block[k] for k in ['x','y','z']]).astype('f8')
            assert np.isfinite(xyz).all()
            total+=np.linalg.norm(xyz@(rotation-gt[:3,:3]).T+translation-gt[:3,3],axis=1).sum();count+=len(xyz)
        calculated=total/count;native=float(r['e_p_m']);assert abs(calculated-native)<1e-6
        audits.append(dict(plot=plot,pair=r['pair'],seed=int(r['seed']),method=r['method'],solver=r['solver'],source_points=count,native_ep_m=native,independent_full_cloud_ep_m=calculated,absolute_difference_m=abs(calculated-native)))
    (OUT/'independent_ep_audit.json').write_text(json.dumps(audits,indent=2)+'\n')
    return audits
def diagnose():
    """Post-hoc examples only; never feed GT residuals into inference."""
    plan();dest=RUNS/'diagnostics/plot2_height';dest.mkdir(parents=True,exist_ok=True);jobs=[]
    for source,target in [(3,1),(4,2)]:
        base=RUNS/'extraction/Plot2';label=f'Plot2_seed31_S{source}-S{target}'
        jobs.append(f'{base}/S{source}/step_seed31_trees.csv {base}/S{target}/step_seed31_trees.csv {DATA}/Plot2/S{source}.ply {DATA}/Plot2/S{source}_S{target}.txt {dest}/{label}.csv {label}')
    batch=dest/'jobs.txt';batch.write_text('\n'.join(jobs)+'\n')
    execute('diagnostic_Plot2_height',[ROOT/'build/cm_register',f'--batch={batch}','--thresholds=0.10','--skip_point_error=1',f'--diagnostics={dest}'])
    summaries=[]
    for source,target in [(3,1),(4,2)]:
        match=[r for r in read_csv(dest/f'Plot2_seed31_S{source}-S{target}_thr0.100000_descriptors.csv') if r['accepted']=='1']
        a=read_csv(RUNS/f'extraction/Plot2/S{source}/step_seed31_trees.csv');b=read_csv(RUNS/f'extraction/Plot2/S{target}/step_seed31_trees.csv')
        ps=np.array([[float(a[int(r['source'])][k]) for k in ['x','y','z']] for r in match]);pt=np.array([[float(b[int(r['target'])][k]) for k in ['x','y','z']] for r in match])
        gt=np.loadtxt(DATA/f'Plot2/S{source}_S{target}.txt');delta=ps@gt[:3,:3].T+gt[:3,3]-pt
        summaries.append(dict(plot=2,seed=31,pair=f'S{source}-S{target}',matches=len(match),median_gt_xy_m=float(np.median(np.linalg.norm(delta[:,:2],axis=1))),median_abs_gt_z_m=float(np.median(abs(delta[:,2]))),median_gt_xyz_m=float(np.median(np.linalg.norm(delta,axis=1)))))
    write_csv(OUT/'posthoc_height_diagnostic.csv',summaries)
def report():
    p=plan();rows=read_csv(RUNS/'rows.csv');assert len(rows)==2400
    inspections=[];extractions=[]
    for plot,scan in itertools.product(range(1,5),range(1,6)):
        log=(RUNS/'logs'/f'inspect_Plot{plot}_S{scan}.log').read_text()
        count,skipped=map(int,re.search(r'points read: (\d+), skipped \(non-finite/unparsable\): (\d+)',log).groups())
        assert skipped==0 and 'matches read+skipped' in log
        origin=re.search(r'0 \(candidate\)\s+(\d+)\s+([\d.]+)\s+([\d.]+)',log)
        offsets=[float(x) for x in re.findall(r'^\s*[xy] [+-][\d.]+\s+\d+\s+[\d.]+\s+([\d.]+)\s*$',log,re.M)]
        assert origin and offsets
        inspections.append(dict(plot=plot,scan=scan,points=count,skipped=skipped,zero_origin_azimuth_top_share=float(origin.group(3)),best_tested_xy_offset_azimuth_top_share=max(offsets),zero_minus_best_xy_offset=float(origin.group(3))-max(offsets)))
        extraction_log=(RUNS/'logs'/f'extract_Plot{plot}_S{scan}.log').read_text()
        assert '\0' not in extraction_log
        assert set(map(int,re.findall(r'completed seed (\d+)',extraction_log)))==set(SEEDS)
        marker=RUNS/'extraction'/f'Plot{plot}'/f'S{scan}'/'complete.json';assert marker.exists()
        for seed in SEEDS:
            path=marker.parent/f'step_seed{seed}_trees.csv';trees=read_csv(path)
            assert all(math.isfinite(float(tree[c])) for tree in trees for c in ['x','y','z'])
            extractions.append(dict(plot=plot,scan=scan,seed=seed,trees=len(trees),sha256=sha(path)))
    write_csv(OUT/'inspection_summary.csv',inspections)
    write_csv(OUT/'extraction_summary.csv',extractions)
    audits=independent_ep_audit(rows)
    summary=[]
    identity=lambda r:(int(r['plot']),int(r['seed']),r['pair'],r['method'],r['solver'])
    assert len({identity(r) for r in rows})==2400,'duplicate cases'
    for plot,method,solver in itertools.product([0,1,2,3,4],METHODS,['svd','svd+ransac']):
        rr=[r for r in rows if (plot==0 or int(r['plot'])==plot) and r['method']==method['name'] and r['solver']==solver];assert len(rr)==(400 if plot==0 else 100)
        good=[float(r['e_p_m']) for r in rr if r['success']=='1']
        for r in rr:assert int(r['success'])==int(r['transform_ok']=='1' and float(r['rot_err_deg'])<2 and float(r['trans_err_m'])<.5)
        summary.append(dict(plot=plot,method=method['name'],solver=solver,cases=len(rr),success=sum(int(r['success']) for r in rr),rate=sum(int(r['success']) for r in rr)/len(rr),mean_success_ep_m=stats.mean(good) if good else '',median_success_ep_m=stats.median(good) if good else ''))
    write_csv(OUT/'summary.csv',summary)
    sections=[]
    for plot,method,section in itertools.product(range(1,5),METHODS,['Effectiveness','Robustness']):
        solver='svd' if method['name']=='reference5cm' else 'svd+ransac'
        rr=[r for r in rows if int(r['plot'])==plot and r['method']==method['name'] and r['solver']==solver and r['section']==section]
        assert len(rr)==(40 if section=='Effectiveness' else 60)
        sections.append(dict(plot=plot,method=method['name'],section=section,cases=len(rr),success=sum(int(r['success']) for r in rr),rate=sum(int(r['success']) for r in rr)/len(rr)))
    write_csv(OUT/'section_summary.csv',sections)
    names={'reference5cm':'5cm參考（我們實作）','baseline10cm':'10cm＋配準RANSAC','recovery10cm':'10cm＋配準RANSAC＋三維補救'}
    primary_solver=lambda method:'svd' if method=='reference5cm' else 'svd+ransac'
    keyed={(int(r['plot']),int(r['seed']),r['pair'],r['method']):r for r in rows if r['solver']==primary_solver(r['method'])}
    paired=[]
    for plot,(old,new) in itertools.product([0,1,2,3,4],[('reference5cm','baseline10cm'),('baseline10cm','recovery10cm'),('reference5cm','recovery10cm')]):
        comparisons=[(r,keyed[(*key[:3],new)]) for key,r in keyed.items() if key[3]==old and (plot==0 or key[0]==plot)]
        shared=[(a,b) for a,b in comparisons if a['success']==b['success']=='1']
        paired.append(dict(plot=plot,old=old,new=new,cases=len(comparisons),old_success=sum(int(a['success']) for a,b in comparisons),new_success=sum(int(b['success']) for a,b in comparisons),gained=sum(a['success']=='0' and b['success']=='1' for a,b in comparisons),lost=sum(a['success']=='1' and b['success']=='0' for a,b in comparisons),shared_success=len(shared),mean_ep_change_m=stats.mean(float(b['e_p_m'])-float(a['e_p_m']) for a,b in shared) if shared else ''))
    write_csv(OUT/'paired_comparison.csv',paired)
    pairs=[]
    for pair,method in itertools.product(p['pairs'],METHODS):
        rr=[keyed[(pair['plot'],seed,f'S{pair["source"]}-S{pair["target"]}',method['name'])] for seed in SEEDS]
        good=[float(r['e_p_m']) for r in rr if r['success']=='1']
        pairs.append(dict(plot=pair['plot'],pair=f'S{pair["source"]}-S{pair["target"]}',section=pair['section'],method=method['name'],cases=10,success=sum(int(r['success']) for r in rr),mean_success_ep_m=stats.mean(good) if good else ''))
    write_csv(OUT/'pair_summary.csv',pairs)
    write_csv(OUT/'registration_rows.csv',rows)
    cross=[]
    eth_path=ROOT/'results/eth_trees/failure_followup_20261025/comparison_full.xlsx'
    park_path=ROOT/'results/whu_park/transfer_20261025/summary.csv'
    if eth_path.exists():
        book=load_workbook(eth_path,read_only=True,data_only=True)
        values=book['新驗證逐次'].iter_rows(values_only=True);headers=next(values)
        eth=[dict(zip(headers,row)) for row in values];book.close()
        for method in ['baseline','recovery']:
            rr=[r for r in eth if r['method']==method and r['solver']=='svd+ransac'];assert len(rr)==150
            ep=[float(r['e_p_m']) for r in rr if int(r['success'])==1]
            cross.append(['ETH原始森林','10cm＋配準RANSAC'+('＋三維補救' if method=='recovery' else ''),'固定ETH參數',6,15,'21–30',150,sum(int(r['success']) for r in rr),stats.mean(ep)*100 if ep else None,'ETH開發後保留seed驗證；同一片森林，不是新森林'])
    if park_path.exists():
        for r in read_csv(park_path):
            if r['solver']!='svd+ransac':continue
            cross.append(['WHU-Park處理版','10cm＋配準RANSAC'+('＋三維補救' if r['method']=='recovery' else ''),r['preprocessing'],32,31,'21–30',int(r['cases']),int(r['success']),100*float(r['mean_ep_success_m']) if r['mean_ep_success_m'] else None,'公開YOHO處理版；不是論文WHU-FGI；密度適配僅法向半徑10→20cm'])
    simple=Workbook();simple.remove(simple.active)
    primary=[r for r in summary if r['solver']==('svd' if r['method']=='reference5cm' else 'svd+ransac')]
    for r in primary:
        if r['plot']==0 and r['method']!='reference5cm':cross.append(['Tongji-Trees四樣區',names[r['method']],'固定ETH參數',20,40,'31–40',r['cases'],r['success'],100*r['mean_success_ep_m'] if r['mean_success_ep_m']!='' else None,'新資料直接測試；未依Tongji真值選參數'])
    write_csv(OUT/'cross_dataset_summary.csv',[dict(zip(['dataset','method','preprocessing','scans','pairs','seeds','cases','success','mean_success_ep_cm','limitations'],r)) for r in cross])
    plot_name=lambda plot:'全部樣區' if plot==0 else f'Plot{plot}'
    ws=table(simple,'先看結果',['樣區','方法','成功數','案例數','成功率','成功案例平均點雲誤差cm'],[[plot_name(r['plot']),names[r['method']],r['success'],r['cases'],r['rate'],100*r['mean_success_ep_m'] if r['mean_success_ep_m']!='' else None] for r in primary],[16,52,17,17,20,35])
    for i in range(2,ws.max_row+1):ws.cell(i,5).number_format='0.0%'
    notes=[['資料','4樣區、20站、40個提供的配對；每對seed31–40，共400個主要案例，未針對Tongji真值調參。'],['站與配對','站是一次掃描；本次每樣區5站，每次配準只取其中2站，沒有用5站共同修正。'],['距離抽樣','沿用論文階梯step：水平距離<5m留50%、5–10m留75%、10m以上全留；每站10個seed重複抽樣。'],['提取','沿用ETH：法向半徑10cm、垂直度0.90、分群20cm、圓柱機率0.99及最多1000次；三組共用同一批樹位。'],['調整1：匹配容差','5cm→10cm，允許樹位間距有較大誤差，也可能加入錯配。指匹配階段，不是點雲降採樣間距。'],['調整2：解變換','SVD→SVD加配準RANSAC，嘗試排除錯配；與樹幹提取時的圓柱RANSAC是不同步驟。'],['調整3：三維補救','初始配對少於6時，從描述子候選建立三維剛體候選，再要求至少6個一對一支持點；不讀取GT。'],['組合比較','精簡表比較5cm/SVD、10cm/RANSAC、10cm/RANSAC加補救；完整表另列兩種solver，可分開看容差與RANSAC。'],['原作差異','5cm參考含我們的DTM及額外擬合檢查，不是原作者程式；論文表I數值另標示為報告值。'],['成功標準','旋轉<2°且平移<0.5m；與ETH和Park相同。點雲誤差為source全部點的平均變換差，只對成功案例彙總。'],['改善退步怎麼看','新增成功越多、退步失敗越少越好；誤差差值=新方法−舊方法，負數較好，僅比較兩法都成功的同一案例。'],['重複次數限制','400案例是40對×10次抽樣，不是400個獨立場景；共用站點的案例也不是獨立樣本。'],['原點','採檔案零原點；保留逐站角度格網檢查供核對，未用GT挑原點；沒有獨立站位metadata。'],['來源限制','僅讀取指定雲端資料夾，未修改雲端資料；私人分享識別碼未寫入Git。']]
    table(simple,'用什麼設定',['項目','說明'],notes,[27,140])
    def comparison_sheet(wb):
        return table(wb,'改善與退步',['樣區','原方法','新方法','案例數','原成功數','新成功數','新增成功','退步失敗','兩法都成功','平均誤差變化cm（負數較好）'],[[plot_name(r['plot']),names[r['old']],names[r['new']],r['cases'],r['old_success'],r['new_success'],r['gained'],r['lost'],r['shared_success'],100*r['mean_ep_change_m'] if r['mean_ep_change_m']!='' else None] for r in paired],[16,48,52,16,16,16,16,16,20,35])
    comparison_sheet(simple)
    def cross_sheet(wb):
        ws=table(wb,'跨資料集比較',['資料集','方法','前處理','站數','配對數','seed','案例數','成功數','成功率','成功案例平均誤差cm','解讀限制'],[r[:8]+[r[7]/r[6]]+r[8:] for r in cross],[26,52,24,14,16,16,16,16,20,32,100])
        for i in range(2,ws.max_row+1):ws.cell(i,9).number_format='0.0%'
    cross_sheet(simple)
    paper_rows=[[f'Plot{plot}',100,PAPER_EP[plot]*100,r['success'],r['cases'],r['rate'],r['mean_success_ep_m']*100 if r['mean_success_ep_m']!='' else None,'論文表I；配對及成功判準未完全核對，不能直接宣稱超越'] for plot in range(1,5) for r in primary if r['plot']==plot and r['method']=='recovery10cm']
    def paper_sheet(wb):
        ws=table(wb,'與論文報告值並列',['樣區','論文成功率%','論文點雲誤差cm','目前成功數','目前案例數','目前成功率','目前成功案例平均誤差cm','比較限制'],paper_rows,[16,22,26,20,20,22,34,100])
        for i in range(2,ws.max_row+1):ws.cell(i,6).number_format='0.0%'
    paper_sheet(simple)
    ws=table(simple,'原資料兩組配對',['樣區','方法','原始分類','案例數','成功數','成功率'],[[r['plot'],names[r['method']],r['section'],r['cases'],r['success'],r['rate']] for r in sections],[16,52,24,18,18,20])
    for i in range(2,ws.max_row+1):ws.cell(i,6).number_format='0.0%'
    simple.save(OUT/'comparison_simple.xlsx')
    full=Workbook();full.remove(full.active);table(full,'設定與限制',['項目','說明'],notes,[27,140])
    comparison_sheet(full);paper_sheet(full);cross_sheet(full)
    table(full,'欄位怎麼讀',['欄位','說明'],[
        ['method / solver','method是5cm參考、10cm基準或10cm補救；solver是SVD或SVD加配準RANSAC。'],
        ['plot / pair / seed','樣區、source到target方向、抽樣亂數；同一組三欄可比較不同方法。plot=0代表全部四樣區。'],
        ['setting=paper','程式內部的設定名稱，不表示原作者程式，也不表示完全重現論文。'],
        ['section','沿用提供的GT文字分組：Effectiveness為S2–S5到S1四對；Robustness為其餘六對。'],
        ['threshold','CN匹配距離容差，單位m；不是體素邊長。'],
        ['n_src_kp / n_tgt_kp','兩站提取的樹位候選數；沒有人工樹位標註，不能當作實際樹木數或提取召回率。'],
        ['n_matches / n_correct_matches','用於配準的配對數／GT事後評估水平XY距離<30cm的配對數；不代表高度也正確，GT沒有用來選配對。'],
        ['ransac_inliers','配準RANSAC保留的支持點數；SVD沒有這個步驟。'],
        ['transform_ok / success','找到變換／通過旋轉與平移標準，1表示是；找到變換仍可能配錯。'],
        ['rot_err_deg / trans_err_m','相對GT的旋轉誤差（度）與平移誤差（m），越小越好。'],
        ['e_p_m','將估計變換及GT分別施於source全部點，平均兩者位置差，單位m；無有限結果時Excel留空。'],
        ['time_ms','原生程式量測的匹配與解變換耗時，不含樹位提取、讀檔或全點誤差評估；不可與論文總時間直接比。'],
        ['r00–r22 / t0–t2','估計旋轉矩陣3×3及平移向量3×1，把source映到target。'],
        ['recovery_triggered / recovery_accepted','補救是否被嘗試／是否取代原配對；空白表示該方法沒開補救。'],
        ['gained / lost','原失敗變成功／原成功變失敗；相同樣區、配對、seed逐一比較。'],
        ['mean_ep_change_m','兩法都成功的案例，新方法減舊方法的平均誤差差值；負值較好。'],
        ['sha256','核對輸入或樹位CSV是否為同一份檔案；原始點雲與私人Drive識別碼未上傳Git。']
    ],[45,140])
    h=list(summary[0]);ws=table(full,'全部方法摘要',h,[[r[k] for k in h] for r in summary],[28]*len(h))
    for i in range(2,ws.max_row+1):ws.cell(i,h.index('rate')+1).number_format='0.0%'
    for title,records in [('原資料分組摘要',sections),('逐對結果',pairs),('逐站格式核驗',inspections),('逐站逐次樹位數',extractions)]:
        keys=list(records[0]);ws=table(full,title,keys,[[r[k] for k in keys] for r in records],[30]*len(keys))
        if 'rate' in keys:
            for i in range(2,ws.max_row+1):ws.cell(i,keys.index('rate')+1).number_format='0.0%'
    if audits:table(full,'全點誤差獨立抽查',list(audits[0]),[[r[k] for k in audits[0]] for r in audits],[30]*len(audits[0]))
    diagnostic=OUT/'posthoc_height_diagnostic.csv'
    if diagnostic.exists():
        records=read_csv(diagnostic);table(full,'事後高度診斷',list(records[0]),[[r[k] for k in records[0]] for r in records],[30]*len(records[0]))
    h=list(rows[0])
    def val(x):
        if x=='':return None
        try:n=float(x);return n if math.isfinite(n) else None
        except ValueError:return x
    table(full,'完整逐次結果',h,[[val(r[k]) for k in h] for r in rows],[35 if k in ['method','label','setting','solver'] else 22 for k in h]);table(full,'輸入檔案雜湊',list(p['inputs'][0]),[[r[k] for k in p['inputs'][0]] for r in p['inputs']],[18,18,25,85]);full.save(OUT/'comparison_full.xlsx')
    files=['comparison_full.xlsx','comparison_simple.xlsx','registration_rows.csv','summary.csv','paired_comparison.csv','pair_summary.csv','cross_dataset_summary.csv','extraction_summary.csv','independent_ep_audit.json']
    files.extend(n for n in ['posthoc_height_diagnostic.csv','extraction_repeat_audit.json'] if (OUT/n).exists())
    (OUT/'validation.json').write_text(json.dumps(dict(rows=2400,scans=20,pairs=40,seeds=SEEDS,extractions=200,finite_keypoints_verified=True,input_point_counts_verified=True,independent_full_cloud_ep_cases=len(audits),duplicate_cases=0,source_folder_read_only=True,no_target_parameter_tuning=True,analysis_script_sha256=sha(Path(__file__)),prior_result_sha256={str(path.relative_to(ROOT)):sha(path) for path in [eth_path,park_path] if path.exists()},files={n:sha(OUT/n) for n in files}),indent=2)+'\n')
    print(json.dumps(primary,indent=2),flush=True)
def main():
    a=argparse.ArgumentParser();a.add_argument('phase',choices=['freeze','extract','register','diagnose','report']);args=a.parse_args();globals()[args.phase]()
if __name__=='__main__':main()
